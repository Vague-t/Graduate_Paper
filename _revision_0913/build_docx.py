from pathlib import Path
import re, json
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
WORK=ROOT/'_revision_0913'
CN='Songti SC'; LAT='Times New Roman'; MATH='STIX Two Math'
FENCE=chr(96)*3
def el(tag,val=None):
    e=OxmlElement(tag)
    if val is not None: e.set(qn('m:val'),str(val))
    return e
def box(tag, children):
    e=el(tag)
    for c in children: e.append(c)
    return e
def mr(text,plain=False):
    e=el('m:r')
    if plain:
        pr=el('m:rPr'); pr.append(el('m:sty','p')); e.append(pr)
    t=el('m:t'); t.text=text; t.set(qn('xml:space'),'preserve'); e.append(t); return e
SYMBOLS={'alpha':'α','beta':'β','gamma':'γ','delta':'δ','eta':'η','nu':'ν','sigma':'σ','tau':'τ','theta':'θ','phi':'φ','psi':'ψ','Pi':'Π','Delta':'Δ','in':'∈','leq':'≤','geq':'≥','neq':'≠','times':'×','cdot':'·','mid':'∣','cup':'∪','succ':'≻','rightarrow':'→','ldots':'…','dots':'…','sum':'∑','prod':'∏'}
SYMBOLS['lambda']='λ'
CALLIGRAPHIC={'A':'𝒜','C':'𝒞','E':'ℰ','G':'𝒢','L':'ℒ','Q':'𝒬','S':'𝒮','T':'𝒯'}
DOUBLE={'R':'ℝ','I':'𝕀'}
class Tex:
    """Strict converter for the manuscript's known TeX vocabulary."""
    def __init__(self,s):
        self.t=re.findall(r'\\[A-Za-z]+|\\.|[{}_^]|.',s.replace('\n','')); self.i=0
    def peek(self): return self.t[self.i] if self.i<len(self.t) else ''
    def pop(self):
        t=self.peek(); self.i+=1; return t
    def skip(self):
        while self.peek() and self.peek().isspace(): self.pop()
    def group(self):
        self.skip()
        if self.peek()=='{':
            self.pop(); result=self.seq('}'); assert self.pop()=='}'; return result
        return self.atom()
    def literal_group(self):
        self.skip()
        if self.peek()!='{': return self.pop()
        self.pop(); depth=1; text=''
        while depth:
            t=self.pop()
            if not t: raise ValueError('Unclosed group')
            if t=='{': depth+=1
            if t=='}': depth-=1
            if depth: text+=t
        return text
    def seq(self,end=None):
        out=[]
        while self.peek() and self.peek()!=end:
            self.skip()
            if not self.peek() or self.peek()==end: break
            base=self.atom(); sub=sup=None; self.skip()
            while self.peek() in ('_','^'):
                which=self.pop(); arg=self.group()
                if which=='_': sub=arg
                else: sup=arg
                self.skip()
            if sub is not None or sup is not None:
                tag='m:sSubSup' if sub is not None and sup is not None else ('m:sSub' if sub is not None else 'm:sSup')
                n=el(tag); n.append(box('m:e',base))
                if sub is not None: n.append(box('m:sub',sub))
                if sup is not None: n.append(box('m:sup',sup))
                base=[n]
            out+=base
        return out
    def atom(self):
        t=self.pop()
        if t=='{':
            a=self.seq('}'); assert self.pop()=='}'; return a
        if not t.startswith('\\'): return [mr(t)]
        cmd=t[1:]
        if cmd in ('frac','dfrac'):
            n=el('m:f'); n.append(box('m:num',self.group())); n.append(box('m:den',self.group())); return [n]
        if cmd in ('operatorname','text','mathrm'): return [mr(self.literal_group(),True)]
        if cmd in ('mathcal','mathbb'):
            text=self.literal_group(); mapper=CALLIGRAPHIC if cmd=='mathcal' else DOUBLE
            return [mr(''.join(mapper.get(c,c) for c in text))]
        if cmd=='hat':
            n=el('m:acc'); pr=el('m:accPr'); pr.append(el('m:chr','̂')); n.append(pr); n.append(box('m:e',self.group())); return [n]
        if cmd=='left':
            begin=self.pop().replace('\\',''); parts=self.seq(r'\right')
            assert self.pop()==r'\right'; end=self.pop().replace('\\','')
            d=el('m:d'); pr=el('m:dPr'); pr.append(el('m:begChr','' if begin=='.' else begin)); pr.append(el('m:endChr','' if end=='.' else end)); d.append(pr); d.append(box('m:e',parts)); return [d]
        if cmd=='begin':
            env=self.literal_group()
            if env!='cases': raise ValueError('Unsupported environment '+env)
            raw=''
            while self.peek() and self.peek()!=r'\end': raw+=self.pop()
            assert self.pop()==r'\end' and self.literal_group()==env
            matrix=el('m:m'); mpr=el('m:mPr'); cols=el('m:mcs'); col=el('m:mc'); cpr=el('m:mcPr'); cpr.append(el('m:count','2')); cpr.append(el('m:mcJc','left')); col.append(cpr); cols.append(col); mpr.append(cols); matrix.append(mpr)
            for row in raw.split(r'\\'):
                if not row.strip(): continue
                tr=el('m:mr')
                for part in row.split('&'): tr.append(box('m:e',Tex(part).seq()))
                matrix.append(tr)
            d=el('m:d'); pr=el('m:dPr'); pr.append(el('m:begChr','{')); pr.append(el('m:endChr','')); d.append(pr); d.append(box('m:e',[matrix])); return [d]
        if cmd=='displaystyle': return []
        if cmd in ('quad','qquad',',',';',' '): return [mr(' ' if cmd!='qquad' else ' ',True)]
        if cmd in ('{','}','_','%'): return [mr(cmd,True)]
        if cmd in ('min','max','arg','log'): return [mr(cmd,True)]
        if cmd in SYMBOLS: return [mr(SYMBOLS[cmd])]
        raise ValueError('Unsupported LaTeX command '+cmd)
math_count=0
def math(p,text):
    global math_count
    root=el('m:oMath')
    try:
        for n in Tex(text).seq(): root.append(n)
    except Exception as exc: raise ValueError(f'{text!r}: {exc}') from exc
    p._p.append(root); math_count+=1
def font_run(r,size=None,bold=None,cn=CN,latin=LAT):
    r.font.name=latin
    if size: r.font.size=Pt(size)
    if bold is not None: r.bold=bold
    r.font.color.rgb=RGBColor(0,0,0)
    r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),cn); return r
def inline(p,text,size=None):
    text=re.sub(r'<!--.*?-->','',text)
    for part in re.split(r'(\$[^$]+\$|\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))',text):
        if not part: continue
        if part.startswith('$') and part.endswith('$'): math(p,part[1:-1]); continue
        if part.startswith('**') and part.endswith('**'):
            first=len(p.runs); inline(p,part[2:-2],size)
            for r in p.runs[first:]: r.bold=True
            continue
        hm=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if hm:
            h=OxmlElement('w:hyperlink'); rid=p.part.relate_to(hm.group(2),'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True); h.set(qn('r:id'),rid)
            r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=hm.group(1); r.append(t); h.append(r); p._p.append(h); continue
        font_run(p.add_run(part),size)
doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297); sec.top_margin=Mm(25); sec.bottom_margin=Mm(25); sec.left_margin=Mm(28); sec.right_margin=Mm(25); sec.header_distance=Mm(13); sec.footer_distance=Mm(13)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name]; s.font.name=LAT; s.font.color.rgb=RGBColor(0,0,0); s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),CN)
    rf=s._element.get_or_add_rPr().rFonts
    for key in list(rf.attrib):
        if key.endswith('Theme'): del rf.attrib[key]
    s.font.italic=False
    if name.startswith('Heading') or name=='Title': rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
    pp=s._element.get_or_add_pPr()
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pp.append(snap)
for style in doc.styles:
    for border in list(style._element.findall('.//'+qn('w:pBdr'))): border.getparent().remove(border)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))): sec._sectPr.remove(grid)
norm=doc.styles['Normal']; norm.font.size=Pt(12); norm.paragraph_format.line_spacing=1.5; norm.paragraph_format.space_after=Pt(4); norm.paragraph_format.first_line_indent=Pt(24)
norm.paragraph_format.line_spacing=Pt(21)
for name,size,before,after in [('Heading 1',16,0,14),('Heading 2',14,12,7),('Heading 3',12,9,5)]:
    s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True; s.paragraph_format.first_line_indent=Pt(0); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.line_spacing=1.3; s.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True; doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Title'].font.size=Pt(24); doc.styles['Title'].font.bold=True; doc.styles['Title'].paragraph_format.first_line_indent=Pt(0); doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
settings=doc.settings.element; mathpr=el('m:mathPr'); mathpr.append(el('m:mathFont',MATH)); settings.append(mathpr)
lang=OxmlElement('w:themeFontLang'); lang.set(qn('w:val'),'en-US'); lang.set(qn('w:eastAsia'),'zh-CN'); settings.append(lang)
lines=(ROOT/'report_0913.md').read_text().splitlines()
p=doc.add_paragraph('武汉大学',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
for r in p.runs: font_run(r,22,True)
p=doc.add_paragraph('学术型硕士学位论文',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in p.runs: font_run(r,18)
p=doc.add_paragraph('基于大语言模型的\n端到端论辩挖掘方法研究',style='Title'); p.paragraph_format.space_before=Pt(45); p.paragraph_format.space_after=Pt(30)
p=doc.add_paragraph('End-to-End Argument Mining\nwith Large Language Models'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Pt(0)
for r in p.runs: font_run(r,15)
for txt in ['计算机科学与技术','0913修订工作稿','2026年9月13日']:
    p=doc.add_paragraph(txt); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.space_before=Pt(15)
toc_map_path=WORK/'toc_pages.json'
toc_pages=json.loads(toc_map_path.read_text()) if toc_map_path.exists() else {}
def toc():
    heads=[line[3:] for line in lines if line.startswith('## ') and ('章 ' in line or '附录' in line or line=='## 参考文献')]
    doc.add_heading('目录',1)
    for name in heads:
        p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.space_after=Pt(14); p.paragraph_format.line_spacing=1.4
        p.paragraph_format.tab_stops.add_tab_stop(Mm(150),WD_TAB_ALIGNMENT.RIGHT); inline(p,name)
        if name in toc_pages: font_run(p.add_run('\t'+str(toc_pages[name])),12)
def table(rows):
    n=len(rows[0]); widths={5:[54,25,24,24,30],4:[64,31,31,31],3:[49,50,58],2:[100,57]}[n]
    compact=rows[0][0]=='编号'
    if rows[0][0]=='编号': widths=[15,53,89]
    if rows[0][0]=='研究问题': widths=[33,40,38,46]
    if rows[0][0]=='语料': widths=[18,52,30,57]
    t=doc.add_table(rows=0,cols=n); t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Mm(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; trpr=t.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,(c,text,w) in enumerate(zip(cells,row,widths)):
            c.width=Mm(w); tcpr=c._tc.get_or_add_tcPr(); v=OxmlElement('w:vAlign'); v.set(qn('w:val'),'center'); tcpr.append(v); margin=OxmlElement('w:tcMar')
            for side,val in [('top',70),('bottom',70),('left',80),('right',80)]:
                z=OxmlElement('w:'+side); z.set(qn('w:w'),str(val)); z.set(qn('w:type'),'dxa'); margin.append(z)
            tcpr.append(margin)
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E7E7'); tcpr.append(sh)
            p=c.paragraphs[0]; p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(0)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(text)>28 else WD_ALIGN_PARAGRAPH.CENTER
            if compact: p.paragraph_format.line_spacing=Pt(15)
            inline(p,text,10 if compact else 10.5)
            for r in p.runs:
                if ri==0: r.bold=True
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
i=next(i for i,l in enumerate(lines) if l=='## 版本说明'); chapter=0; eq_counter=0; in_appendix=False
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if line.startswith('## 第一章'):
        toc(); sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
        pg=OxmlElement('w:pgNumType'); pg.set(qn('w:start'),'1'); sec._sectPr.append(pg)
        sec.header.is_linked_to_previous=False; sec.footer.is_linked_to_previous=False
        hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER; hp.paragraph_format.first_line_indent=Pt(0); font_run(hp.add_run('武汉大学硕士学位论文　0913修订工作稿'),9)
        fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fp.paragraph_format.first_line_indent=Pt(0); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); fp._p.append(f)
    hm=re.match(r'^(#{2,4}) (.*)',line)
    if hm:
        title=hm.group(2)
        if title.startswith('附录'): in_appendix=True
        if '章 ' in title: chapter+=1; eq_counter=0
        p=doc.add_heading(title,len(hm.group(1))-1)
        if title.startswith('第一章'): p.paragraph_format.page_break_before=False
        continue
    if line.startswith(FENCE):
        kind=line[3:]; code=[]
        while i<len(lines) and not lines[i].startswith(FENCE): code.append(lines[i]); i+=1
        i+=1; p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.left_indent=Mm(3); p.paragraph_format.line_spacing=1.1; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(9); p.paragraph_format.keep_together=True
        if kind=='flow': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        font_run(p.add_run('\n'.join(code)),11 if kind=='flow' else 9.5,latin=LAT if kind=='flow' else 'Courier New'); continue
    if line.startswith('$$'):
        tex=line[2:]
        while not tex.endswith('$$'): tex+='\n'+lines[i]; i+=1
        tex=tex[:-2]; p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.line_spacing=1.2; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(9); p.paragraph_format.keep_together=True; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        math(p,tex); eq_counter+=1; font_run(p.add_run(f'　（{chapter}-{eq_counter}）'),10); continue
    if line.startswith('|'):
        raw=[line]
        while i<len(lines) and lines[i].strip().startswith('|'): raw.append(lines[i].strip()); i+=1
        rows=[[c.strip() for c in r.strip('|').split('|')] for r in raw if not re.match(r'^\|[\s:|-]+\|$',r)]
        assert len({len(r) for r in rows})==1,raw
        table(rows); continue
    p=doc.add_paragraph()
    if in_appendix: p.paragraph_format.line_spacing=Pt(19)
    if line.startswith(('第一章介绍研究问题','本章在明确标签映射')): p.paragraph_format.keep_together=True
    if re.match(r'^(表|图|算法)[A-C0-9]+-\d+\s',line):
        p.paragraph_format.keep_with_next=True; p.paragraph_format.first_line_indent=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; inline(p,line,11); continue
    if line.startswith('待核验'): p.paragraph_format.first_line_indent=Pt(0); inline(p,line,11)
    elif re.match(r'^\[\d+\] ',line):
        p.paragraph_format.first_line_indent=Mm(-8); p.paragraph_format.left_indent=Mm(8); p.paragraph_format.line_spacing=Pt(14); p.paragraph_format.space_after=Pt(8); p.paragraph_format.keep_together=True; inline(p,line,10.5)
    else:
        if re.match(r'^\d+\.',line) or line.startswith('关键词') or line.startswith('Keywords'): p.paragraph_format.first_line_indent=Pt(0)
        inline(p,line,11 if in_appendix else None)
doc.core_properties.title='基于大语言模型的端到端论辩挖掘方法研究'; doc.core_properties.subject='0913修订工作稿'; doc.core_properties.author=''
doc.save(ROOT/'report_0913.docx')
(WORK/'docx_build.json').write_text(json.dumps({'native_equation_count':math_count,'table_count':len(doc.tables),'paragraph_count':len(doc.paragraphs),'cjk_font':CN,'math_font':MATH},ensure_ascii=False,indent=2))
print('Built report_0913.docx',math_count,'native equations',len(doc.tables),'tables')
