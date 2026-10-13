from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm
from docx.oxml.ns import qn
d=Document()
s=d.sections[0]
s.page_width=Mm(210); s.page_height=Mm(297)
for font in ['Songti SC','STSong','Arial Unicode MS','Hiragino Sans GB']:
    p=d.add_paragraph()
    r=p.add_run(font+'  武汉大学计算机科学与技术硕士学位论文 论辩挖掘 结构合法性')
    r.font.name=font; r.font.size=Pt(14)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
d.save(Path(__file__).parent/'font_smoke.docx')
