from pathlib import Path
from copy import deepcopy
from hashlib import sha256
import json, re, zipfile
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parent
SOURCE = BASE.parents[1] / 'manuscript_review_test.docx'
EXPECTED = '6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'
assert sha256(SOURCE.read_bytes()).hexdigest() == EXPECTED
BLOCKS = json.loads((BASE / 'source_ch01_blocks.json').read_text())
ORIGINAL = {x['body_index']: x for x in BLOCKS}

# All edits are confined to the extracted first chapter. References retain their
# source numbering; no experimental records or numerical results are generated.
EDITS = {
23: ['论辩挖掘将观点、理由及其连接方式转化为结构化对象[1]。严格抽取要求组件的原文跨度和类型正确，关系还须满足端点、方向及类别一致。组件边界一旦变化，所有引用该组件的关系身份都会随之改变。因此，组件与关系之间的依赖既影响监督能否沿用，也影响预测图能否完整修正。'],
24: ['以AAEC训练文档essay001为例，Premise T4的原文为“What we acquired from team work is not only how to achieve the same goal with others but more importantly, how to get along with others”，标注为支持Claim T3：“through cooperation, children can learn about interpersonal skills which are significant in the future life of all students”。同文档中的T5、T6也分别支持T3。若T3的边界或类型发生变化，这三条关系都需要重新判定，不能仅凭显示编号沿用。'],
27: ['训练时，局部改写若改变T3或其理由的含义，原有关系监督就可能失效；推理时，若T4被抽取为较短片段，T4→T3的严格对象身份也随之变化。前者要求判断改写后哪些监督仍可使用，后者要求同时处理组件修改与关联关系。本文围绕这两类依赖分别设计学习与修正方法，并在受控比较中评价其作用。'],
28: ['方法的价值由对应比较决定。第三章区分数据构造与学习目标的作用，第四章区分修改范围、共同决策与求解方式的作用，第五章再检验两阶段收益是否重叠。比较同时关注恢复、误删及资源成本，使结论对应到具体配置和适用条件。'],
44: ['论辩挖掘中的直接近邻是质量导向合成QOS。该方法将上下文和组件组织为槽位化输入，在改写后恢复跨度并继承组件类型及关系标注；其方法节说明，论辩关系并未显式提供给生成模型，而依赖语义保持使原关系继续成立[35]。这一做法为局部改写提供了基础，也提出了一个需要继续检验的问题：哪些标签确实在改写后保持，哪些只能视为未知。'],
45: ['QOS与同文提出的多样性导向合成DOS分别提供质量和多样性的参照[35]。本文的QOS式适配继承槽位化改写思想，不等同于原论文两种方案或完整系统的复现。2026年10月公开的对抗强化学习合成研究交替优化生成器与判别器，生成新的文本—结构标注对[49]。因此，结构质量与多样性的联合考虑已有直接近邻。本文研究的是保留原图监督的局部改写，其数据构造与学习目标分别在训练规模和下游抽取器受控的条件下比较。'],
47: ['增强还可能改变模型对表面线索的依赖。话语标记补全的效果因语料而异[40]，弱监督论辩挖掘也需要处理组件边界和关系噪声[18]。因此，本文同时评价标签漂移、表达多样性、接受覆盖和训练成本，并在同一批增强样本上比较不同学习目标。关系保持作为生成与审核目标，其可靠程度由因素审核和下游任务评价共同判断。'],
54: ['本文的直接基础是ICIC 2026提出的多候选重排序与有界修正方法MCR-BR[34]。原方法已包含候选跨度替换、类型调整和关系修改。本文进一步关注组件修改后的完整依赖处理：哪些关系必须重新评价，组件与关系的修改应在什么条件下共同接受。集合动作与求解路径作为实现选择分别检验。'],
60: ['表1-1从输入条件和方法机制比较最接近的研究。已有工作为结构保持增强、多候选选择及对象修正提供了基础。本文分别研究改写前后监督对应的维护，以及组件变化后关联关系的完整重建，并为两者设置相应对照。'],
63: [
'相对槽位化增强，本文增加的审核单位是可对应的组件与关系因素。关系审核覆盖改写集合内部、跨集合边界的合法端点组合，并包含无关系状态NONE；标签变化、原标签分歧或未审核的因素不参与相应的增强侧辅助监督与对齐，原文仍承担完整图生成监督。固定增强视图、有效因素掩码和辅助头后，成对项只改变对应因素的学习约束，不改变样本内容。表3-4分别比较可靠对应构造与QOS式适配、完整成对与NoPair，以区分数据构造和学习目标的作用。',
'相对已有的组件关系联合预测与MCR-BR，本文细化的是生成后修改的事务规约。组件身份改变时，失效旧关系的删除与新关系的重建共同计入一次事务；在固定候选和对象评分下，事务须满足依赖闭合、分项接受条件及预算限制，并按规定字段保留修改记录。这样，组件与关系的变化能够共同评价、共同接受并供回放核对。表4-4以I1对RelationOnly考察组件修改范围，以Single对Sequential考察共同接受方式；集合动作及局部、全局求解的作用另由对应比较判断。',
'外部对照进一步覆盖结构解析、文本生成、结构感知生成、动作生成与大模型系统，选择依据及转换规则见附录G。系统层面的主要复现候选为Kawarada等的文本生成方法与UniASA；GMAM、DENIM作为结构生成补充，QOS／DOS作为数据构造近邻。各方法按输入、表示及代码入口的可核对程度选择，并统一接入AAEC stance、CDCP完整源集合及预测坐标的共同评价规则。其他语言和任务范围见附录A.3。'
],
68: ['组件与关系之间的依赖是两项研究的共同出发点。训练阶段改变文本表达，维护经过检查的监督对应；推理阶段保持原文，修改预测图并完整处理关联关系。图1-2概括二者的问题与方法对应，测试预测不作为训练金标。'],
73: ['RQ1：局部改写后，怎样识别可用于学习的组件与关系对应？在相同来源、变体规模和更新预算下，可靠对应构造、因素级辅助监督与成对学习分别产生什么作用，其收益能否抵偿筛选、拒绝和审核成本？'],
74: ['RQ2：给定候选空间与冻结评分，怎样完整处理组件修改的关系依赖，并控制误改与计算成本？单对象与集合动作、局部与全局求解分别适用于什么规模和预算，哪些差异来自动作范围，哪些来自共同决策？候选空间之外的遗漏单独归因。'],
75: ['两阶段四格比较检验相对ICIC继承配置的条件增量，对象成员分析进一步区分重叠恢复、独有恢复和新增损失。它们用于解释两项研究的关系，不另列为第三项算法创新。'],
79: ['（1）可靠对应因素上的关系感知成对学习。围绕局部改写建立关系影响域和标注输运规则，以逐因素审核状态确定哪些组件与关系可以参与学习。原文提供完整图生成监督，变体只在有效因素上提供辅助监督与分布对齐。默认因素掩码与完整影响域整例筛选分别评价；数据构造与学习目标分开比较，其中成对目标采用同数据消融。'],
80: ['（2）候选与评分约束下的依赖闭合修正。将组件变化与全部受影响关系的重建组织为统一事务，明确对象身份、缺失分数处理、分项接受条件和编辑预算。在这一规约下比较纯关系修改、分步决策、单对象共同决策和集合动作，并检查全局与局部求解的适用条件。依赖闭合规定需要共同处理的对象，实现是否值得保留还取决于误改和资源成本。'],
81: ['候选接口、训练辅助头和推理评分器用于实现上述研究。训练辅助头与推理评分器具有不同的监督对象和使用阶段，第二章统一说明其接口。'],
82: [
'本文将贡献与证据分为三个层次。研究设计层面，提出可靠对应因素的学习规约与依赖闭合的事务规约，并给出可检验的控制比较；系统效果层面，在规定的数据、候选和预算条件下报告整套训练、整套推理及两阶段系统的条件增量；机制归因层面，分别检验成对项、共同接受、集合动作和求解路径的独立作用。',
'第三、四章报告的比较支持配置层面的条件改善，但尚未确立成对项及单对象共同决策的稳定独立优势；集合局部动作未显示额外收益，主测试的I1又全部进入全局路径，因此这些结果不足以支持集合动作或局部求解在主测试中的独立优势。两项规约构成本文的方法内容，各项机制的效果与必要性则按表1-2所列比较判断。'
],
85: ['主要关系F1比较族共14项，定义见2.6.3，其余分析列为探索性。表1-2规定每项主张的证据来源，不预设所有机制、数据集或次级指标都应得到显著正结果；结论随对应比较的支持范围表述。'],
89: ['本文以ICIC 2026发表的MCR-BR方法[34]为基础。表1-3明确已有机制与本研究深化内容，避免将继承能力计为新增贡献。'],
92: ['MCR-BR已有组件类型和候选跨度修改能力。本文的深化集中在可靠因素上的学习，以及组件修改后的依赖完整性和共同决策规约。LoRA、提示、辅助头和散度函数作为继承基础或实现手段，不单独计为原创贡献；具体归属见附录C.1。'],
97: ['第二章统一任务表示、数据、配置与统计协议，为各项比较建立共同基础；第三章定义可靠对应与成对学习，并报告训练比较和因素审核；第四章定义依赖闭合修正，分析评分、候选覆盖、决策比较、资源成本与案例；第五章通过四格和对象成员分析说明两阶段的作用关系；第六章总结配置收益、机制证据及适用条件。附录提供数据坐标、实现验收、实验登记、外部对照、逐种子记录与案例定位。']
}

TABLE84 = [
['需要回答的问题','承担证明的比较','不能由此单独推出','报告位置'],
['整套新训练是否有用','T1/I0对T0/I0','成对项是全部收益来源','表3-3、表5-3'],
['成对学习是否有额外作用','同视图、同掩码、同辅助头下完整成对对NoPair','增强构造本身更优','表3-4'],
['可靠对应构造是否优于近邻构造','同规模、同辅助头、无成对项下可靠对应对QOS式适配','成对学习的作用','表3-4'],
['整个新推理配置是否有用','T1/I1对T1/I0','收益来自某一事务机制','表5-3'],
['修改范围与共同接受各有何作用','同候选、同评分下I1对RelationOnly；Single对Sequential','允许组件修改的收益即共同接受的收益','表4-3、表4-4'],
['修正是否优于等预算重生成','T1/I1对完整图重生成与重排序','独立重采样已实现等token调用','表4-4、表4-9'],
['开发风险限制是否值得保留','T1/I1对取消开发风险限制','开发保留率保证测试安全','表4-4、表4-6'],
['集合动作或局部求解是否值得保留','单对象、集合局部、同目标全局；分别核对可行域和预算','差异全部来自搜索效率','表4-4、表4-9'],
['审核与评分是否可靠','按使用单位分层审核；分开报告评分质量、逐次接受与终态净变化','组件审核代表关系审核；策略指标代表概率质量；净变化代表过程风险','表3-6、表4-7、表4-8a、表4-8b']
]

TABLE91 = [
['研究层次','ICIC已公开机制','本文深化内容','结果归属'],
['训练与表示','结论专项适配、线索提示、句子辅助监督','可靠对应构造、因素掩码与成对学习','原表支持MCR-BR训练配置'],
['候选生成与选择','多候选组件、关系构建、整图重排序','保留初始图接口，建立固定候选空间','候选质量与决策效果分别分析'],
['组件及关系修正','类型调整、候选跨度替换、删边与局部重预测','完整依赖重建与事务接受，比较单对象及集合动作','新增机制分别采用控制比较'],
['评分与评价','覆盖、语义、合法性混合评分及非下降接受','完整对象正确性、计数效用与风险分析','新效用及原生超边评价单独报告']
]
TABLES = {84: TABLE84, 91: TABLE91}
NOTES = {
23: '强化组件身份变化与关系依赖的因果衔接，保留原生例证。',
44: 'Q-02：文献段保留近邻机制，将本文最小差异集中到表1-1后。',
63: 'Q-02：新增两段最小差异，明确审核因素、掩码、固定视图的成对目标，以及依赖闭合、接受条件和事务记录；绑定表3-4与表4-4。',
79: '压缩方法细节，保留两项研究内容，接口说明转接第二章。',
82: 'Q-01：将贡献分为研究规约、配置条件效果、独立机制增量，并如实保留尚未确立的机制优势。',
84: 'Q-01：区分I1/RelationOnly的动作范围比较与Single/Sequential的共同接受比较，避免混淆归因。',
91: '保留继承归属，避免把集合动作预设为已证实增量。'
}

source_doc = Document(SOURCE)
doc = Document()
for attr in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']:
    setattr(doc.sections[0],attr,getattr(source_doc.sections[0],attr))
usable_inches=(doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)/914400
kept={}
changed_paragraphs = {}
from docx.text.paragraph import Paragraph
from docx.table import Table
for block in BLOCKS:
    idx=block['body_index']
    if block['type']=='p':
        texts=EDITS.get(idx,[block.get('text','')])
        if idx in [25,69]:
            p=doc.add_paragraph()
            p.add_run().add_picture(str(BASE/('rId22.png' if idx==25 else 'rId31.png')),width=Inches(usable_inches))
            kept[idx]=p._p
            continue
        for text in texts:
            if not text:continue
            style='Title' if idx==20 else 'Heading 1' if idx in [22,31,67,96] else 'Heading 2' if idx in [34,42,50,58,72,78,88] else 'Caption' if re.match(r'^[图表]1-\d+ ',text) else 'Body Text'
            p=doc.add_paragraph(text,style=style)
            kept.setdefault(idx,p._p)
            if idx in EDITS:changed_paragraphs.setdefault(idx,[]).append(p)
    elif block['type']=='tbl':
        rows=TABLES.get(idx,block['table'])
        t=doc.add_table(rows=len(rows),cols=len(rows[0]))
        kept[idx]=t._tbl
        for row,values in zip(t.rows,rows):
            for cell,value in zip(row.cells,values):cell.text=value

# Retain the source page size and hierarchy. The review excerpt uses legible
# Chinese fonts and black headings, without changing the full manuscript.
for style in doc.styles:
    if style.type == 1:
        style.font.name = 'Arial Unicode MS'
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
        for border in style._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Arial Unicode MS'
        run.font.color.rgb = RGBColor(0,0,0)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    if p.style.style_id in ['FirstParagraph','BodyText','Normal'] and p.text:
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs: r.font.size = Pt(11)
    if p.style.style_id.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(7)
        for r in p.runs: r.font.size = Pt(16 if p.text == '第一章 绪论' else 13 if re.match(r'^1\.[1-4] ',p.text) else 11.5)
    if p.style.style_id=='Title':
        p.paragraph_format.space_after=Pt(12)
        for r in p.runs:r.font.size=Pt(16);r.bold=True
    if re.match(r'^[图表]1-\d+ ', p.text):
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = p.text.startswith('表')
        for r in p.runs: r.font.size = Pt(10)
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
for idx in [63,85,92]:
    Paragraph(kept[idx],doc._body).paragraph_format.space_before=Pt(6)

for ti,t in enumerate(doc.tables):
    t.autofit = False
    widths = ([1.10,1.52,1.65,1.91] if ti==0 else [1.14,2.10,1.70,1.24] if ti==1 else [1.02,1.74,1.77,1.65])
    widths=[w*usable_inches/sum(widths) for w in widths]
    for j,w in enumerate(widths):
        t.columns[j].width = Inches(w)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    for old in list(pr.findall(qn('w:tblBorders'))): pr.remove(old)
    pr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for old in list(trpr.findall(qn('w:trHeight'))): trpr.remove(old)
        for old in list(trpr.findall(qn('w:tblHeader')))+list(trpr.findall(qn('w:cantSplit'))):trpr.remove(old)
        nosplit=OxmlElement('w:cantSplit'); trpr.append(nosplit)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge,amount in [('top',80),('bottom',80),('left',85),('right',85)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(amount));e.set(qn('w:type'),'dxa');margins.append(e)
            for old in list(tcpr.findall(qn('w:tcMar'))):tcpr.remove(old)
            tcpr.insert_element_before(margins, 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark')
            for old in list(tcpr.findall(qn('w:shd'))):tcpr.remove(old)
            if ri==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7E6E6');tcpr.insert_element_before(fill, 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark')
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.1
                p.paragraph_format.keep_with_next=False
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name='Arial Unicode MS';r.font.size=Pt(9.5);r.font.color.rgb=RGBColor(0,0,0);r.bold=ri==0
                    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')

# The source disclosure accompanies this standalone excerpt as a header,
# outside the thesis body. It is not an additional experimental claim.
for section in doc.sections:
    h=section.header.paragraphs[0]
    h.text='test稿件  第一章审阅版  沿用原稿模拟说明  未新增实验数值'
    h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in h.runs:r.font.size=Pt(8);r.font.name='Arial Unicode MS';r.font.color.rgb=RGBColor(0,0,0)
    f=section.footer.paragraphs[0]
    f.text='第一章审阅版  ·  '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in f.runs:r.font.size=Pt(8)

for idx,msg in NOTES.items():
    if idx in changed_paragraphs:
        doc.add_comment(changed_paragraphs[idx][0].runs, msg, author='修订说明', initials='REV')
    elif idx in TABLES:
        p=Table(kept[idx],doc._body).rows[0].cells[0].paragraphs[0]
        doc.add_comment(p.runs,msg,author='修订说明',initials='REV')
OUTPUT = BASE/'第一章_修订稿_v1.docx'
doc.save(OUTPUT)

def render_md():
    out=[]
    for b in BLOCKS:
        idx=b['body_index']
        if b['type']=='p':
            texts=EDITS.get(idx,[b.get('text','')])
            for text in texts:
                if not text:
                    if idx in [25,69]:
                        file='rId22.png' if idx==25 else 'rId31.png'
                        out.append(f'![图1-{1 if idx==25 else 2}]({BASE/file})')
                    continue
                if idx==20: text='# '+text
                elif idx in [22,31,67,96]:text='## '+text
                elif idx in [34,42,50,58,72,78,88]:text='### '+text
                out.append(text)
        elif b['type']=='tbl':
            rows=TABLES.get(idx,b['table'])
            out.append('\n'.join(['|'+'|'.join(rows[0])+'|','|'+'|'.join(['---']*len(rows[0]))+'|']+['|'+'|'.join(r)+'|' for r in rows[1:]]))
    return '\n\n'.join(out)+'\n'
after=render_md()
(BASE/'第一章_修订后全文.md').write_text(after)
patches=[]
for idx,texts in EDITS.items():
    patches.append({'body_index':idx,'before':ORIGINAL[idx]['text'],'after':texts})
for idx,rows in TABLES.items():
    patches.append({'body_index':idx,'before':ORIGINAL[idx]['table'],'after':rows})
patches.sort(key=lambda x:x['body_index'])
(BASE/'第一章_逐项补丁.json').write_text(json.dumps({'source_sha256':EXPECTED,'scope':'chapter 1 only','new_experimental_values':0,'changes':patches},ensure_ascii=False,indent=2))
diff=['# 第一章逐段修改对照','本文件逐项列出实际文本修改。未列出的正文、两幅例图和引文编号沿用原稿。']
for item in patches:
    idx=item['body_index']
    diff += [f'## 原稿正文块 {idx}', '修改前：']
    if isinstance(item['before'],str):diff.append(item['before'])
    else:diff.append('\n'.join(' | '.join(r) for r in item['before']))
    diff.append('修改后：')
    if idx in TABLES:diff.append('\n'.join(' | '.join(r) for r in item['after']))
    else:diff.extend(item['after'])
(BASE/'第一章_逐段修改对照.md').write_text('\n\n'.join(diff)+'\n')

original_md=(BASE/'第一章_修改前原文.md').read_text()
citations=lambda s:set(re.findall(r'\[(\d+)\]',s))
assert citations(original_md)==citations(after), (citations(original_md)-citations(after))
assert sha256(SOURCE.read_bytes()).hexdigest()==EXPECTED
with zipfile.ZipFile(OUTPUT) as z:
    assert z.testzip() is None
    for name in ['word/document.xml','word/comments.xml']:
        etree.fromstring(z.read(name))
docx_text='\n'.join(p.text for p in doc.paragraphs)
assert '第二章 任务表示' not in docx_text
for texts in EDITS.values():
    for t in texts:assert t in docx_text
expected_ptexts=[text for block in BLOCKS if block['type']=='p' for text in EDITS.get(block['body_index'],[block.get('text','')]) if text]
assert [p.text for p in doc.paragraphs if p.text] == expected_ptexts
expected_tables=[TABLES.get(block['body_index'],block['table']) for block in BLOCKS if block['type']=='tbl']
assert [[[cell.text for cell in row.cells] for row in t.rows] for t in doc.tables]==expected_tables
assert len(doc.inline_shapes)==2 and len(doc.tables)==3
qa={'source_unchanged':True,'source_sha256':EXPECTED,'output_sha256':sha256(OUTPUT.read_bytes()).hexdigest(),'edited_source_paragraphs':len(EDITS),'edited_tables':len(TABLES),'figures_preserved':len(doc.inline_shapes),'tables':len(doc.tables),'source_citation_ids':sorted(citations(original_md),key=int),'citation_ids_preserved':True,'new_experimental_values':0,'scope':'chapter 1 excerpt only','zip_xml_valid':True,'docx_markdown_text_match':True,'visual_review':'pending'}
(BASE/'qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps(qa,ensure_ascii=False,indent=2))
