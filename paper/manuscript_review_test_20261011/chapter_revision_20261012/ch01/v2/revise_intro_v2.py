from pathlib import Path
from copy import deepcopy
import hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).resolve().parent
BASE = OUT.parent
SOURCE = BASE.parents[1] / 'manuscript_review_test.docx'
V1 = BASE / '第一章_修订稿_v1.docx'
V1_MD = BASE / '第一章_修订后全文.md'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
source_sha, v1_sha, v1_md_sha = sha(SOURCE), sha(V1), sha(V1_MD)
assert source_sha == '6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'

PARAS = [
    '论辩是人们表达观点、说明理由和回应不同意见的重要方式，广泛存在于议论文写作与公共议题讨论等语言活动中。理解这类文本，需要联系作者的主张与相应理由，并辨别不同论述之间的支持或反驳关系。面对内容分散、表达多样的文本，人工梳理论证过程需要逐一分析语句及其上下文，耗费较多时间。因此，自动识别文本中的论辩结构，是自然语言理解研究中具有实际价值的问题[1]。',
    '论辩挖掘（Argument Mining，AM）旨在从自然语言文本中识别论辩组件及其相互关系，将非结构化的论述转化为可分析的结构[1]。其中，论辩组件是承担观点、理由等功能的文本片段，组件之间的关系则描述这些片段如何共同构成论证。通过确定组件在原文中的位置、判断其论辩角色并识别关系，可以得到文本对应的论辩图。在写作辅助场景中，这种结构可为分析理由是否围绕观点展开提供依据；在公共意见分析中，它有助于整理不同立场及其理由，为后续比较和人工研判提供信息。',
    '以论辩标注作文语料AAEC中的essay001为例，作者提出合作有助于儿童学习人际交往技能，并从团队合作中的相处经验、沟通与妥协以及生活技能等方面说明理由。图1-1展示了其中一组支持关系：T3表示被支持的观点，T4、T5和T6分别表示支撑该观点的理由。论辩挖掘需要从原文中识别这些片段，区分观点与理由的角色，并恢复理由指向观点的连接。该例说明，文本中的论辩结构由组件及其关系共同构成，二者需要结合上下文进行判断。',
    '围绕上述任务，已有研究从组件识别与关系分类出发，逐步探索对论辩结构进行联合建模的方法[2][4]。随着预训练语言模型的发展，生成式方法进一步将组件边界、组件类型和关系组织为统一输出，使从原始文本直接生成论辩结构成为一种可行的技术路线[7]。这一变化为利用上下文信息和统一处理多个子任务提供了条件，但输出形式的统一并不意味着结构判断已经可靠。训练数据能否提供有效监督，以及预测结果中的局部错误如何得到恰当修正，仍是需要研究的问题。',
    '在训练数据方面，论辩结构标注需要同时确定文本片段、论辩角色及其连接，数据构建成本较高[1][2]。利用语言模型改写已有文本，可以增加表达形式，为扩充训练样本提供途径[15]；然而，文本表达发生变化后，原有标注是否仍然成立需要检查。两段文字即使主题相同，也可能在立场、论述范围或理由强度上存在差异，从而影响组件的角色以及对其他组件的支持关系。已有槽位化增强方法通过保持语义来沿用标注[35]，但对于组件与关系相互依赖的任务，还需要进一步判断哪些监督信息能够可靠保留。',
    '在预测结果方面，组件识别与关系判断之间的依赖增加了错误处理的难度。组件的边界或类型错误可能同时影响多条关系；相应地，修改一个组件后，与其相连的关系也需要重新判断。例如，图1-1中的T3与三个理由相连，若只调整T3而直接保留原有关系，修改后的图仍可能与文本含义不一致。另一方面，过多的重新生成或反复修订会增加计算开销，也可能破坏已经正确的部分。因此，如何在控制修改范围和计算成本的同时，协调处理组件与关联关系，是提升论辩结构预测可靠性需要面对的问题。',
    '对上述问题开展研究，具有方法与应用两方面的意义。在方法层面，论辩挖掘同时涉及文本语义、组件角色和关系结构。考察文本改写后监督信息的保留条件，有助于明确数据增强在结构化任务中的使用边界；研究组件变化对关联关系的影响，则有助于分析局部修正与完整结构之间的联系。这些问题能够为论辩挖掘的训练与预测方法提供更具体的设计依据，也可为其他包含对象依赖的结构化抽取任务提供研究思路。',
    '在应用层面，更准确的组件识别与关系恢复，有望为议论文分析、写作反馈和公共意见整理提供更可靠的结构信息。例如，写作反馈可以结合观点及其相连理由展开，意见整理也可以保留不同主张背后的理由，便于使用者检查和比较。研究数据利用效率与修正成本，还有助于评估相关方法在标注资源和计算资源受限条件下的适用性。本文据此从训练阶段的监督保持和推理阶段的结构修正两个方面展开研究，具体研究内容在1.3节介绍。',
]
CAPTION = '图1-1 AAEC作文中的观点与理由及其支持关系（essay001局部结构，中文为阅读提示）'
OLD_12 = '相关研究从组件识别逐步扩展到联合结构预测，并形成生成建模、数据增强和预测后修正等路线。本文按这些机制组织综述，重点关注输入是否包含金标组件、关系依赖如何进入模型，以及修改后如何判定结构质量。AAEC和CDCP的任务表示及评价规则统一在第二章说明。'
NEW_12 = '围绕论辩结构的自动识别与恢复，相关研究涉及端到端建模、训练数据构造和预测结果修正等方面。本节首先介绍端到端论辩挖掘与生成建模的发展，随后梳理关系保持数据增强和结构化预测修正方法，最后结合已有工作的特点与不足，说明本文的研究位置。'
OLD_13 = '组件与关系之间的依赖是两项研究的共同出发点。训练阶段改变文本表达，维护经过检查的监督对应；推理阶段保持原文，修改预测图并完整处理关联关系。图1-2概括二者的问题与方法对应，测试预测不作为训练金标。'
NEW_13 = '针对训练数据扩充中的监督保持问题，以及预测图修正中的组件关系依赖问题，本文分别从训练和推理两个阶段开展研究。训练阶段关注文本改写后可用于学习的可靠对应，推理阶段关注组件修改与关联关系重建的协调处理。两项研究的关系如图1-2所示，以下进一步说明研究问题、主要内容及其与已有工作的联系。'

doc = Document(V1)
ps = doc.paragraphs
start = next(i for i,p in enumerate(ps) if p.text.startswith('1.1 '))
end = next(i for i,p in enumerate(ps) if p.text == '1.2 国内外研究现状')
heading = ps[start]
heading.text = '1.1 研究背景与意义'
for r in heading.runs:
    r.font.size = Pt(13)
body_template = deepcopy(ps[start+1]._p)
picture = deepcopy(next(p._p for p in ps[start+1:end] if p._p.xpath('.//w:drawing')))
caption_template = deepcopy(next(p._p for p in ps[start+1:end] if p.text.startswith('图1-1 ')))
for p in ps[start+1:end]:
    p._p.getparent().remove(p._p)
anchor = heading._p
added=[]
def insert_text(text, template):
    global anchor
    node=deepcopy(template)
    for child in list(node):
        if child.tag != qn('w:pPr'): node.remove(child)
    p=Paragraph(node,doc._body)
    r=p.add_run(text)
    r.font.size=Pt(10 if text.startswith('图1-1 ') else 11)
    p.paragraph_format.first_line_indent=Pt(0 if text.startswith('图1-1 ') else 22)
    anchor.addnext(node); anchor=node
    added.append(p)
    return p
for i,text in enumerate(PARAS):
    insert_text(text,body_template)
    if i==2:
        anchor.addnext(picture);anchor=picture
        insert_text(CAPTION,caption_template)
for p in doc.paragraphs:
    if p.text in [OLD_12,OLD_13]:
        p.text={OLD_12:NEW_12,OLD_13:NEW_13}[p.text]
        for r in p.runs:r.font.size=Pt(11)

# Remove obsolete first-section comments while retaining all other v1 notes.
live_ids={n.get(qn('w:id')) for n in doc._element.xpath('.//w:commentReference')}
for c in list(doc.comments._comments_part.element):
    if c.get(qn('w:id')) not in live_ids:c.getparent().remove(c)
doc.add_comment(added[0].runs,'按参照论文的绪论顺序，先交代研究背景与需求，再引入任务、困难和研究意义。',author='修订说明',initials='REV')
doc.add_comment(added[2].runs,'示例用于解释论辩组件和支持关系；原稿中的标签输运、对象身份和比较协议不在背景部分展开。',author='修订说明',initials='REV')
doc.add_comment(added[-2].runs,'新增方法层面的研究意义；应用意义按潜在用途表述，不作为已完成的应用验证。',author='修订说明',initials='REV')
OUTPUT = OUT/'第一章_修订稿_v2.docx'
doc.save(OUTPUT)

intro=['## 1.1 研究背景与意义']
for i,t in enumerate(PARAS):
    intro.append(t)
    if i==2:intro += [f'![图1-1]({BASE/"rId22.png"})',CAPTION]
intro_md='\n\n'.join(intro)+'\n'
oldmd=V1_MD.read_text()
prefix, tail=oldmd.split('## 1.1 ',1)
oldintro,rest=tail.split('## 1.2 国内外研究现状',1)
oldintro='## 1.1 '+oldintro.strip()+'\n'
newmd=prefix+intro_md+'\n## 1.2 国内外研究现状'+rest
newmd=newmd.replace(OLD_12,NEW_12).replace(OLD_13,NEW_13)
(OUT/'第一章_修订后全文_v2.md').write_text(newmd)
(OUT/'1.1_修订后全文_v2.md').write_text(intro_md)
(OUT/'1.1_逐段修改对照_v2.md').write_text('# 第一章第二轮修改对照\n\n## 修改前 v1\n\n'+oldintro+'\n## 修改后 v2\n\n'+intro_md+'\n## 衔接段修改\n\n### 1.2节开头\n\n修改前：'+OLD_12+'\n\n修改后：'+NEW_12+'\n\n### 1.3节开头\n\n修改前：'+OLD_13+'\n\n修改后：'+NEW_13+'\n')

# Verify version preservation, exact edited text and every unaffected block.
assert sha(SOURCE)==source_sha and sha(V1)==v1_sha and sha(V1_MD)==v1_md_sha
v1doc=Document(V1)
def suffix_nodes(d):
    nodes=list(d._element.body)
    pos=next(i for i,n in enumerate(nodes) if ''.join(n.itertext()).startswith('1.2 国内外研究现状'))
    return nodes[pos:]
a,b=suffix_nodes(v1doc),suffix_nodes(doc)
assert len(a)==len(b)
for x,y in zip(a,b):
    old=''.join(x.xpath('.//w:t/text()'))
    if old in [OLD_12,OLD_13]:continue
    assert etree.tostring(x)==etree.tostring(y),old[:60]
assert len(doc.tables)==3 and len(doc.inline_shapes)==2
for t in PARAS+[CAPTION,NEW_12,NEW_13]:assert t in [p.text for p in doc.paragraphs]
assert set(re.findall(r'\[(\d+)\]',newmd))==set(re.findall(r'\[(\d+)\]',oldmd))
with zipfile.ZipFile(OUTPUT) as z:
    assert z.testzip() is None
    for name in ['word/document.xml','word/comments.xml']:etree.fromstring(z.read(name))
qa={'source_sha256':source_sha,'v1_sha256':v1_sha,'output_sha256':sha(OUTPUT),'source_and_v1_unchanged':True,'scope':'1.1 rewritten; 1.2 and 1.3 introductory paragraphs revised; all remaining chapter 1 body blocks unchanged','new_experimental_values':0,'intro_paragraphs':len(PARAS),'intro_chinese_characters':sum(len(re.findall(r'[\u4e00-\u9fff]',t)) for t in PARAS),'figures':2,'tables':3,'citation_ids_preserved':True,'unaffected_ooxml_blocks_identical':True,'zip_xml_valid':True,'visual_review':'pending'}
(OUT/'qa_v2.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps(qa,ensure_ascii=False,indent=2))
