from pathlib import Path
from copy import deepcopy
import hashlib,json,re,zipfile
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT=Path(__file__).resolve().parent
V2=OUT.parent/'v2'
CH2=OUT.parent.parent/'ch02'
SOURCE=OUT.parents[2]/'manuscript_review_test.docx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
baseline=json.loads((V2/'delivery_manifest_v2.json').read_text())
preserved={SOURCE:sha(SOURCE),**{V2/n:h for n,h in baseline['files'].items()},CH2/'第二章_补充材料_v1.md':sha(CH2/'第二章_补充材料_v1.md')}
for path,digest in preserved.items():assert sha(path)==digest
assert preserved[SOURCE]=='6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'
oldmd=(V2/'第一章_修订后全文_v2.md').read_text()
oldconclusion=next(p for p in oldmd.split('\n\n') if p.startswith('第三、四章报告的比较'))
newconclusion='第三、四章报告的比较支持配置层面的条件改善，但尚未确立成对项及单对象共同决策的稳定独立优势；集合局部动作未显示额外收益，主测试的I1又全部进入全局路径，因此这些结果不足以支持集合动作或局部求解在主测试中的独立优势。两项规约构成本文的方法内容，各项机制的效果与必要性由相应控制比较判断。'
oldresponsibility=next(p for p in oldmd.split('\n\n') if p.startswith('主要关系F1比较族共14项'))
newresponsibility='为区分配置效果与机制作用，本文分别比较整套训练、成对目标、整套推理、组件修改范围及共同接受，并以等预算重生成和开发风险限制检验替代方案与使用条件。主要关系F1比较族包含7组比较在两种语料上的14项检验，具体定义见2.6.3节，控制条件与归因边界见补充表S2-4；其他分析列为探索性。各项比较不预设显著正结果，结论按对应证据的支持范围表述。'
oldinheritance=next(p for p in oldmd.split('\n\n') if p.startswith('本文以ICIC 2026发表'))
newinheritance='本文以ICIC 2026发表的MCR-BR方法[34]为基础。表1-2概括已有机制与本文研究内容。'
oldcaption='表1-3 MCR-BR已公开机制与本文深化内容'
newcaption='表1-2 MCR-BR已有机制与本文研究内容'
oldstructure=next(p for p in oldmd.split('\n\n') if p.startswith('第二章统一任务表示')).strip()
newstructure=oldstructure.replace('第二章统一任务表示、数据、配置与统计协议，为各项比较建立共同基础','第二章介绍论辩挖掘的任务表示、相关理论与技术，并统一评价和公共实验协议，为各项比较建立共同基础')
replacements={oldconclusion:newconclusion,oldresponsibility:newresponsibility,oldinheritance:newinheritance,oldcaption:newcaption,oldstructure:newstructure}
rows=[
 ['研究层次','MCR-BR已有机制','本文研究内容'],
 ['训练与表示','结论专项适配、线索提示、句子辅助监督','可靠对应构造、因素掩码与成对学习'],
 ['候选生成与选择','多候选组件、关系构建与整图重排序','保留初始图接口，明确固定候选及完整对象身份'],
 ['组件与关系修正','类型调整、候选跨度替换、删边与局部重预测','依赖完整重建与事务接受，比较单对象及集合动作'],
 ['评分与评价','覆盖、语义、合法性混合评分及非下降接受','完整对象正确性评分、计数效用与风险分析'],
]

doc=Document(V2/'第一章_修订稿_v2.docx')
original_nodes=list(doc._element.body)
unaffected={}
changed=set()
affected_paragraphs=[]
for p in list(doc.paragraphs):
 if p.text in replacements:
  old=p.text;changed.add(p._p);p.text=replacements[old]
  for run in p.runs:run.font.size=Pt(10 if old==oldcaption else 11)
  affected_paragraphs.append(p)
assert len(affected_paragraphs)==len(replacements)
caption=next(p for p in doc.paragraphs if p.text=='表1-2 研究问题与承担证明责任的比较')
responsibility_table=doc.tables[1]
inheritance_table=doc.tables[2]
removed={caption._p,responsibility_table._tbl,inheritance_table._tbl}
for n in original_nodes:
 if n not in changed and n not in removed:unaffected[n]=etree.tostring(n)
caption._p.getparent().remove(caption._p)
responsibility_table._tbl.getparent().remove(responsibility_table._tbl)

# Replace the inherited four-column table with a readable three-column comparison.
t=doc.add_table(rows=len(rows),cols=3);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
width=(doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)/914400
widths=[width*.18,width*.39,width*.43]
for i,w in enumerate(widths):t.columns[i].width=Inches(w)
pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 b=OxmlElement('w:'+side)
 for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+key),val)
 borders.append(b)
pr.append(borders)
for i,row in enumerate(t.rows):
 rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
 if i==0:rp.append(OxmlElement('w:tblHeader'))
 for j,c in enumerate(row.cells):
  c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cp=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
  for side,v in [('top',75),('bottom',75),('left',90),('right',90)]:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');m.append(e)
  cp.append(m)
  if i==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');cp.append(sh)
  p=c.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
  r=p.add_run(rows[i][j]);r.font.name='Arial Unicode MS';r.font.size=Pt(9.5);r.bold=(i==0)
  r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
inheritance_table._tbl.addprevious(t._tbl)
inheritance_table._tbl.getparent().remove(inheritance_table._tbl)

# Keep only comments whose anchors still exist, then annotate the two confirmed edits.
live_ids={n.get(qn('w:id')) for n in doc._element.xpath('.//w:commentReference')}
for c in list(doc.comments._comments_part.element):
 if c.get(qn('w:id')) not in live_ids:c.getparent().remove(c)
summary=next(p for p in doc.paragraphs if p.text==newresponsibility)
doc.add_comment(summary.runs,'按已确认方案，将原表1-2的详细比较责任后移至第二章补充表S2-4；绪论保留研究问题、比较族与归因边界的概括。',author='修订说明',initials='REV')
table_intro=next(p for p in doc.paragraphs if p.text==newinheritance)
doc.add_comment(table_intro.runs,'原表1-3压缩为三列四项，顺序改为表1-2。继承与新增内容继续明确区分；删去的结果归属内容由实验章节及附录C.1承担。',author='修订说明',initials='REV')
output=OUT/'第一章_修订稿_v3.docx';doc.save(output)

# Apply the identical scoped changes to the complete Markdown chapter.
newmd=oldmd
removed_md=newmd.split('表1-2 研究问题与承担证明责任的比较\n\n',1)[1].split('\n\n'+oldresponsibility,1)[0]
newmd=newmd.replace('表1-2 研究问题与承担证明责任的比较\n\n'+removed_md+'\n\n','',1)
oldtable=newmd.split(oldcaption+'\n\n',1)[1].split('\n\n',1)[0]
newtable='\n'.join(['|'+'|'.join(rows[0])+'|','|---|---|---|']+['|'+'|'.join(row)+'|' for row in rows[1:]])
newmd=newmd.replace(oldtable,newtable,1)
for old,new in replacements.items():assert old in newmd;newmd=newmd.replace(old,new,1)
(OUT/'第一章_修订后全文_v3.md').write_text(newmd)
patches=[dict(kind='paragraph',before=a,after=b) for a,b in replacements.items()]
patches += [dict(kind='table_moved',before=removed_md,after='第二章补充表S2-4；绪论保留概括段落'),dict(kind='table_replaced',before=oldtable,after=newtable)]
(OUT/'第一章_v2至v3_逐项修改.json').write_text(json.dumps(patches,ensure_ascii=False,indent=2))
comparison=['# 第一章v2至v3修改对照','本轮采用已确认的表格调整方案，未重新改写1.1和1.2。']
for i,p in enumerate(patches,1):comparison += [f'## 修改{i}', '修改前：\n\n'+p['before'],'修改后：\n\n'+p['after']]
(OUT/'第一章_v2至v3_修改对照.md').write_text('\n\n'.join(comparison)+'\n')

for n,data in unaffected.items():assert etree.tostring(n)==data
for path,digest in preserved.items():assert sha(path)==digest
assert len(doc.tables)==2 and len(doc.inline_shapes)==2
for expected in replacements.values():assert expected in [p.text for p in doc.paragraphs]
assert [[c.text for c in row.cells] for row in doc.tables[1].rows]==rows
assert '表1-3' not in newmd
assert newmd.count('表1-2')==2
assert newmd.count('补充表S2-4')==1
assert set(re.findall(r'\[(\d+)\]',newmd))==set(re.findall(r'\[(\d+)\]',oldmd))
assert newmd.split('### 1.3.2',1)[0]==oldmd.split('### 1.3.2',1)[0]
with zipfile.ZipFile(output) as z:
 assert z.testzip() is None
 for name in ['word/document.xml','word/comments.xml']:etree.fromstring(z.read(name))
qa={'source_sha256':preserved[SOURCE],'v2_sha256':baseline['docx_sha256'],'v3_sha256':sha(output),'prior_versions_unchanged':True,'chapter2_supplement_unchanged':True,'unaffected_ooxml_blocks_identical':len(unaffected),'sections_1_1_and_1_2_unchanged':True,'edited_paragraphs':len(replacements),'tables':2,'figures':2,'citation_ids_preserved':True,'old_table_1_3_references_remaining':0,'new_experimental_values':0,'zip_xml_valid':True,'visual_review':'pending'}
(OUT/'qa_v3.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps(qa,ensure_ascii=False,indent=2))
