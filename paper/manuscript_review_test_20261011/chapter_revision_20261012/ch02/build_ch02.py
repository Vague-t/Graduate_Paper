from pathlib import Path
import re,json,hashlib,zipfile
from copy import deepcopy
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

BASE=Path(__file__).resolve().parent
SOURCE=BASE.parents[1]/'manuscript_review_test.docx'
SOURCE_SHA='6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_SHA

def el(tag,*children,**attrs):
    n=OxmlElement('m:'+tag)
    for k,v in attrs.items():n.set(qn('m:'+k),v)
    for child in children:
        if isinstance(child,list):n.extend(child)
        else:n.append(child)
    return n
def mr(text,normal=False):
    run=el('r')
    if normal:run.append(el('rPr',el('nor')))
    t=el('t');t.text=text;t.set(qn('xml:space'),'preserve');run.append(t)
    return run
def box(tag,nodes):return el(tag,nodes)
def scripts(base,sub=None,sup=None):
    if sub is not None and sup is not None:return el('sSubSup',box('e',[base]),box('sub',sub),box('sup',sup))
    if sub is not None:return el('sSub',box('e',[base]),box('sub',sub))
    if sup is not None:return el('sSup',box('e',[base]),box('sup',sup))
    return base

class MathParser:
    SYM={'theta':'θ','alpha':'α','rho':'ρ','phi':'φ','ell':'ℓ','Delta':'Δ','kappa':'κ','delta':'δ','in':'∈','leq':'≤','geq':'≥','neq':'≠','mid':'|','Vert':'‖','times':'×','ldots':'…','cdots':'⋯','log':'log','quad':'  ','qquad':'    ','left':'','right':'','!':'',',':' ',';':'  ','{':'{','}':'}','|':'‖'}
    def __init__(self,s):self.s=s;self.i=0
    def ws(self):
        while self.i<len(self.s) and self.s[self.i].isspace():self.i+=1
    def group(self):
        self.ws()
        if self.i<len(self.s) and self.s[self.i]=='{':
            self.i+=1;v=self.expr('}');assert self.i<len(self.s) and self.s[self.i]=='}';self.i+=1;return v
        return [self.atom()]
    def postfix(self):
        sub=sup=None
        while True:
            self.ws()
            if self.i>=len(self.s) or self.s[self.i] not in '_^':break
            op=self.s[self.i];self.i+=1;v=self.group()
            if op=='_':sub=v
            else:sup=v
        return sub,sup
    def atom(self):
        self.ws()
        if self.s[self.i]=='{':return el('box',box('e',self.group()))
        if self.s[self.i]=='\\':
            self.i+=1
            m=re.match(r'[A-Za-z]+|.',self.s[self.i:]);cmd=m.group();self.i+=len(cmd)
            if cmd=='frac':return el('f',box('num',self.group()),box('den',self.group()))
            if cmd=='sqrt':return el('rad',el('radPr',el('degHide',val='1')),el('deg'),box('e',self.group()))
            if cmd in ['hat','widehat','bar']:
                return el('acc',el('accPr',el('chr',val='̂' if cmd!='bar' else '̄')),box('e',self.group()))
            if cmd in ['mathrm','operatorname','mathsf','mathbf','mathcal','mathbb']:
                nodes=self.group()
                if cmd in ['mathrm','operatorname','mathsf'] and all(n.tag==qn('m:r') for n in nodes):
                    return mr(''.join(n.find(qn('m:t')).text or '' for n in nodes),normal=True)
                if cmd=='mathbb':
                    for t in nodes:
                        for x in t.iter(qn('m:t')):
                            if x.text=='R':x.text='ℝ'
                if cmd=='mathcal':
                    for n in nodes:
                        for run in ([n] if n.tag==qn('m:r') else list(n.iter(qn('m:r')))):
                            run.insert(0,el('rPr',el('scr',val='script')))
                if cmd in ['mathrm','operatorname','mathsf']:
                    for n in nodes:
                        for run in ([n] if n.tag==qn('m:r') else list(n.iter(qn('m:r')))):
                            if run.find(qn('m:rPr')) is None:run.insert(0,el('rPr',el('nor')))
                return el('box',box('e',nodes))
            if cmd not in self.SYM:raise ValueError('Unsupported TeX command '+cmd)
            return mr(self.SYM[cmd],normal=cmd=='log')
        ch=self.s[self.i];self.i+=1
        return mr(ch)
    def expr(self,stop=None):
        out=[]
        while True:
            self.ws()
            if self.i>=len(self.s) or (stop and self.s[self.i]==stop):break
            m=re.match(r'\\(sum|prod)(?![A-Za-z])',self.s[self.i:])
            if m:
                self.i+=len(m.group());sub,sup=self.postfix()
                pr=el('naryPr',el('chr',val='∑' if m.group(1)=='sum' else '∏'),el('limLoc',val='undOvr'))
                if sub is None:pr.append(el('subHide',val='1'))
                if sup is None:pr.append(el('supHide',val='1'))
                out.append(el('nary',pr,box('sub',sub or []),box('sup',sup or []),box('e',self.expr(stop))))
                break
            atom=self.atom();sub,sup=self.postfix();out.append(scripts(atom,sub,sup))
        return out

def math_node(s):
    p=MathParser(s);nodes=p.expr();assert p.i==len(s)
    return el('oMath',nodes)

source_doc=Document(SOURCE)
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.82);sec.bottom_margin=Inches(.78)
sec.left_margin=sec.right_margin=Inches(.82)
sec.header_distance=sec.footer_distance=Inches(.35)
usable=6.63
for style in doc.styles:
    if style.type==1:
        style.font.name='Arial Unicode MS';style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
        for x in style._element.xpath('.//w:pBdr'):x.getparent().remove(x)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.18
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.first_line_indent=Pt(21)
for name,size in [('Title',16),('Heading 1',13),('Heading 2',11.5)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
header.add_run('test稿件  第二章审阅版  沿用原稿模拟说明').font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第二章审阅版 · ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

eq_count=0;md_tables=[];expected_text=[]
def add_inline(p,text):
    # Native OMML for inline math; code identifiers remain plain readable text.
    pieces=re.split(r'(\$[^$]+\$)',text)
    for piece in pieces:
        if piece.startswith('$') and piece.endswith('$'):p._p.append(math_node(piece[1:-1]))
        elif piece:
            piece=piece.replace('`','')
            p.add_run(piece)
def table(rows):
    n=len(rows[0]);t=doc.add_table(rows=len(rows),cols=n);t.autofit=False
    if n==8:widths=[.62,.56,.64,.64,.63,.71,.64,.57]
    elif n==5:widths=[1.55,1.15,1.28,.9,1.2]
    elif n==4:widths=[.65,1.58,1.4,2.2]
    elif n==3:widths=[1.15,2.40,2.7]
    elif n==2:widths=[1.32,5]
    elif n==7:widths=[.6,.63,.63,.63,1.03,1.04,.83]
    else:widths=[1]*n
    widths=[w*usable/sum(widths) for w in widths]
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    for j,w in enumerate(widths):t.columns[j].width=Inches(w)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge,amt in [('top',50),('bottom',50),('left',80),('right',80)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(amt));e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            if i==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E7E6E6');tcpr.append(shd)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or n in [7,8] else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p,rows[i][j])
            for r in p.runs:r.font.size=Pt(9 if n>=7 else 9.5);r.bold=i==0
    md_tables.append(rows)
    return t

def build(md):
    global eq_count
    lines=md.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line=='$$':
            buf=[]
            while i<len(lines) and lines[i].strip()!='$$':buf.append(lines[i].strip());i+=1
            i+=1;tex=' '.join(buf);tag=re.search(r'\\tag\{([^}]+)\}',tex)
            tex=re.sub(r'\\tag\{[^}]+\}','',tex).strip()
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(7)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p._p.append(math_node(tex))
            if tag:p.add_run('   （'+tag.group(1)+'）').font.size=Pt(9)
            eq_count+=1;continue
        if line.startswith('|'):
            rows=[line.strip('|').split('|')]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=lines[i].strip().strip('|').split('|');i+=1
                if all(re.fullmatch(r':?-+:?',c.strip()) for c in row):continue
                rows.append(row)
            table(rows);continue
        img=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if img:
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(img.group(2),width=Inches(usable));continue
        if line.startswith('# '):style='Title';text=line[2:]
        elif line.startswith('## '):style='Heading 1';text=line[3:]
        elif line.startswith('### '):style='Heading 2';text=line[4:]
        else:style='Normal';text=line
        p=doc.add_paragraph(style=style);add_inline(p,text);expected_text.append(text)
        next_line=next((s.strip() for s in lines[i:] if s.strip()),'')
        if next_line=='$$':p.paragraph_format.keep_with_next=True
        if re.match(r'^(?:表2-\d+|补充表S2-\d+|图2-\d+)\s',text):
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=not text.startswith('图')
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5)
            for r in p.runs:r.font.size=Pt(9.5);r.bold=True

build((BASE/'第二章_修订后全文_v1.md').read_text())
doc.add_page_break()
build((BASE/'第二章_补充材料_v1.md').read_text())
output=BASE/'第二章_修订稿_v1.docx';doc.save(output)
with zipfile.ZipFile(output) as z:
    assert z.testzip() is None
    root=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    assert not any('$' in t or '\\frac' in t for t in root.xpath('.//w:t/text()',namespaces=ns))
    math_total=len(root.xpath('.//m:oMath',namespaces=ns))
assert eq_count==10
assert len(doc.tables)==8
assert len(doc.inline_shapes)==1
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_SHA
qa={'source_sha256':SOURCE_SHA,'source_unchanged':True,'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'display_equations':eq_count,'native_math_nodes':math_total,'tables':len(doc.tables),'figures':len(doc.inline_shapes),'zip_xml_valid':True,'raw_latex_absent_from_word_text':True,'visual_review':'pending'}
(BASE/'document_qa.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2));print(json.dumps(qa,ensure_ascii=False,indent=2))
