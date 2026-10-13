from pathlib import Path
from copy import deepcopy
import hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from pypdf import PdfReader

B = Path(__file__).resolve().parent
C1, C2 = B/'ch01', B/'ch02'
O1, O2, REPORT = C1/'v4', C2/'v2', B/'heading_consolidation_20261012'
for p in [O1,O2,REPORT]: p.mkdir(exist_ok=True)
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
source = B.parent/'manuscript_review_test.docx'
inputs = [source,C1/'v3/第一章_修订稿_v3.docx',C1/'v3/第一章_修订后全文_v3.md',C2/'第二章_修订稿_v1.docx',C2/'第二章_修订后全文_v1.md',C2/'第二章_补充材料_v1.md']
preserved = {str(p):sha(p) for p in inputs}
assert preserved[str(source)] == '6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'

# None removes a heading; a body transition uses the Normal paragraph style.
h1 = {
 '1.2.4 综合评述与本文研究位置': None,
 '1.3.1 研究问题': ('Normal','围绕上述两个阶段，本文提出以下研究问题。'),
 '1.3.2 主要研究内容': ('Normal','针对上述研究问题，本文开展以下两项研究工作。'),
 '1.3.3 与MCR-BR的关系': None,
}
h2 = {
 '2.1.1 文档与论辩组件': ('Heading 2','2.1.1 论辩组件与关系'),
 '2.1.2 二元关系与多前提关系': None,
 '2.2 对象身份与结构化表示': ('Heading 2','2.1.2 字符坐标与对象身份'),
 '2.2.1 字符坐标与分词坐标的无损转换': None,
 '2.2.2 组件键、关系键与标注对应': None,
 '2.2.3 对象、候选与预测的层次关系': None,
 '2.3 预训练语言模型与学习方法': ('Heading 1','2.2 预训练语言模型与学习方法'),
 '2.3.1 注意力机制与自回归生成': ('Heading 2','2.2.1 注意力机制与自回归生成'),
 '2.3.2 参数高效适配与连续提示': ('Heading 2','2.2.2 参数高效适配与连续提示'),
 '2.3.3 数据增强与因素一致性学习': ('Heading 2','2.2.3 数据增强与因素一致性学习'),
 '2.4 候选评分与结构化决策': ('Heading 1','2.3 候选评分与结构化决策'),
 '2.4.1 多候选生成与初始结构选择': None,
 '2.4.2 完整对象评分与概率校准': None,
 '2.4.3 结构约束、搜索范围与计算预算': None,
 '2.5 评价指标与计数原则': ('Heading 1','2.4 评价指标与计数原则'),
 '2.5.1 严格对象匹配与micro-F1': None,
 '2.5.2 重复、失败与结果关联': None,
 '2.6 公共实验条件与统计分析': ('Heading 1','2.5 公共实验条件与统计分析'),
 '2.6.1 数据划分与使用原则': None,
 '2.6.2 公共配置与成本口径': None,
 '2.6.3 配对重采样与多重比较': None,
 '2.7 训练与推理的阶段接口': ('Heading 1','2.6 训练与推理的阶段接口'),
 '2.8 本章小结': ('Heading 1','2.7 本章小结'),
}
r1 = {
 '不同语料的任务表示和训练来源会影响方法比较。':'综合上述研究，方法比较还需考虑语料任务表示和训练来源的差异。',
 '具体定义见2.6.3节':'具体定义见2.5节',
}
r2 = {
 '以2.1、2.2节的字符身份进行评价':'以2.1节的字符身份进行评价',
 '对于每个训练种子，先汇总全部测试文档的整数计数并计算micro-F1，再报告三个种子的均值与样本标准差。':'在上述公共实验条件下，统计分析先对每个训练种子汇总全部测试文档的整数计数并计算micro-F1，再报告三个种子的均值与样本标准差。',
 '用途及其规模见正文2.6.1。':'用途及其规模见正文2.5节。',
 '这张表承接原表1-2的证据责任内容，作为从绪论移至实验说明的候选位置；已确认的第一章v2在本轮不被覆盖。':'这张表承接原绪论中研究问题与比较责任的详细说明；第一章保留概括性表述，并在2.5节统一说明主要比较族。',
}

def md_edit(text, heads, replacements):
 out=[]
 for line in text.splitlines():
  m=re.match(r'^#{2,3} (.+)$',line)
  if m and m.group(1) in heads:
   edit=heads[m.group(1)]
   if edit is None: continue
   style,t=edit
   line=('## ' if style=='Heading 1' else '### ' if style=='Heading 2' else '')+t
  out.append(line)
 text='\n'.join(out)+'\n'
 for a,b in replacements.items():text=text.replace(a,b)
 return re.sub(r'\n{3,}','\n\n',text)

def canonical(nodes):
 return [etree.tostring(n,method='c14n',exclusive=True) for n in nodes]

def doc_edit(src, dest, heads, replacements):
 doc=Document(src)
 math_before=canonical(doc._element.xpath('.//m:oMath'))
 tables_before=canonical([t._tbl for t in doc.tables])
 body_sample=next(p for p in doc.paragraphs if p.style.name in ['Normal','Body Text'] and len(p.text)>100)
 normal_pr=deepcopy(body_sample._p.pPr)
 seen=[]
 for p in list(doc.paragraphs):
  if p.text not in heads:continue
  seen.append(p.text);edit=heads[p.text]
  if edit is None:
   p._p.getparent().remove(p._p);continue
  style,text=edit;p.text=text;p.style=style
  if style=='Normal':
   if p._p.pPr is not None:p._p.remove(p._p.pPr)
   if normal_pr is not None:p._p.insert(0,deepcopy(normal_pr))
   p.style=body_sample.style
   for r in p.runs:r.bold=False;r.font.size=Pt(11)
  else:
   for r in p.runs:r.font.size=None;r.bold=None
 assert set(seen)==set(heads)
 replacement_counts={}
 for a,b in replacements.items():
  count=0
  for t in doc._element.xpath('.//w:t'):
   if t.text and a in t.text:count+=t.text.count(a);t.text=t.text.replace(a,b)
  assert count>0,(a,'not found in visible Word text')
  replacement_counts[a]=count
 live={e.get(qn('w:id')) for e in doc._element.xpath('.//w:commentReference')}
 if len(doc.comments):
  for c in list(doc.comments._comments_part.element):
   if c.get(qn('w:id')) not in live:c.getparent().remove(c)
 assert canonical(doc._element.xpath('.//m:oMath'))==math_before
 assert canonical([t._tbl for t in doc.tables])==tables_before
 doc.save(dest)
 with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dest) as zout:
  assert zout.testzip() is None
  for name in zout.namelist():
   if name.endswith('.xml'):etree.fromstring(zout.read(name))
  for name in zin.namelist():
   if name.startswith('word/media/'):assert zin.read(name)==zout.read(name)
 return {'native_equations_preserved':len(math_before),'tables_preserved':len(doc.tables),'figures_preserved':len(doc.inline_shapes),'heading_edits':len(seen),'replacement_counts':replacement_counts,'zip_xml_valid':True}

old1=inputs[2].read_text();old2=inputs[4].read_text();olds=inputs[5].read_text()
new1=md_edit(old1,h1,r1);new2=md_edit(old2,h2,r2);news=md_edit(olds,{},r2)
md1=O1/'第一章_修订后全文_v4.md';md2=O2/'第二章_修订后全文_v2.md';sup=O2/'第二章_补充材料_v2.md'
md1.write_text(new1);md2.write_text(new2);sup.write_text(news)
doc1=O1/'第一章_修订稿_v4.docx';doc2=O2/'第二章_修订稿_v2.docx'
qa1=doc_edit(inputs[1],doc1,h1,r1);qa2=doc_edit(inputs[3],doc2,h2,r2)

def headings(text):return [l for l in text.splitlines() if re.match(r'^#{2,3} [12]\.\d',l)]
def counts(text):return {'second':len(re.findall(r'^## [12]\.\d+ ',text,re.M)),'third':len(re.findall(r'^### [12]\.\d+\.\d+ ',text,re.M))}

assert counts(new1)=={'second':4,'third':3}
assert counts(new2)=={'second':7,'third':5}
for old,new in [(old1,new1),(old2,new2),(olds,news)]:
 assert set(re.findall(r'\[(\d+)\]',old))==set(re.findall(r'\[(\d+)\]',new))
 assert re.findall(r'\$\$.*?\$\$',old,re.S)==re.findall(r'\$\$.*?\$\$',new,re.S)
 assert [l for l in old.splitlines() if l.startswith('|')]==[l for l in new.splitlines() if l.startswith('|')]

# Bind the comparison to supplied PDF versions and fresh body heading extraction.
refs=json.loads((C1/'v2/reference_audit.json').read_text())
ranges={'temporal_kg':(25,31),'contrastive':(21,30),'style_transfer':(25,32),'transformer':(18,27),'prior_0405':(26,33)}
second_counts={'temporal_kg':4,'contrastive':4,'style_transfer':5,'transformer':4,'prior_0405':8}
expected={'temporal_kg':4,'contrastive':7,'style_transfer':0,'transformer':7,'prior_0405':3}
audits=[]
for ref in refs:
 p=Path(ref['path']);assert sha(p)==ref['sha256']
 lo,hi=ranges[ref['key']];reader=PdfReader(p);found=[]
 for n in range(lo,hi+1):
  for line in (reader.pages[n-1].extract_text() or '').splitlines():
   m=re.match(r'^\s*(2\.\d+\.\d+)(?:\s|[^\d.])',line)
   if m:found.append({'physical_page':n,'number':m.group(1),'text':line.strip() if ref['key']!='temporal_kg' else m.group(1)})
 assert len(found)==expected[ref['key']],(ref['key'],found)
 audits.append({'key':ref['key'],'path':str(p),'sha256':ref['sha256'],'chapter2_physical_pages':[lo,hi],'second_level_count':second_counts[ref['key']],'third_level_count':len(found),'third_level_headings':found})
(REPORT/'reference_heading_audit.json').write_text(json.dumps(audits,ensure_ascii=False,indent=2))

mapping={
 '1.2.4':'1.2末尾的综合评述段落','1.3.1':'1.3中的研究问题段落','1.3.2':'1.3中的研究内容条目','1.3.3':'1.3中的继承关系说明与表1-2',
 '2.1':'2.1','2.1.1':'2.1.1','2.1.2':'2.1.1','2.2':'2.1.2','2.2.1':'2.1.2','2.2.2':'2.1.2','2.2.3':'2.1.2',
 '2.3':'2.2','2.3.1':'2.2.1','2.3.2':'2.2.2','2.3.3':'2.2.3',
 '2.4':'2.3','2.4.1':'2.3','2.4.2':'2.3','2.4.3':'2.3',
 '2.5':'2.4','2.5.1':'2.4','2.5.2':'2.4','2.6':'2.5','2.6.1':'2.5','2.6.2':'2.5','2.6.3':'2.5','2.7':'2.6','2.8':'2.7',
}
(REPORT/'章节编号映射.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2))

lines=['# 第一章与第二章标题收拢对照','本轮以第一章v3和第二章v1为基线，分别生成第一章v4和第二章v2。旧稿保留。','## 参照依据','以下为已核对PDF的第二章结构。页码为PDF物理页码；参照论文体现写作习惯，不构成学校强制格式。','|参考论文|第二章所在页|二级标题|三级标题|','|---|---|---:|---:|']
for a in audits:
 lines.append(f"|[{Path(a['path']).stem}]({a['path']})|{a['chapter2_physical_pages'][0]}—{a['chapter2_physical_pages'][1]}|{a['second_level_count']}|{a['third_level_count']}|")
lines += ['三级标题用于区分真正不同的任务类型或技术分支。篇幅较短、前后相依的概念、评价规则和配置说明可以由连续段落组织，不需要逐一编号。',
 '## 本轮结构变化','|章节|二级标题修改前→后|三级标题修改前→后|','|---|---:|---:|','|第一章v3→v4|4→4|7→3|','|第二章v1→v2|8→7|16→5|',
 '第一章保留研究现状的三个方向。综合评述接在1.2末尾，1.3通过过渡句连接研究问题、两项研究工作和MCR-BR继承关系。取消的标题没有改成加粗小标题。',
 '第二章将原2.1和2.2合为任务与结构表示，仅保留“组件与关系”和“字符坐标与对象身份”两个分支。模型与学习方法保留三个技术分支。候选评分、评价计数及公共实验条件各自采用连贯段落，删去短小的三级标题。',
 '## 修改后完整目录','### 第一章','```text',*['  '*(l.count('#')-2)+l.lstrip('# ') for l in headings(new1)],'```','### 第二章','```text',*['  '*(l.count('#')-2)+l.lstrip('# ') for l in headings(new2)],'```',
 '## 编号与引用对应','|旧位置|新位置|','|---|---|']
lines += [f'|{a}|{b}|' for a,b in mapping.items()]
lines += ['本轮已同步第一章中主要比较族的引用、第二章字符身份的引用、补充材料的数据用途引用。图号、表号和公式号保持不变。后续各章并入总稿时，须按上表逐项核对语义后更新章节引用，不能用全局数字替换。',
 '## 逐项修改与衔接','第一章：取消1.2.4；1.3.1、1.3.2分别改为普通过渡句，取消1.3.3；综合评述增加“综合上述研究”的衔接；2.6.3节引用更新为2.5节。',
 '第二章：合并任务定义和对象身份；保留并重编号五个必要分支；删除候选、评价和公共协议下的八个三级标题；统计段落增加承接句；补充材料删除已经过时的“候选迁移位置”说明，改为已落实的引用。',
 '## 缺陷修复情况','|问题|本轮状态|说明|','|---|---|---|','|第一章标题过细|已修复|7个三级标题收拢为3个，1.3不再分三级标题|','|第二章标题过细及任务表示重复分节|已修复|16个三级标题收拢为5个，二级标题8个收拢为7个|','|收拢后的段落连接与章节引用|已修复当前交付两章|保留正文、补充过渡并更新交叉引用|','|Q03字符边界与对象身份|保持上一版修复|原生例子、公式与规范键定义保留|','|Q04统计证据缺口|仍为部分修复|统计方法与复算接口保留，未新增主实验逐文档计数或完成主结果复算|',
 '## 内容保留范围','本轮处理标题层级、相邻过渡及章节引用。两章全部表格、插图、公式和文献编号保留；实验数值未修改。第二章补充材料继续保留test稿件的模拟数值说明。第三章及后续正文不在本轮修改范围内。',
 '## 交付与检查','第一章和第二章分别提供完整正文Markdown及Word，另提供本对照说明。Word完成渲染后逐页检查中文、表格、公式、图注和分页。']
report='\n\n'.join(lines)+'\n'
# Join table and fenced-outline rows without unintended blank lines.
report=re.sub(r'(?m)(\|[^\n]*\|)\n\n(?=\|)',r'\1\n',report)
report=re.sub(r'```text\n\n(.*?)\n\n```',lambda m:'```text\n'+m.group(1).replace('\n\n','\n')+'\n```',report,flags=re.S)
(REPORT/'第一章与第二章_标题收拢对照与修复说明.md').write_text(report)

for out,chapter,text,old,qa in [(O1,'第一章',new1,old1,qa1),(O2,'第二章',new2,old2,qa2)]:
 version='v4' if chapter=='第一章' else 'v2'
 full=f'# {chapter}{version}完整正文与变更说明\n\n'+text+'\n\n'+report
 if chapter=='第二章':full+='\n\n'+news
 (out/f'{chapter}_{version}_完整正文与变更说明.md').write_text(full)
 qa.update({'before_heading_counts':counts(old),'after_heading_counts':counts(text),'citation_ids_preserved':True,'experimental_values_changed':False,'source_versions_unchanged':True,'visual_review':'pending'})
 (out/f'qa_{version}.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))

for path,digest in preserved.items():assert sha(Path(path))==digest
manifest={'input_sha256':preserved,'outputs':{str(p):sha(p) for p in [md1,doc1,md2,doc2,sup]},'chapter1_counts':counts(new1),'chapter2_counts':counts(new2),'visual_review':'pending'}
(REPORT/'delivery_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps({'chapter1':qa1,'chapter2':qa2},ensure_ascii=False,indent=2))
