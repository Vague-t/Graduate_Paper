from pathlib import Path
from collections import Counter
from copy import deepcopy
import json,re,hashlib,zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from normalize_ch03_ch04_math import normalize_math_layout

B=Path(__file__).resolve().parent;O=B/'ch03_ch04_revision_v1'
S=B.parent/'manuscript_review_test.docx';src=Document(S);els=list(src._element.body)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(S)=='6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'
text=lambda e:''.join(x.text or '' for x in e.iter() if x.tag in [qn('w:t'),qn('m:t')])
canon=lambda e:etree.tostring(e,method='c14n',exclusive=True)
spec=json.loads((O/'修订块清单.json').read_text());qa={}
tablemap=json.loads((O/'表号映射.json').read_text())
tablepat=re.compile(r'表('+ '|'.join(re.escape(k) for k in sorted(tablemap,key=len,reverse=True))+r')(?![\da-zA-Z])')
def tabreplace(t):return tablepat.sub(lambda m:'表'+tablemap[m.group(1)],t)
def tbrows(e):return [[text(c) for c in r.findall(qn('w:tc'))] for r in e.findall(qn('w:tr'))]
for ch,name,lo,hi in [(3,'第三章',208,393),(4,'第四章',393,682)]:
 d=Document(O/f'{name}_修订稿_v1.docx');md=(O/f'{name}_修订后全文_v1.md').read_text()
 audit=json.loads((O/f'ch{ch}_draft_audit.json').read_text())
 inds=audit['source_indices'];assert len(inds)==len(set(inds))
 oldtables=[tbrows(e) for e in els[lo:hi] if e.tag==qn('w:tbl')]
 newtables=[tbrows(t._tbl) for t in d.tables]
 for t in oldtables:assert t in newtables,('source table modified',ch,t[0])
 def normalize_source_math(m):
  return normalize_math_layout(deepcopy(m))
 mathsrc=Counter(canon(normalize_source_math(m)) for i in inds for m in els[i].xpath('.//m:oMath'))
 normalized_differences=sum(any('\\' in (t.text or '') for t in m.iter(qn('m:t'))) for i in inds for m in els[i].xpath('.//m:oMath'))
 mathout=Counter(canon(m) for m in d._element.xpath('.//m:oMath'))
 assert not (mathsrc-mathout),('native equation changed',ch)
 # Every Word paragraph with no native math or drawing must be recoverable from Markdown.
 norm=lambda s:re.sub(r'\s+','',s)
 failures=[]
 for p in d.paragraphs:
  if p._p.xpath('.//m:oMath') or p._p.xpath('.//w:drawing') or not p.text.strip():continue
  if norm(p.text) not in norm(md):failures.append(p.text)
 assert not failures,failures
 for t in d.tables:
  for row in t.rows:
   for cell in row.cells:
    if not cell._tc.xpath('.//m:oMath'):assert norm(cell.text) in norm(md),(ch,cell.text)
 headings=[p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
 expected=[s.lstrip('# ') for s in md.splitlines() if re.match(r'^#{2,3} [34]\.',s)]
 assert headings==expected
 assert 'Qwen2.4' not in md and '3.4.3节' not in md
 if ch==3:assert 'Qwen2.5-7B-Instruct' in md
 oldcites=set(re.findall(r'\[(\d+)\]',''.join(text(e) for e in els[lo:hi])))
 newcites=set(re.findall(r'\[(\d+)\]',md));assert oldcites==newcites,(ch,oldcites-newcites,newcites-oldcites)
 captions=[p.text for p in d.paragraphs if re.match(r'^表[34]-\d+\s',p.text)]
 numbers=[int(re.match(r'^表[34]-(\d+)',s).group(1)) for s in captions]
 assert numbers==list(range(1,len(d.tables)+1)),(ch,numbers)
 pdf=O/f'_render_final_ch{ch}'/f'{name}_修订稿_v1.pdf';reader=PdfReader(pdf)
 pages=[p.extract_text() for p in reader.pages]
 assert all(len(t.strip())>30 for t in pages)
 assert not any('\ufffd' in t for t in pages)
 assert not any('¿' in t for t in pages)
 oldhead=[]
 for e in els[lo:hi]:
  s=text(e)
  if re.match(rf'^{ch}\.\d+(?:\.\d+)* ',s):oldhead.append(s)
 counter=lambda xs:{'second':sum(bool(re.match(r'^\d\.\d+ ',s)) for s in xs),
                    'third':sum(bool(re.match(r'^\d\.\d+\.\d+ ',s)) for s in xs),
                    'fourth':sum(bool(re.match(r'^\d\.\d+\.\d+\.\d+ ',s)) for s in xs)}
 qa[str(ch)]={'old_heading_counts':counter(oldhead),'new_heading_counts':counter(headings),
  'retained_source_tables':len(oldtables),'new_teaching_tables':len(newtables)-len(oldtables),
  'source_table_cells_exact':True,'native_equations_preserved_after_format_normalization':sum(mathsrc.values()),
  'setminus_rendering_normalized_equations':normalized_differences,
  'new_native_equations':sum(mathout.values())-sum(mathsrc.values()),
  'citations_preserved':sorted(oldcites,key=int),'word_markdown_prose_match':True,
  'table_numbers_sequential':True,'pdf_pages':len(reader.pages),'pdf_replacement_glyphs':0,
  'docx_sha256':sha(O/f'{name}_修订稿_v1.docx'),'headings':headings}

# Reference paths are exact retained versions; excerpts stay in local review evidence.
refs=json.loads((B/'ch01/v2/reference_audit.json').read_text())
toc_ranges={'temporal_kg':(9,10),'contrastive':(11,12),'style_transfer':(10,12),'transformer':(7,8),'prior_0405':(10,11)}
descriptions={
 'temporal_kg':'以引言和问题定义引出总体模型，方法模块后接实验、消融、敏感性或案例分析；用于第三章的构造与学习分层、第四章的问题到框架过渡。',
 'contrastive':'两个方法章均保留引言、若干核心模块、实验设计与结果分析、小结；用于本稿七个二级标题的组织，避免把每项接口规则单列成节。',
 'style_transfer':'围绕两个主要方法模块展开，随后集中报告实验对比；用于第三章“增强构造—成对学习”主线及第四章“接受准则—依赖事务”的连接。',
 'transformer':'第三、四章以研究动机与方法概述开篇，按关键模块展开并集中分析复杂度，实验另设第五章；仅借鉴其问题到方法的叙述次序，不机械套用章内实验布局。',
 'prior_0405':'先给形式建模、模型或生成框架，再说明训练策略与实验；其第四章实验部分较简略，只作方法层次参考，不作为实验完整性标准。'}
refaudit=[]
for r in refs:
 p=Path(r['path']);assert sha(p)==r['sha256']
 cache=B.parents[1]/'tmp/pdfs/comparison_1009'/f"{r['key']}_text.txt"
 chunks=re.split(r'===== PHYSICAL PAGE (\d+) =====',cache.read_text());pages={int(chunks[i]):chunks[i+1] for i in range(1,len(chunks),2)}
 lo,hi=toc_ranges[r['key']];toc='\n'.join(pages[n] for n in range(lo,hi+1))
 rec=dict(r);rec['toc_physical_pages']=[lo,hi];rec['application']=descriptions[r['key']]
 rec['chapter3_4_toc_excerpt']=[s.strip() for s in toc.splitlines() if re.match(r'^\s*[34]\.\d',s)]
 refaudit.append(rec)
(O/'参考论文结构核对.json').write_text(json.dumps(refaudit,ensure_ascii=False,indent=2))
(O/'qa_final.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))

report=['# 第三章与第四章修订说明和缺陷状态',
 '本轮分别形成第三章v1、第四章v1，依据原test稿件及其逐章质量分析修订。原稿、已确认的第一章v4和第二章v2均保留；第五章及后续章节未修改。',
 f'修订基线：`{S}`；SHA-256：`{sha(S)}`。',
 '## 参考论文结构及本轮采用方式',
 '五份参考论文均核对到具体PDF版本。它们提供组织方式参考，不构成学校强制格式，也不要求各章标题数量相同。',
 '|参考论文|核对目录页（PDF物理页）|采用方式|','|---|---|---|']
for r in refaudit:report.append(f"|[{Path(r['path']).stem}]({r['path']})|{r['toc_physical_pages'][0]}—{r['toc_physical_pages'][1]}|{r['application']}|")
report+=['第三章采用“问题—框架—增强构造—学习目标—算法与性质—实验—小结”，将审查规则组织到构造模块，将损失设计组织到学习模块。第四章采用“问题—候选—评分与接受—依赖闭合—求解与性质—实验—小结”，把原4.4移到原4.3之前，先让读者知道接受依据，再理解事务为什么要闭合。',
 '## 标题与内容调整','|章节|二级标题|三级标题|四级标题|','|---|---|---|---|']
for ch in [3,4]:
 a=qa[str(ch)];old=a['old_heading_counts'];new=a['new_heading_counts']
 report.append(f"|第{ch}章|{old['second']} → {new['second']}|{old['third']} → {new['third']}|{old['fourth']} → {new['fourth']}|")
report += [
 '第三章增加问题背景和整体接口说明；统一φ的字符身份；合并审核与CDCP语义未知规则；合并辅助头与对齐定义；将性质放在算法之后连续说明；把分层审核和加权漂移放到同一节，再讨论成本。新增一个接纳及一个拒绝的教学增强过程，说明全部坐标、掩码和权重。',
 '第四章合并五个候选小节为两个逻辑单元；合并评分与风险选择，取消4.4.2.1；把冲突闭合、完整关系重建和事务示例组织到新4.4；局部与全局求解先后展开，再集中讨论性质与路由；实际CPU成功、失败和原生语料案例放在同一节。',
 '正文没有将取消的编号换成同样密集的加粗小标题。局部和全局数学约束、端点依赖边界、累计与净编辑差异、缺分和超时语义等关键内容继续保留。',
 '## 缺陷修复情况','|问题|当前状态|具体处理与剩余边界|','|---|---|---|',
 '|Q03 字符身份在第三四章的延续|已修正文形式|φ改为文档版本与字符半开区间的映射；候选跨度u展开为(d,b,e)，缓存和关系键引用同一身份；未将此编辑称为实际模型验收|',
 '|Q05 缺少真实增强全链条|示例与接口已补 真实证据仍缺|新增构造的部分掩码接纳与否定拒绝案例，坐标和归一化可复算；教学记录不是模型生成记录|',
 '|Q06 分层审核上界解释|已修复风险措辞 真实分层区间待补|Wilson改为工作性近似；明确AAEC跨边界层上界5.74%超过5%；π=1只描述固定池；不据总体源组数伪造逐层聚类区间|',
 '|Q07 成对项独立收益|已修复解释|将固定视图的T1I0−NoPair置于解释中心，保留0.0600／0.0720；平均JS和整图质量分开，不把总配置增益归于成对项|',
 '|Q08 主修正轨迹缺失|教学重放已补 真实证据仍缺|构造一次接受、一次缺分拒绝、一次总效用过门但关系分项下降的拒绝；实际CPU规则继续独立报告，未包装为Qwen或ILP轨迹|',
 '|Q09 共同接受及集合动作独立增益|已修复论证边界|显式比较RelationOnly、Sequential、Single、SetLocal；集合可达性只证明表达能力，未确认的收益继续保留|',
 '|Q10 大规模局部分支|已修复适用范围 真实大池验证待补|全部690个模拟文档—种子输入走Global；2500阈值及附录H扫描标明合成条件，不称真实大池部署验收|',
 '|m01 组件ID不一致|已修复|将c1/c2/c3改为表中T4/T3/T7，并强调两种关系分别审核|',
 '|m02 训练曲线统计编码不清|已从正文移除误导项|未找到原图对应逐种子输入与绘图代码，无法核定均值和色带，故移出正文；原图保留供核查，不擅自补图注定义|',
 '|m03 影响域与审核子集易混|已修复说明|增加R_A⊆J_A⊆D_A⊆Ω_sem及U_sem排除关系，结合原图3-1解释默认与完整筛选的覆盖边界|',
 '|m04 集合反例等于接受阈值|已修复构造并验证|旧效用0.60、新效用0.90，Δ=0.30>0.20；单独删除−0.30，合并必须删除两旧对象；另附B.6替换文本|',
 '|Q04 主结果统计复算输入|仍保留上一章未完成项|结果表沿用模拟摘要，不把教学枚举或文档QA当成主实验复算|',
 '“已修复”用于文本、形式规约和构造例子的缺陷；需要真实目标流程数据的问题，明确保留其未验证状态。本轮新增的数值只属于显式教学构造，原主结果表的数值没有调整。',
 '## 数值与图表处理',
 '第三章保留8张原表，新增2张教学表；第四章保留16张原表，新增2张教学表。原表单元格逐一保持，表号按当前出现顺序统一重排：第三章为表3-1至表3-10，第四章为表4-1至表4-18。映射保存在表号映射.json。',
 '原图3-3移出正文，文件“原图3-3_聚合与色带定义待核对.png”保留。其余6幅插图保持原像素。未为原图虚构逐种子序列、误差带或训练日志。',
 '## 完整目录']
for ch in [3,4]:report += [f'### 第{ch}章','```text',*qa[str(ch)]['headings'],'```']
report += ['## 表号与章节衔接','表格编号与章节移动会影响其他章节的交叉引用；本轮只更新第三、四章内部引用。第一、二章已经确认的版本保持只读，后续合并总稿时应按映射逐项同步第五章、附录、图表目录及前文引用。',
 '|原表号或新增示例|本版表号|','|---|---|']
for a,b in tablemap.items():report.append(f'|表{a}|表{b}|')
report += ['章节精确映射见章节引用映射.json。主要调整为原4.4→新4.3、原4.3→新4.4；原4.5.6→新4.5.2；原4.5.3—4.5.5→新4.5.3；原4.5.7→新4.5.4；原4.6.5—4.6.8→新4.6.4。',
 '## 复算与文档检查',
 'replay_ch03_ch04_examples.py已执行：全部字符切片往返、md与监督权重说明、16种关系选择的枚举、缺分与负分项拒绝、预算差集和严格阈值反例均通过。该执行仅为教学算术和小规模枚举，不包含LLM训练、语义验证器调用或生产ILP运行。',
 f"原24张表的内容均与输入一致；第三章保留{qa['3']['native_equations_preserved_after_format_normalization']}个原生数学对象，新增φ及集合关系的原生公式；第四章保留{qa['4']['native_equations_preserved_after_format_normalization']}个原生数学对象。原公式中的差集反斜杠规范为数学符号∖；局部和全局约束的空矩阵占位改为逐行公式排列，逐一检查数学符号与顺序不变；辅助监督公式移除转换异常的描述性下括号，其组件项和关系项继续由紧随公式的正文解释。这三项处理修复了转换器显示红色问号的问题；其他保留数学对象的OOXML逐一一致。未将公式转成图片。Word与Markdown正文、标题及无公式表格单元格匹配，文献编号集合不变。",
 f"第三章渲染为{qa['3']['pdf_pages']}页，第四章{qa['4']['pdf_pages']}页，已逐页检查中文、公式、表格、插图及分页，并完成异常公式的修复复查。PDF中未检出替代字符或公式转换问号；交付清单记录本次检查所绑定的文件摘要。PDF和页面PNG仅作内部排版检查。",
 '## 本轮交付及停止点',
 '完整修订后正文见本目录的两份Markdown及Word。综合审阅文件将本说明与两章全文连续展示，便于审阅。完成本轮后停在第三、四章，等待用户review，再处理第五章及后续整合。']
reporttext='\n\n'.join(report)+'\n'
reporttext=re.sub(r'(?m)(^\|.*\|)\n\n(?=\|)',r'\1\n',reporttext)
(O/'第三章与第四章_变更点与缺陷修复说明.md').write_text(reporttext)
full=reporttext+'\n---\n\n'+(O/'第三章_修订后全文_v1.md').read_text()+'\n---\n\n'+(O/'第四章_修订后全文_v1.md').read_text()
(O/'第三章与第四章_完整修订稿与变更说明.md').write_text(full)

# Version-bound before/after patch ledger (the source indexes are zero-based OOXML body indexes).
patches=[]
for ch in [3,4]:
 blocks=spec[f'ch{ch}'];kept={x['index'] for x in blocks if x['kind']=='source'}
 lo,hi=(208,393) if ch==3 else (393,682)
 for i in range(lo,hi):
  old=text(els[i])
  if old and i not in kept:patches.append({'chapter':ch,'source_body_index':i,'original':old,'action':'rewritten, merged, or moved out; see revised block sequence'})
(O/'逐项修改与新增内容.json').write_text(json.dumps({'replaced_or_removed_source_blocks':patches,'new_revised_blocks':{f'ch{c}':[x for x in spec[f'ch{c}'] if x['kind']!='source'] for c in [3,4]}},ensure_ascii=False,indent=2))
patchmd=['# 第三章与第四章逐项修改对照','本表按原稿位置列出被重写或取消的段落与标题。完整替换后的段落按新目录连续列在第二部分；公式和表格请以完整修订稿为准。','## 原稿被改写或取消的内容']
for p in patches:patchmd += [f"### 第{p['chapter']}章 原正文块{p['source_body_index']}",p['original']]
patchmd+=['## 新增或重写内容按修订后顺序']
for ch in [3,4]:
 patchmd.append(f'### 第{ch}章')
 for x in spec[f'ch{ch}']:
  if x['kind']=='heading':patchmd.append('#### '+x['text'])
  elif x['kind']=='text':patchmd.append(tabreplace(x['text']))
  elif x['kind']=='table_new':patchmd.append('\n'.join('|'+ '|'.join(tabreplace(c) for c in row)+'|' for row in [x['rows'][0],['---']*len(x['rows'][0]),*x['rows'][1:]]))
(O/'第三章与第四章_逐项修改对照.md').write_text('\n\n'.join(patchmd)+'\n')

b6='''# 附录B6数值反例替换文本

本文件供后续总稿整合使用，本轮不覆盖原附录。

取两个兼容旧组件c1=[0,4)、c2=[4,8)，以及与二者均冲突的新组件b=[0,8)。三者类型相同，当前与候选关系集合均为空。令κC=1、wC=wR=1、δ=0.20，qC(c1)=qC(c2)=0.65、qC(b)=0.95。旧效用为0.30+0.30=0.60，新效用为0.90，故ΔΦC=0.30、ΔΦR=0，总增量严格大于δ。组件预算至少为3时，2→1集合事务可以接受。

单独删除任一旧组件的效用增量为−0.30，不能通过严格改进门。插入b或仅声明替换一个旧组件，都必须经冲突闭合删除另一个旧组件，其实际形态仍是2→1，因而不属于单对象邻域。该构造说明集合动作的可达性，不证明真实F1改善。

原例0.8→1.0的增量为0.20，只能在δ严格小于0.20时通过；默认δ=0.20时必须拒绝。不要将严格大于改成大于等于来迁就原例。
'''
(O/'附录B6_数值反例修订补丁.md').write_text(b6)
print(json.dumps({c:{k:v for k,v in a.items() if k!='headings'} for c,a in qa.items()},ensure_ascii=False,indent=2))
