from pathlib import Path
from copy import deepcopy
from io import BytesIO
import re,json,hashlib,zipfile,math,sys
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from normalize_ch03_ch04_math import normalize_math_layout

B=Path(__file__).resolve().parent
SRC=B.parent/'manuscript_review_test.docx'
OUT=B/'ch03_ch04_revision_v1'; OUT.mkdir(exist_ok=True)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(SRC)=='6de7116c7d5e6b76a06104d9f3a787c98dc7168504e7813146c8a1a2b9d92ab5'
original=Document(SRC); elements=list(original._element.body)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
# Reuse the already validated native equation builder without executing the old authoring body.
code=(B/'ch02/build_ch02.py').read_text().split('source_doc=Document(SOURCE)')[0]
env={'__file__':str(B/'ch02/build_ch02.py')};exec(code,env);math_node=env['math_node']
env['MathParser'].SYM.update({'to':'→','mapsto':'↦','prime':'′'})
plain=lambda e: ''.join(e.xpath('.//w:t/text()|.//m:t/text()',namespaces=ns)) if isinstance(e,etree._Element) and not hasattr(e,'pPr') else ''.join(e.itertext())
def txt(e):return ''.join(t.text or '' for t in e.iter() if t.tag in [qn('w:t'),qn('m:t')])
def code_text(e):
 return ''.join((n.text or '') if n.tag==qn('w:t') else '\n' if n.tag==qn('w:br') else '\t' if n.tag==qn('w:tab') else '' for n in e.iter())
def sty(e):
 s=e.find('w:pPr/w:pStyle',ns)
 return s.get(qn('w:val')) if s is not None else 'BodyText'

# All targets are final section numbers. Original table and figure numbers stay stable.
mapping={
 '2.2':'2.1.2','2.3':'2.2','2.4':'2.3','2.5':'2.4','2.6':'2.5','2.6.1':'2.5','2.6.2':'2.5','2.6.3':'2.5','2.7':'2.6','2.8':'2.7',
 '3.3.4':'3.3.3','3.4.2':'3.4.1','3.4.3':'3.4.2','3.5.1':'3.5','3.5.2':'3.5',
 '3.6.2':'3.6.2','3.6.3':'3.6.2','3.6.4':'3.6.3','3.6.5':'3.6.3','3.6.6':'3.3.4',
 '4.2.1':'4.2.1','4.2.2':'4.2.1','4.2.3':'4.2.2','4.2.4':'4.2.2','4.2.5':'4.2.2',
 '4.3':'4.4','4.3.1':'4.4.1','4.3.2':'4.4.2',
 '4.4':'4.3','4.4.1':'4.3.1','4.4.2':'4.3.1','4.4.2.1':'4.3.1','4.4.3':'4.3.2','4.4.4':'4.3.2','4.4.5':'4.4.3',
 '4.5.1':'4.5.1','4.5.2':'4.5.1','4.5.3':'4.5.3','4.5.4':'4.5.3','4.5.5':'4.5.3','4.5.6':'4.5.2','4.5.7':'4.5.4',
 '4.6.1':'4.6.1','4.6.2':'4.6.1','4.6.3':'4.6.2','4.6.4':'4.6.3','4.6.5':'4.6.4','4.6.6':'4.6.4','4.6.7':'4.6.4','4.6.8':'4.6.4','4.6.9':'4.6.5'}
refpat=re.compile(r'(?<![A-Za-z0-9_.])(?:'+ '|'.join(re.escape(k) for k in sorted(mapping,key=len,reverse=True))+r')(?![\d.])')
def refreplace(t):return refpat.sub(lambda m:mapping[m.group()],t)
tablemap={'3-A':'3-2','3-B':'3-3','3-1':'3-1','3-2':'3-4','3-3':'3-5','3-4':'3-6','3-5':'3-7','3-6':'3-8','3-8':'3-9','3-7':'3-10',
 '4-A':'4-1','4-B':'4-2','4-1':'4-3','4-2':'4-4','4-3':'4-5','4-4':'4-6','4-5':'4-7','4-6':'4-8','4-7':'4-9','4-8a':'4-10','4-8b':'4-11','4-9':'4-12','4-10':'4-13','4-11':'4-14','4-12':'4-15','4-13':'4-16','4-14':'4-17','4-15':'4-18'}
tablepat=re.compile(r'表('+ '|'.join(re.escape(k) for k in sorted(tablemap,key=len,reverse=True))+r')(?![\da-zA-Z])')
def tabreplace(t):return tablepat.sub(lambda m:'表'+tablemap[m.group(1)],t)
captionpat=re.compile(r'^(?:图|表|算法)[34]-[\dAB]+[ab]?\s+')

heads={
 214:'3.2 总体框架与基础模型',
 234:'3.3 关系影响域与可靠增强视图',
 236:'3.3.1 改写范围与关系因素',256:'3.3.2 条件生成与字符坐标迁移',264:'3.3.3 审核掩码与语义未知因素',273:None,
 287:'3.4.1 辅助预测与可靠因素对齐',297:None,311:'3.4.2 源文档归一化与参数更新',325:None,332:None,
 341:'3.6.1 实验设置与受控比较',348:'3.6.2 主结果与机制消融',354:None,
 361:'3.6.3 分层审核与语义漂移',375:None,382:None,
 406:'4.2.1 完整候选生成与初始图选择',413:None,423:'4.2.2 固定对象候选与资源边界',430:None,437:None,
 468:'4.3 对象评分与接受准则',471:'4.3.1 固定评分器与开发风险选择',479:None,485:None,
 490:'4.3.2 计数效用与完整事务接受',500:None,507:None,
 445:'4.4 依赖闭合与事务构造',448:'4.4.1 组件冲突闭合',455:'4.4.2 完整关系重建',
 516:'4.5.1 局部求解与有界执行',529:None,542:'4.5.3 方法性质与适用条件',549:None,555:None,
 563:'4.5.2 全局对象选择',581:'4.5.4 规模路由与返回语义',
 593:'4.6.1 主结果与受控比较',600:None,610:'4.6.2 候选覆盖与修正空间',617:'4.6.3 评分质量与修正代价',
 637:'4.6.4 实际规则诊断与原生关系案例',647:None,652:None,658:None,
 666:'4.6.5 弱评分条件与回退分析'}

repl={
211:[
 '在端到端论辩挖掘中，组件文字既承载自身的论辩角色，也为与其他组件之间的关系提供依据。因此，局部改写即使保持句意相近，也可能改变支持方向、削弱跨边界理由，或使原本没有标注关系的因素获得新的解释。若直接将全部旧标签复制到增强文本，增加的样本可能同时引入错误监督。',
 '本章研究在组件位置变化和关系语义不确定的条件下，如何构造可追踪的增强视图，并将其中通过检查的监督用于训练。方法首先定义改写的关系影响域，在字符坐标上迁移全部组件；随后分别审核组件与关系，对变化、待定及未审核的因素设置掩码；最后在固定视图和掩码上比较无成对、组件成对与完整成对学习。',
 '本章依次介绍总体框架、增强构造、学习目标及算法性质，再通过受控比较讨论整套训练配置、成对目标和审核成本。需要区分两项判断：可靠因素接口明确了哪些标签参与学习；成对项是否带来独立收益，仍须由固定数据条件下的比较回答。',
 '本章沿用test稿件的模拟评审用途。主结果、审核统计与成本数值为原稿的模拟材料；新增的完整增强示例为明确标注的教学构造，均不能代替目标模型的真实运行证据。表3-1单独保留真实语料的原文与标注身份。'],
240:['表内T4、T5、T6各自以supports指向T3，T4→T7没有标注边。以局部观察集合A={T3,T4}解释因素分类时，T4→T3属于内部正因素，T5→T3和T6→T3属于跨边界正因素，T4→T7属于跨边界NONE。该观察集合用于说明定义，不等于默认自动选择结果：若默认中心为T3，所有指向T3的完整源成员T4、T5、T6都会进入改写集合。'],
259:['$\\phi:C\\to C^{\\prime},\\quad(d,b_c,e_c,t_c)\\mapsto(d^{\\prime},b_c^{\\prime},e_c^{\\prime},t_c).$'],
260:['组件采用第二章定义的规范身份：文档版本、原文字符半开区间和原生类型。增强全文是另一个文本版本，故记为d′；显示ID仅建立追踪对应，不能替代坐标或版本身份。回填后更新所有组件的字符区间，再派生评价token与模型子词表示；只读组件也会随前文长度变化而移动。φ是组件之间的双射，不是字符或token位置之间的双射。关系及超边的全部端点经过同一φ迁移。槽位重叠、坐标往返不一致或端点无法无损输运时，拒绝该视图并保留原始样本。'],
261:['表3-1提供了迁移前的真实原生对象；3.3.4节另用一个可逐项计算的教学文档展示回填、φ、审核结果及监督掩码。二者的证据身份分别保留，不将教学改写称为模型实际输出。'],
269:['每篇最多四次尝试、最多两个接受视图；记录模型、提示、原始判断、掩码、对应、拒绝原因及全部成本。原文和变体按源文档共同分组。因素掩码是本章声明的默认设计，其漂移和覆盖由3.6.3节独立诊断；不能依据测试结果反向选择掩码策略。'],
281:['增强视图通过检查后，训练还需利用“不同表达对应相同论辩功能”的可靠对应。以表3-1为例，若T4被改写且通过组件检查，其Premise标签可以参与增强侧监督；T4→T3的supports与T4→T7的NONE则须各自通过关系检查，不能由组件通过推定二者都保持。成对目标比较这些可靠因素的类别分布，不要求新旧文本具有相同词序、长度或坐标。两个视图也可能共同预测错误，因此分布对齐必须与金标锚定配合。'],
294:['其中fC、fR为两层感知机，Pool为源成员表示的算术均值。η包含源成员数、源起点的最小值和均值、源终点的最大值和均值，以及目标起止位置；这里的位置是由规范字符区间派生的评价token坐标，并除以文档token数。它们仅为输入特征，不参与规范对象去重，亦不替代字符残差。η不包含待预测的关系标签；二元关系使用单一源表示，超边以完整源集合和目标作为监督对象，状态包含NONE。两个辅助头在开发评价和测试时移除。'],
351:['在表3-3的模拟结果中，T1/I0相对T0/I0的关系F1增加3.63 pp和4.89 pp，两项主要比较的Holm p均为0.0120。该比较评估整套训练配置，不能将全部差值归因于成对项。表内“±”表示三个种子条件下的样本标准差，反映该模拟条件中的跨种子离散程度，不能据此扩大到真实模型效果。'],
357:['判断成对目标的独立作用，应优先查看同一视图、掩码和辅助头条件下的T1I0−NoPair。在模拟结果中，两语料的点估计分别为1.98 pp和2.51 pp，单项95%区间不含零；但14项主要比较共同校正后，Holm p为0.0600和0.0720，未达到预设0.05标准。未调整的单项区间与家族校正p并不矛盾。辅助头平均JS下降也不能推出每个因素或完整生成图更准确，故本章保留成对正则化的研究价值，不将其表述为已经独立证实的关键增益。组件项和关系项的差值分别计算，不相加置信区间。'],
358:['可靠因素构造相对普通增强的探索差值为正，相对QOS式适配的区间则包含零或边界零，尚不足以确认普遍优越性。完整筛选与中心选择没有形成跨语料一致优势。因此，因素掩码、可选成对目标和可选完整筛选应按各自的受控比较解释；当前模拟材料主要用于检验研究设计及报告口径。'],
367:[
 '表3-5和表3-6的Wilson区间按因素级二项比例计算，仅作为工作性近似。来自同一源文档的多个视图或因素可能相关，且各层抽样比例不同；在未提供逐层源文档数及按源聚类、按抽样设计计算的区间前，不用这些区间承诺层内不安全率低于5%。一致率与κ只反映审核者间一致，不等同于语义准确率，NONE层的类别偏斜也可能压低κ。',
 'AAEC跨边界正关系层在123个被接受因素中有2个不安全因素，比例为1.63%，Wilson上界为5.74%，即使采用这一近似也未达到“上界不超过5%”的条件。总体加权不安全率0.81%不能替代该层结论。π=1表示审核覆盖了该固定因素池的全层，其比例描述这一有限池；推断未来生成批次仍需要新的来源与相应抽样设计。'],
370:[
 '表3-7沿用模拟资源配置：源访问为5N，更新步为5×ceil(N/32)；基础适配和带辅助头配置的可训练参数分别为44,310,528和63,460,376，基础权重冻结。相同有效更新步数不意味着相同总成本，增强的生成、审核和训练开销应共同报告。',
 '在模拟成本表中，T1I0的GPU时间由继承配置的1.69／2.86小时增加至3.53／5.98小时，约为原来的2.09倍。生成和审核开销由共享增强视图的组共同承担；T1I0相对CompPair仍增加训练时间。因此，采用成对目标需要同时考虑独立收益的不确定性与额外成本，不能仅凭总F1较高忽略这一代价。训练损失曲线用于诊断优化过程，不能单独证明严格抽取收益；真实复核应保留逐种子日志及明确的聚合定义。'],
378:['表3-8的不安全率以被接受因素为分母，漂移率以可判定且原标注有效的因素为分母；undecided与original_disagreement不并入稳定。原稿模拟口径采用逆抽样概率加权，并按源文档聚类重采样2000次。总体源组数为257和401，不能据此补出每一层的源组数或区间。真实审核仍需提供因素总体、抽样标记、源文档与视图身份、裁决及训练掩码，才能复核加权与聚类计算。'],
383:['真实语料案例与教学增强示例承担不同作用。表3-1的essay001说明原生因素的身份与依赖，3.3.4节的构造示例说明掩码和归一化的计算；二者均不能证明目标增强模型已经生成并接受了同样的视图。完整实证仍需补交一条部分掩码后接纳的实际记录和一条因语义变化而拒绝的实际记录。'],
384:['CDCP 00662的原生双源reason在4.6.4节展示。两个单源真子集既不能匹配完整超边，也不能仅据此获得语义NONE标签。这一边界贯穿训练审核和第四章完整对象评分。'],
389:['本章将局部改写后的监督保持分解为字符身份迁移、语义审核及可靠因素学习。组件双射保证标注对象可追踪，关系掩码使不确定监督不会重新进入增强侧损失，源文档归一化控制了视图数量造成的权重差异。模拟比较给出了整套训练配置的正向结果，但尚未在主要比较族内独立确认成对项的收益。教学示例已明确接口计算；真实增强记录、分层聚类审核及目标模型实验仍是验证这些方法效果所需的证据。'],
396:[
 '第三章处理训练文本改变后哪些监督可以继承，本章转向推理阶段：在原文保持不变时，如何修改预测组件及其关系。组件边界或类型一旦改变，引用旧端点的关系就成为不同的严格对象。仅依据组件置信度先行提交，可能遗漏整条关系失效的代价；仅依据图中平均覆盖分数接受，又可能因分母变化产生表面增益。',
 'MCR-BR已支持类型与跨度修改[34]。在此基础上，本章以固定候选宇宙和完整对象评分为前提，将冲突闭合、关系重建、分项效用及编辑预算组成统一接受规则。I1表示依赖闭合决策及固定路由的完整配置，BR-J专指集合局部子程序；共同接受、集合动作与求解路径分别通过对照评价。',
 '下文先说明候选和初始图，再给出评分及接受准则，随后用完整事务解释依赖闭合，最后介绍局部与全局求解、方法性质和实验分析。本章沿用test稿件的模拟主结果；教学事务只检验计算规则，CPU程序案例保留为独立的实际诊断，三者不互相替代。'],
397:['AAEC essay098的CPU实际预测同时包含一条正确supports关系及多条误边。固定规则能够减少误边，也可能删除原有正确关系，4.6.4节并列展示这两个结果。该现象说明修正需要记录收益和代价，但不等于本章目标模型、固定对象评分器或ILP已在该文档上得到验证。'],
441:['原稿所述CasiMedicos实际诊断复用既有的C层候选，未启用没有额外聚合覆盖收益的D层边界扩展。其证据限于该有限提案诊断；CDCP完整源组提案器仍属于待实现并经训练验收的目标配置，不因主结果表完整而视为已经落实。'],
475:['评分器采用固定文本编码、跨度MLP与有向端点组合MLP，参数不与生成器共享。它学习真实候选中完整对象的正确性；第三章辅助头则学习给定监督因素的类型与关系状态。两者的参数、标签含义与用途均有区别，参数分离本身不保证误差独立，参见第二章2.3节。'],
464:['在上述规约中，跨度或类型改变都会产生新的语义键。旧键所关联的关系全部计入删除，新键参与的合法完整关系另行计入加入。4.4.3节给出可复算的构造事务，说明这一规则如何进入接受和预算；实际任务效果仍由事后严格评价判断。'],
514:['局部路径先固定一个通过冲突检查的组件提议，再为受影响关系求合格补全；全局路径同时决定全部组件与关系。两条路径共用冻结评分、分项门及整图约束，但局部累计编辑与全局净编辑的可行域不同。以下依次给出局部执行、全局模型、方法性质及规模路由。'],
551:['可取两个兼容旧组件c1=[0,4)、c2=[4,8)，候选新组件b=[0,8)与二者均冲突。三者类型相同、关系集合为空，κC=1、wC=wR=1。令qC(c1)=qC(c2)=0.65、qC(b)=0.95，则旧组件效用为0.30+0.30=0.60，新组件效用为0.90，整体增量为0.30，严格超过默认δ=0.20；关系增量为0，实际成本为dC=3、dR=0。',
 '删除任何一个旧组件的增量均为−0.30，不能先通过严格改进门；直接插入b或以b替换一个旧组件，都必须经冲突闭合同时删除另一个旧组件，因此实际成为2→1事务，不属于单对象邻域。允许2→1动作且组件预算至少为3时，整体替换可以接受。该构造的关键在于整体增量严格大于δ，且单对象动作不能通过同一接受门；等于δ的增量仍应拒绝。'],
584:['规模判定取n=|V|+|P|，即裁剪后含受保护对象的组件与关系决策变量数；无环序变量、约束数量、CDCP源连续性约束数量及破同分调用另行记录。当前模拟配置规定n≤2500使用全局路径，超过该值使用有界局部路径。附录H的2.71／4.12秒P95及3000变量超时均属于合成规模扫描，仅用于展示阈值选择规则，不作为真实大池部署验收。本章主评价全部走全局路径，真实大池分支仍须单独验证。'],
594:['各对照在模拟主条件中固定T1检查点、候选快照、对象评分和合法化规则。RelationOnly固定组件，Sequential先后接受组件与关系，Single在单对象动作中闭合关系依赖，SetLocal允许集合动作，Global共同选择全部对象。I1按规模路由，但所报告690个文档—种子输入均小于2500变量，全部走Global。因此本节结论限定于这一规模与评分条件，不能从Single或SetLocal独立消融推断大池路由已经验收。'],
605:['模拟结果中，I1相对I0的AAEC、CDCP关系F1增量分别为4.48 pp、8.37 pp，两项主要比较通过Holm校正。RelationOnly已达到72.61和50.67，完整I1为73.29和52.02；组件可变性的附加点估计为正，但对应校正p为0.1924和0.1239，未独立确认其收益。Single−Sequential的区间跨零，SetLocal与Single的报告结果相同。因此整套配置增益、共同接受的规约价值及集合动作的表达能力必须分别解释。',
 '表4-4的95%区间是单项未调整区间，主要比较p按第二章2.5节的14项比较族作Holm校正。区间不含零但校正p超过0.05并不构成计算矛盾；真实复核还需要逐文档联合计数，仅有摘要统计不足以重建联合重采样分布。'],
638:['以下保留原稿中的实际CPU程序诊断。官方测试文档essay098的全文哈希为f68e231bdf7576fe87d57e801350c2b0ef2f513d51e29938b223671a392f528c；CPU-BIO-SGD在训练侧拟合组件和关系模型。种子17原始输出含1条正确关系和54条误边。后处理只使用训练侧类型—关系签名，并为同一预测源保留中点距离最近的目标，平局按语义键排序，评价前固定规则。该处理不读取测试金标决定删边，也没有执行本章的完整对象评分与ILP事务。'],
649:['两个案例来自同一实际处理流程的事后挑选，只解释成功与失败机制，不能估计发生率。原稿指向完整690条CPU诊断记录；它们与本章模拟主评价的文档—种子数相同，仍是不同运行和证据来源。真实主方法轨迹尚需补交初始键、候选来源、分数、闭合差集、预算、求解返回和提交后的严格计数，不能用此CPU记录替代。'],
667:['以下回到原稿的模拟评分条件分析。主条件与弱评分版本v0的检查点、候选和评分快照不同，分别报告，不作为同一冻结输入的单因素增量。v0在κR=1时不满足开发侧原TP损失≤5%的选择规则，输出保持初始图，关系F1为39.81；强制修正仅用于说明失败。'],
678:['本章将组件修改及关系身份变化纳入同一个可复核决策，给出了固定对象评分、完整差集计费、分项门与回退语义。局部和全局路径分别具有累计与净编辑预算，不能将它们的质量差异全部归因于搜索效率。模拟主比较显示整套推理配置的正向结果，但未独立确认共同接受、集合动作或相对等预算重生成的收益。全部主池走全局路径，大池局部路由和目标模型的真实事务仍待验证；构造事务说明规则可计算，CPU成功与失败案例说明实际后处理存在收益与误删，两类材料均不扩大主方法的证据边界。']}

def P(text):return {'kind':'text','text':text}
def H(text,level=2):return {'kind':'heading','text':text,'level':level}
def T(caption,rows):return [P(caption),{'kind':'table_new','rows':rows}]
aug=json.loads((OUT/'augmentation_example.json').read_text())
augblocks=[H('3.3.4 从改写到训练的完整示例'),
 P('以下为构造的教学文档，审核状态也按教学目的设定；它不来自目标模型调用。用a、b、c、d追踪四个组件，原始正关系为b→a、c→b，均标为supports，b→d与c→d为NONE。为便于手算，本例声明的因素宇宙仅含这四项，不作为AAEC全量因素枚举的实例。给定中心a，一跳完整入源规则得到A={a,b}；影响域含前三项，c→b是跨边界正因素，c→d位于域外。'),
 *T('表3-A 教学示例的原文 改写与全部组件迁移',[
 ['ID及类型','原字符区间与文字','新字符区间与文字'],
 *[[k+' / '+aug['old'][k]['key'][3],f"[{aug['old'][k]['key'][1]},{aug['old'][k]['key'][2]}) "+aug['old'][k]['text'],f"[{aug['new'][k]['key'][1]},{aug['new'][k]['key'][2]}) "+aug['new'][k]['text']] for k in aug['old']]]),
 P('四段组件文字以一个ASCII空格连接，坐标从0开始，使用半开区间。新文本仅改写a、b；c、d的文字保持不变，其坐标仍随前文长度变化更新。φ分别把原文版本demo_aug下的四个规范键映射至demo_aug_v1下的新键，关系的每一个端点同步迁移。坐标与字符切片的往返计算保存在augmentation_example.json。'),
 *T('表3-B 教学示例的审核状态与损失入口',[
 ['监督对象','构造审核状态','掩码','增强侧用途'],
 ['a b c d组件','含义与角色均通过','mC=1','全部参与头监督；a b c参与组件对齐'],
 ['b→a supports','原新标签一致','mR=1','正关系监督与成对项'],
 ['c→b supports','关系待定','mR=0','不进入增强关系监督或对齐，不改成NONE'],
 ['b→d NONE','原新标签一致','mR=1','NONE监督与成对项'],
 ['域外c→d NONE','不在影响域','mR=0','保留原文监督 不进入增强侧关系损失']]),
 P('该视图通过组件门，虽有一项关系待定，仍由默认因素掩码策略接纳。设这篇源文档仅有一个接受视图，则md=1，头监督为(S原文+Sm增强)/2。组件对齐集合A+={a,b,c}，每项先取1/3；d只参与可靠NONE端点，不因此进入组件对齐。关系对齐的正关系组和NONE组各有一项，组权重各为1/2；乘默认λR=0.5后，两项各占0.25。原文完整生成及句子监督的权重不变，增强输运图不作为完整生成答案。'),
 P('再考虑第二次构造尝试，将a改为“Teamwork does not help children learn social skills.”。新增否定改变原组件含义，整视图在组件门被拒绝，不进入Views，不增加md，也不参与任何增强损失；原文监督及第一条已接受视图保留。如果随后另有一个真实接受视图，md才增至2，头监督的每视图权重变为1/3，成对项再按两个接受视图平均。'),
 P('该示例使坐标、掩码与权重可被复算，却不验证模型生成或语义审核的实际可靠性。真实机制验证仍需提供相同字段的目标流程原始输出，尤其是“部分掩码接纳”与“语义变化拒绝”各一条。')]

txblocks=[H('4.4.3 一次完整事务与两种拒绝情形'),
 P('下面采用独立的构造文档demo_tx。字符区间、类型、候选来源、分数和事后金标均为教学设定，不是Qwen预测或ILP运行轨迹。它的用途是从固定输入复算闭合、分项增量、预算与返回规则；伴随脚本枚举全部16种关系选择，仅核对这个小例子的数学决策。'),
 *T('表4-A 构造事务的组件与关系评分',[
 ['对象','完整对象定义','来源及初始状态','固定q'],
 ['a','(demo_tx,0,8,Claim)','构造初始对象 已选','0.90'],
 ['b','(demo_tx,10,20,Premise)','构造初始对象 已选','0.55'],
 ['c','(demo_tx,22,32,Premise)','构造初始对象 已选','0.85'],
 ['b′','(demo_tx,10,21,Premise)','构造补充提议 未选','0.80'],
 ['r1','(b,a,supports)','构造初始关系 已选','0.55'],
 ['r2','(c,b,supports)','构造初始关系 已选','0.45'],
 ['r3','(b′,a,supports)','冻结池内新关系 未选','0.80'],
 ['r4','(c,b′,supports)','冻结池内新关系 未选','0.70']]),
 P('关系表中的a、b、b′、c均展开为同表完整组件键后再形成关系键；显示ID不参与身份判断。初始C={a,b,c}、R={r1,r2}。名义提议为B={b′}、D0=∅；由于b′与b重叠，冲突闭合得到D={b}，实际动作成为1→1，dC=2。受影响关系空间Ptx={r1,r2,r3,r4}覆盖b的全部入出关系及b′的新候选。新组件状态下r1、r2必须删除，r3、r4可以共同选择；这条链满足本例的端点、非重叠、标签互斥与无环规则。'),
 P('取κC=κR=wC=wR=1、δ=0.20，则a(q)=2q−1。组件增量为(2×0.80−1)−(2×0.55−1)=0.50；关系旧效用为0.10−0.10=0，新效用为0.60+0.40=1.00，故ΔΦR=1.00、ΔΦ=1.50。两个分项非负，总增量严格超过0.20。删除r1、r2和加入r3、r4共计dR=4。若剩余预算为组件4、关系8，提交后分别为2、4；仅在全部复核通过后一次性提交C′和R′。'),
 *T('表4-B 同一闭合事务的接受与拒绝',[
 ['情形','输入变化','检查结果','返回与编辑账本'],
 ['接受构造','使用表4-A的全部分数','ΔC=0.50 ΔR=1.00 总计1.50','原子提交 累计增加C=2 R=4'],
 ['缺分构造','r4没有有效分数','事务空间评分不齐备 不把缺分设为0','拒绝并保持原图 语义编辑不增加'],
 ['关系分项下降构造','r1 r2 r3改为q=0.60 r4改为0.40','最优补全只选r3 ΔC=0.50 ΔR=−0.20 总计0.30','即使总增量过门仍拒绝 原图不变']]),
 P('拒绝不退还已经消耗的候选评分、枚举或求解资源，但不增加语义编辑。构造金标另设为组件{a,b′,c}及关系{r3,r4}，只用于提交后的演示计数：组件TP/FP/FN从2/1/1变为3/0/0，关系从0/2/2变为2/0/0。金标未参与接受条件；由于金标本身也是构造的，这些变化不作为真实质量证据。'),
 P('小例子的状态、分数、预算及拒绝原因保存在transaction_examples.json，复算脚本为replay_ch03_ch04_examples.py。真实主流程仍须另交生成候选原始输出、冻结评分版本、求解器状态及证书、事务日志和事后计数。4.6.4节的CPU实际案例使用另一种规则，不附加本例分数冒充主方法执行。')]

# Sequence edits consolidate headings without hiding paragraphs behind new bold subheadings.
def selected(indices):
 out=[]
 for i in indices:
  e=elements[i]
  if e.tag not in [qn('w:p'),qn('w:tbl')]:continue
  if e.tag==qn('w:p') and not txt(e) and not e.xpath('.//w:drawing'):continue
  if i in heads:
   if heads[i] is not None:out.append(H(heads[i],1 if sty(e)=='Heading3' else 2))
   continue
  if i in repl:out += [P(t) for t in repl[i]];continue
  out.append({'kind':'source','index':i})
 return out

ch3=selected(range(208,276+1))+augblocks+selected(range(280,368))
# Audit drift belongs next to its sampling design; resource cost follows it.
ch3+=selected(range(375,380))+[H('3.6.4 训练成本与方法局限')]+selected(range(368,371))+selected(range(383,385))+selected(range(388,390))
ch4=selected(range(393,442))+selected(range(468,505))+selected(range(445,465))+txblocks
ch4+=selected(range(513,540))+selected(range(563,579))+selected(range(542,561))+selected(range(581,588))+selected(range(591,679))
pos=next(i for i,x in enumerate(ch4) if x.get('index')==427)+1
ch4[pos:pos]=[P('这里u表示文档版本与字符半开区间组成的规范跨度(d,b,e)，故带类型候选(u,t)展开后为(d,b,e,t)。token范围与残差只用于派生表示和无损往返，不以相同token外包区间合并不同字符对象；关系键再由完整源组件键集合、目标组件键及原生标签构成。候选去重、评分缓存、冲突检查及严格评价均引用这一身份。')]
# Make the scope hierarchy explicit at the end of the mask definition.
pos=next(i for i,x in enumerate(ch3) if x.get('kind')=='source' and x.get('index')==276)+1
ch3[pos:pos]=[P('将通过审核且端点组件也通过检查的关系集合记为$R_A$，则各层满足'),
 P('$R_A ⊆ J_A ⊆ D_A ⊆ Ω_{sem},\\quad Ω_{sem}=Ω∖U_{sem}.$'),
 P('完整影响域$D_A$是定义范围，$J_A$是实际审核范围，$R_A$才是增强关系损失的入口。默认$J_A$通常小于$D_A$；完整筛选才令$J_A=D_A$，语义未知$U_{sem}$始终在两者之外。图3-1的两个分支应按这一集合关系理解。')]
# Framework paragraph before implementation inherited details.
pos=next(i for i,x in enumerate(ch3) if x.get('text')=='3.2 总体框架与基础模型')+1
ch3[pos:pos]=[P('整体框架包括增强构造与参数学习两部分。构造阶段读取训练文档和原生标注，输出新全文、全部组件的对应φ以及可靠因素掩码；学习阶段保留原文完整生成监督，并在增强侧只学习通过检查的因素。图3-2展示二者进入共享生成器后的连接。类型条件LoRA、线索提示和句子监督构成继承基础，本章新增部分是关系影响域约束、可靠监督接口及受控的成对学习目标。')]

# Native OMML to readable Markdown mathematics, retaining document-native math in Word.
def local(e):return etree.QName(e).localname
def child(e,n):return e.find('m:'+n,ns)
def val(e,path,default=''):
 x=e.find(path,ns);return x.get(qn('m:val'),default) if x is not None else default
def mx(e):
 if e is None:return ''
 k=local(e)
 if k.endswith('Pr') or k in ['rPr','ctrlPr']:return ''
 if k=='t':return (e.text or '').replace('−','-')
 if k in ['sSub','sSup','sSubSup']:
  s='{'+mx(child(e,'e'))+'}'
  if k!='sSup':s+='_{'+mx(child(e,'sub'))+'}'
  if k!='sSub':s+='^{'+mx(child(e,'sup'))+'}'
  return s
 if k=='f':return '\\frac{'+mx(child(e,'num'))+'}{'+mx(child(e,'den'))+'}'
 if k=='rad':return '\\sqrt'+('['+mx(child(e,'deg'))+']' if mx(child(e,'deg')) else '')+'{'+mx(child(e,'e'))+'}'
 if k=='d':return '\\left'+val(e,'m:dPr/m:begChr','(')+mx(child(e,'e'))+'\\right'+val(e,'m:dPr/m:endChr',')')
 if k=='nary':
  s={'∑':'\\sum','∏':'\\prod','⋃':'\\bigcup','∫':'\\int'}.get(val(e,'m:naryPr/m:chr','∫'),val(e,'m:naryPr/m:chr','∫'))
  if val(e,'m:naryPr/m:subHide')!='1' and mx(child(e,'sub')):s+='_{'+mx(child(e,'sub'))+'}'
  if val(e,'m:naryPr/m:supHide')!='1' and mx(child(e,'sup')):s+='^{'+mx(child(e,'sup'))+'}'
  return s+' '+mx(child(e,'e'))
 if k in ['acc','bar']:return '\\'+('bar' if k=='bar' or val(e,'m:accPr/m:chr')=='̄' else 'hat')+'{'+mx(child(e,'e'))+'}'
 if k in ['limLow','limUpp']:return mx(child(e,'e'))+('_{' if k=='limLow' else '^{')+mx(child(e,'lim'))+'}'
 if k=='func':return mx(child(e,'fName'))+' '+mx(child(e,'e'))
 if k=='eqArr':return '\\begin{aligned}'+r' \\ '.join(mx(x) for x in e if local(x)=='e')+'\\end{aligned}'
 if k=='m':return '\\begin{matrix}'+r' \\ '.join(' & '.join(mx(c) for c in row if local(c)=='e') for row in e if local(row)=='mr')+'\\end{matrix}'
 return ''.join(mx(c) for c in e)
def pm(e):
 def walk(n):
  k=local(n)
  if k=='oMath':return '$'+mx(n)+'$'
  if k=='t' and etree.QName(n).namespace==ns['w']:return n.text or ''
  if k=='br':return '\n'
  if k=='tab':return '    '
  if k in ['pPr','rPr','drawing','bookmarkStart','bookmarkEnd']:return ''
  return ''.join(walk(c) for c in n)
 return walk(e)

def setup(chapter):
 d=Document();s=d.sections[0]
 s.page_width=Inches(8.27);s.page_height=Inches(11.69)
 s.top_margin=Inches(.76);s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.78)
 s.header_distance=s.footer_distance=Inches(.30)
 for st in d.styles:
  if st.type==1:
   st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
   st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
   for z in st._element.xpath('.//w:pBdr'):z.getparent().remove(z)
 normal=d.styles['Normal'];normal.font.size=Pt(10.5)
 pf=normal.paragraph_format;pf.line_spacing=1.16;pf.space_after=Pt(6);pf.first_line_indent=Pt(21)
 for name,size in [('Title',16),('Heading 1',13),('Heading 2',11.5)]:
  st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
  pf=st.paragraph_format;pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.space_before=Pt(12);pf.space_after=Pt(6)
 for name in ['SourceCode','Caption']:
  if name not in d.styles:d.styles.add_style(name,1)
  st=d.styles[name];st.font.name='Arial Unicode MS';st.font.size=Pt(8.5 if name=='SourceCode' else 10)
  st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.line_spacing=1.05
  st.paragraph_format.space_after=Pt(1 if name=='SourceCode' else 6)
 s.header.paragraphs[0].text=f'test稿件  第{chapter}章修订审阅版  主结果沿用模拟说明'
 s.header.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
 for r in s.header.paragraphs[0].runs:r.font.size=Pt(8)
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 return d

def normalize_p(e,style='Normal',small=False):
 # Only paragraph/run formatting changes; math subtrees retain source OOXML.
 old=e.find(qn('w:pPr'))
 if old is not None:e.remove(old)
 pr=OxmlElement('w:pPr');s=OxmlElement('w:pStyle');s.set(qn('w:val'),style.replace(' ',''));pr.append(s);e.insert(0,pr)
 for z in list(e.iter()):
  if local(z) in ['commentRangeStart','commentRangeEnd','commentReference','bookmarkStart','bookmarkEnd']:
   z.getparent().remove(z)
 for r in e.findall('.//w:r',ns):
  rp=r.find(qn('w:rPr'))
  if rp is not None:r.remove(rp)
 return e

def add_text(d,s,style='Normal'):
 p=d.add_paragraph(style=style)
 for piece in re.split(r'(\$[^$]+\$)',s):
  if piece.startswith('$') and piece.endswith('$'):p._p.append(math_node(piece[1:-1]))
  else:p.add_run(piece)
 if captionpat.match(s):
  p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if s.startswith('$') and s.endswith('$'):
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
 return p

def add_table(d,rows,source=None):
 n=len(rows[0]);t=d.add_table(rows=len(rows),cols=n);t.autofit=False
 if n==3:weights=[1.25,2.65,2.8]
 elif n==4 and max(len(c) for row in rows for c in row)>65:weights=[.9,1.0,1.05,3.76]
 elif n==4:weights=[1.15,2.45,1.6,1.51]
 elif n==5:weights=[1.85,1.215,1.215,1.215,1.215]
 elif n==7:weights=[.58,1.69,.52,1.03,1.12,.52,1.25]
 elif n==8:weights=[.75,.7,.85,.85,.7,1.0,.85,1.01]
 elif n==9:weights=[.54,.85,.63,.6,.65,1.25,.67,.78,.74]
 elif n==10:weights=[1.16,.44,.51,.59,.52,.57,.55,.99,.72,.66]
 else:weights=[1]*n
 widths=[v*6.71/sum(weights) for v in weights]
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 pr=t._tbl.tblPr;bd=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');bd.append(b)
 pr.append(bd)
 sourcerows=source.findall(qn('w:tr')) if source is not None else None
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge,v in [('top',65),('bottom',65),('left',65),('right',65)]:
    a=OxmlElement('w:'+edge);a.set(qn('w:w'),str(v));a.set(qn('w:type'),'dxa');mar.append(a)
   tcpr.append(mar)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');tcpr.append(sh)
   if source is None:cell.text=rows[i][j]
   else:
    for p in list(cell._tc):
     if p.tag!=qn('w:tcPr'):cell._tc.remove(p)
    sc=sourcerows[i].findall(qn('w:tc'))[j]
    for p in sc.findall(qn('w:p')):cell._tc.append(normalize_p(deepcopy(p)))
   for p in cell.paragraphs:
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(rows[i][j])<22 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.size=Pt(8 if n>=8 else 8.5 if n>=5 else 9);r.bold=i==0
 return t

def render_chapter(ch,blocks):
 d=setup(ch);md=[];used=[];stats={'source_equations':0,'new_equations':0,'source_tables':0,'new_tables':0,'images':0}
 for x in blocks:
  if x['kind']=='heading':
   add_text(d,x['text'],'Heading '+str(x['level']));md.append('#'*(x['level']+1)+' '+x['text']);continue
  if x['kind']=='text':
   s=tabreplace(x['text']);add_text(d,s);md.append(s);stats['new_equations']+=len(re.findall(r'\$[^$]+\$',s));continue
  if x['kind']=='table_new':
   rows=[[tabreplace(cell) for cell in row] for row in x['rows']];add_table(d,rows);stats['new_tables']+=1
   md.append('\n'.join(['|'+'|'.join(r)+'|' for r in [rows[0],['---']*len(rows[0]),*rows[1:]]]));continue
  i=x['index'];e=deepcopy(elements[i]);used.append(i)
  normalize_math_layout(e)
  stats['source_equations']+=len(e.xpath('.//m:oMath'))
  if e.tag==qn('w:tbl'):
   rows=[[txt(c) for c in r.findall(qn('w:tc'))] for r in e.findall(qn('w:tr'))]
   add_table(d,rows,e);stats['source_tables']+=1
   mdrows=[['<br>'.join(pm(p) for p in c.findall(qn('w:p'))).replace('|',r'\|').replace('\n','<br>') for c in r.findall(qn('w:tc'))] for r in e.findall(qn('w:tr'))]
   md.append('\n'.join(['|'+'|'.join(r)+'|' for r in [mdrows[0],['---']*len(rows[0]),*mdrows[1:]]]));continue
  for t in e.findall('.//w:t',ns):
   if t.text:t.text=tabreplace(refreplace(t.text))
  if e.xpath('.//w:drawing'):
   for blip in e.xpath('.//a:blip'):
    rid=blip.get(qn('r:embed'));part=original.part.related_parts[rid]
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
    imgpath=OUT/f'ch{ch}_{rid}.png';imgpath.write_bytes(part.blob)
    p.add_run().add_picture(BytesIO(part.blob),width=Inches(6.65));md.append(f'![原稿图示]({imgpath})');stats['images']+=1
   continue
  st=sty(e);s=txt(e)
  if st=='Heading2':
   add_text(d,s,'Title');md.append('# '+s);continue
  if st in ['Heading3','Heading4']:
   level=1 if st=='Heading3' else 2;add_text(d,s,'Heading '+str(level));md.append('#'*(level+1)+' '+s);continue
  if st=='SourceCode':
   ct=code_text(e)
   for line in ct.splitlines():add_text(d,line,'SourceCode')
   md.append('```text\n'+ct+'\n```');continue
  normalize_p(e);d._element.body.insert(len(d._element.body)-1,e)
  pp=d.paragraphs[-1]
  if captionpat.match(s):
   pp.paragraph_format.first_line_indent=Pt(0);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
   pp.paragraph_format.keep_with_next=s.startswith(('表','算法'))
  if e.xpath('./m:oMathPara'):
   pp.paragraph_format.first_line_indent=Pt(0);pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
  md.append(pm(e))
 title='第三章' if ch==3 else '第四章'
 path=OUT/f'{title}_修订稿_v1.docx';d.save(path)
 markdown='\n\n'.join(md)+'\n';(OUT/f'{title}_修订后全文_v1.md').write_text(markdown)
 stats['total_native_equations']=len(d._element.xpath('.//m:oMath'))
 stats['headings']=[{'style':p.style.name,'text':p.text} for p in d.paragraphs if p.style.name.startswith('Heading')]
 stats['source_indices']=used;stats['word_paragraphs']=len(d.paragraphs)
 stats['source_sha256']=sha(SRC);stats['docx_sha256']=sha(path)
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith('.xml'):etree.fromstring(z.read(name))
 (OUT/f'ch{ch}_draft_audit.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2))
 return markdown,stats

if len(sys.argv)==1 or sys.argv[1]=='3':md3,s3=render_chapter(3,ch3)
else:s3=json.loads((OUT/'ch3_draft_audit.json').read_text())
if len(sys.argv)==1 or sys.argv[1]=='4':md4,s4=render_chapter(4,ch4)
else:s4=json.loads((OUT/'ch4_draft_audit.json').read_text())
(OUT/'章节引用映射.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2))
(OUT/'表号映射.json').write_text(json.dumps(tablemap,ensure_ascii=False,indent=2))
# Original uncertain chart remains available for later source reconciliation.
(OUT/'原图3-3_聚合与色带定义待核对.png').write_bytes((B/'ch03/rId86.png').read_bytes())
(OUT/'修订块清单.json').write_text(json.dumps({'ch3':ch3,'ch4':ch4},ensure_ascii=False,indent=2))
print(json.dumps({'ch3':{k:v for k,v in s3.items() if k not in ['headings','source_indices']},'ch4':{k:v for k,v in s4.items() if k not in ['headings','source_indices']}},ensure_ascii=False))
