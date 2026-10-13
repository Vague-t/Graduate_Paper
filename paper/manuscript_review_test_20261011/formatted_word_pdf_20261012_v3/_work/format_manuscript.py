from pathlib import Path
from copy import deepcopy
import json, hashlib, re
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

O=Path(__file__).resolve().parents[1]
B=O.parent
SOURCE=B/'integrated_revision_20261012_v2/论文整合修订稿_图示公式与术语_v2.docx'
OUT=O/'论文规范排版稿_v3.docx'
doc=Document(SOURCE)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','a':'http://schemas.openxmlformats.org/drawingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
xp=lambda e,s:etree.XPath(s,namespaces=ns)(e)
txt=lambda e:''.join(n.text or '' for n in xp(e,'.//w:t|.//m:t'))
wt=lambda e:''.join(n.text or '' for n in xp(e,'.//w:t'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
edits=[]

def replace(old,new):
    hits=[p for p in doc.paragraphs if old in wt(p._p)]
    assert len(hits)==1,(old,len(hits))
    p=hits[0];before=txt(p._p)
    # Materialize the old explanatory REF fields before removing their phrases.
    for f in xp(p._p,'.//w:fldSimple'):
        if 'REF Equation_' in f.get(qn('w:instr'),''):
            at=p._p.index(f)
            for c in list(f):p._p.insert(at,c);at+=1
            p._p.remove(f)
    nodes=xp(p._p,'.//w:t');s=''.join(n.text or '' for n in nodes)
    start=s.index(old);end=start+len(old);offset=0
    for n in nodes:
        t=n.text or '';a=offset;b=a+len(t);offset=b
        if b<=start or a>=end:continue
        n.text=t[:max(0,start-a)]+(new if a<=start<b else '')+t[max(0,end-a):]
        n.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    edits.append({'before':before,'after':txt(p._p)})

for a,b in [
 ('；组件掩码与关系掩码分别记录，具体判定见3.3.3节。','；组件掩码与关系掩码分别记录。'),
 ('式（3-1）在第二章式（2-6）的LoRA参数化上区分两组低秩矩阵：下标g','两组低秩矩阵的下标g'),
 ('教师强制训练沿用2.2.1节的序列条件，同样只使用','教师强制训练只使用'),
 ('NONE（无边状态沿用2.3节的定义）','NONE'),
 ('金标超边（多前提关系，见2.1.1节）','金标超边'),
 ('组件采用第二章定义的规范身份','组件采用规范身份'),
 ('式（3-7）中的MeanPool','MeanPool'),
 ('式（3-8）的关系头','关系头'),
 ('式（3-10）中的JS散度沿用第二章式（2-7），用于衡量两个视图的预测分布差异；两侧均保留梯度','对齐时两侧均保留梯度'),
 ('式（3-4）的完整影响域','完整影响域'),
 ('总损失见式（3-14）：','总损失为：'),
]:replace(a,b)

def run_font(r,size=12,east='宋体',west='Times New Roman',bold=None):
    pr=r.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');r.insert(0,pr)
    rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    rf.attrib.clear()
    for k,v in [('ascii',west),('hAnsi',west),('eastAsia',east),('cs',west)]:rf.set(qn('w:'+k),v)
    for k,val in [('sz',str(round(size*2))),('szCs',str(round(size*2))),('color','000000')]:
        el=pr.find(qn('w:'+k))
        if el is None:el=OxmlElement('w:'+k);pr.append(el)
        el.attrib.clear();el.set(qn('w:val'),val)
    if bold is not None:
        el=pr.find(qn('w:b'))
        if el is None:el=OxmlElement('w:b');pr.append(el)
        el.set(qn('w:val'),'1' if bold else '0')

def style(name,size,east='宋体',bold=False):
    st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().get_or_add_rFonts();rf.attrib.clear()
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia',east)]:rf.set(qn('w:'+k),v)
    pf=st.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5;pf.widow_control=True
    for e in xp(st.element,'.//w:pBdr'):e.getparent().remove(e)
    return st

for name,size,east,bold in [('Normal',12,'宋体',False),('Body Text',12,'宋体',False),('Title',22,'黑体',True),('Heading 1',18,'黑体',True),('Heading 2',15,'黑体',True),('Heading 3',12,'黑体',True),('Heading 4',12,'黑体',True),('Caption',10.5,'宋体',True),('Algorithm',10.5,'宋体',False),('TOC 1',11,'宋体',True),('TOC 2',10.5,'宋体',False),('TOC 3',10.5,'宋体',False),('Front Heading',18,'黑体',True),('Running Header',9,'宋体',False),('Bibliography',10.5,'宋体',False)]:style(name,size,east,bold)

# Resolve dangling Pandoc paragraph styles before formatting or PDF conversion.
known_styles={s.style_id for s in doc.styles}
for p in doc.paragraphs:
    raw=p._p.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
    if raw is not None and raw.get(qn('w:val')) not in known_styles:p.style='Body Text'
    if p.style.name.startswith('Heading') and re.match(r'^[A-J]\.\d+\.\d+\b',p.text):p.style='Heading 3'
for name in ('Normal','Body Text'):
    doc.styles[name].paragraph_format.line_spacing=Pt(22)
for s in doc.sections:
    for grid in list(s._sectPr.findall(qn('w:docGrid'))):s._sectPr.remove(grid)

width=Cm(16)
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
sec.header_distance=sec.footer_distance=Cm(1.25)

# Separate the cover while retaining the existing title, date and material status.
title=next(p for p in doc.paragraphs if p.text=='面向关系感知学习与依赖闭合修正的端到端论辩挖掘方法研究')
title.style='Title'
date=next(p for p in doc.paragraphs if p.text=='2026年10月')
status=next(p for p in doc.paragraphs if p.text.startswith('评审功能测试用完整模拟稿'))
english=next(p for p in doc.paragraphs if p.text.startswith('English Title:'))
title._p.addnext(english._p);english._p.addnext(date._p);date._p.addnext(status._p)
abstract=next(p for p in doc.paragraphs if p.text=='摘要')
english_abstract=next(p for p in doc.paragraphs if p.text=='Abstract')
first_chapter=next(p for p in doc.paragraphs if p.text.startswith('第一章 '))

in_refs=False;in_algorithm=False
for p in doc.paragraphs:
    text=txt(p._p).strip();pf=p.paragraph_format
    pf.left_indent=pf.right_indent=Pt(0);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=Pt(22)
    pf.first_line_indent=Pt(24);pf.widow_control=True;pf.keep_together=False;pf.keep_with_next=False
    if p.style.name.startswith('Heading'):
        level=int(p.style.name[-1]);pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.keep_together=True;pf.line_spacing=1.2
        pf.space_before=Pt(18 if level==1 else 12);pf.space_after=Pt(12 if level==1 else 6)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
        size=18 if level==1 else 15 if level==2 else 12
        for r in xp(p._p,'.//w:r'):run_font(r,size,'黑体',bold=True)
        if level==1:pf.page_break_before=True
        if text=='参考文献':in_refs=True
        if text.startswith('附录'):in_refs=False
    else:
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in xp(p._p,'.//w:r'):run_font(r,12)
        if any(len(txt(e))>=30 for e in xp(p._p,'./m:oMath')) or re.sub(r'\s','',text).startswith(('该效用具有明确的代价解释','取κC=','4.4.2节的','以4.4.2节的','依次执行','逐文档计数的主键')):
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if in_refs and text!='参考文献':
        p.style='Bibliography';pf.first_line_indent=Pt(-21);pf.left_indent=Pt(21);pf.line_spacing=1.25;pf.space_after=Pt(6)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.keep_together=True
        for r in xp(p._p,'.//w:r'):run_font(r,10.5)
    if p.style.name=='Caption' or re.match(r'^(图\d-\d+|表[\dA-J]-\d+[ab]?|补充表S2-\d+)\s',text):
        p.style='Caption';pf.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing=1.2;pf.space_before=Pt(6);pf.space_after=Pt(6);pf.keep_together=True;pf.keep_with_next=text.startswith(('表','补充表'))
        for r in xp(p._p,'.//w:r'):run_font(r,10.5,bold=True)
    if text.startswith('算法3-1 ') or text.startswith('算法4-1 ') or text.startswith('算法4-2 '):in_algorithm=True
    if text.startswith('每篇文档的参考预算') or xp(p._p,'.//w:drawing'):in_algorithm=False
    if in_algorithm:
        p.style='Algorithm';pf.first_line_indent=Pt(0);pf.left_indent=Cm(.2);pf.line_spacing=1.2;pf.space_after=Pt(2)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in xp(p._p,'.//w:r'):run_font(r,10.5)
        if text.startswith('算法'):
            pf.keep_with_next=True;pf.space_before=Pt(6);pf.space_after=Pt(6);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if xp(p._p,'.//w:drawing'):
        pf.first_line_indent=Pt(0);pf.line_spacing=1;pf.space_before=Pt(6);pf.space_after=Pt(0);pf.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if xp(p._p,'.//w:bookmarkStart[starts-with(@w:name,"Equation_")]'):
        pf.first_line_indent=Pt(0);pf.line_spacing=1.15;pf.space_before=Pt(8);pf.space_after=Pt(8);pf.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;tabs=pf.tab_stops;tabs.clear_all()
        tabs.add_tab_stop(Cm(7.45),WD_TAB_ALIGNMENT.CENTER);tabs.add_tab_stop(width,WD_TAB_ALIGNMENT.RIGHT)
        for r in xp(p._p,'.//w:r'):run_font(r,10.5)
        token=''.join(n.text or '' for n in xp(p._p,'.//m:t'))
        size=11.5 if len(token)<80 or xp(p._p,'.//m:eqArr') else 10 if len(token)<125 else 9.5
        for r in xp(p._p,'.//m:r'):
            wr=r.find(qn('w:rPr'))
            if wr is None:wr=OxmlElement('w:rPr');r.insert(0,wr)
            sz=wr.find(qn('w:sz'))
            if sz is None:sz=OxmlElement('w:sz');wr.append(sz)
            sz.set(qn('w:val'),str(round(size*2)))
    if text=='B.6 固定评分性质的证明':pf.page_break_before=False

for p in (abstract,english_abstract):
    p.style='Heading 1';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf=p.paragraph_format
    pf.page_break_before=p is english_abstract;pf.first_line_indent=Pt(0);pf.space_before=Pt(18);pf.space_after=Pt(18);pf.keep_with_next=True
    for r in xp(p._p,'.//w:r'):run_font(r,18,'黑体',bold=True)
for p in (title,english,date,status):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.keep_with_next=False;pf.keep_together=True;pf.page_break_before=False
title.paragraph_format.space_before=Pt(105);title.paragraph_format.space_after=Pt(28);title.paragraph_format.line_spacing=1.4
english.paragraph_format.space_after=Pt(70);english.paragraph_format.line_spacing=1.4
date.paragraph_format.space_after=Pt(55)
status.paragraph_format.line_spacing=1.5;status.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for r in xp(title._p,'.//w:r'):run_font(r,22,'黑体',bold=True)
for r in xp(english._p,'.//w:r'):run_font(r,14,bold=False)
for r in xp(date._p,'.//w:r'):run_font(r,14,bold=False)
for r in xp(status._p,'.//w:r'):run_font(r,10.5,bold=False)
title.text='面向关系感知学习与依赖闭合修正的\n端到端论辩挖掘方法研究'
for r in xp(title._p,'.//w:r'):run_font(r,22,'黑体',bold=True)
for p in doc.paragraphs:
    if p.text.startswith(('关键词：','Keywords:')):
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.space_before=Pt(12)
        p.paragraph_format.line_spacing=Pt(18)
        for r in xp(p._p,'.//w:r'):run_font(r,10.5)

# Set all table widths, retain each original relative column allocation.
for table in doc.tables:
    grid=table._tbl.find(qn('w:tblGrid'));cols=list(grid)
    old=[int(c.get(qn('w:w'))) for c in cols];total=sum(old)
    sizes=[round(9071*v/total) for v in old]
    before=table._tbl.getprevious()
    caption=txt(before) if before is not None else ''
    for prefix, allocation in {
        '表4-8 ':[730,1360,790,790,790,790,790,1450,1581],
        '表4-9 ':[640,1100,420,670,670,800,700,720,800,760,800,991],
        '表4-12 ':[650,1120,720,720,720,900,900,1190,1190,961],
        '表4-13 ':[850,1900,1900,1100,1100,1100,1121],
        '表A-5 ':[730,1090,3625,3626],
    }.items():
        if caption.startswith(prefix):
            assert len(allocation)==len(cols) and sum(allocation)==9071
            sizes=allocation
    for c,w in zip(cols,sizes):c.set(qn('w:w'),str(w))
    pr=table._tbl.find(qn('w:tblPr'))
    tw=pr.find(qn('w:tblW'))
    if tw is None:tw=OxmlElement('w:tblW');pr.append(tw)
    tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9071')
    for ri,row in enumerate(table.rows):
        rp=row._tr.get_or_add_trPr()
        if rp.find(qn('w:cantSplit')) is None:rp.append(OxmlElement('w:cantSplit'))
        if ri==0 and rp.find(qn('w:tblHeader')) is None:rp.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cp=cell._tc.get_or_add_tcPr();cw=cp.find(qn('w:tcW'))
            if cw is not None:cw.set(qn('w:w'),str(sizes[min(ci,len(sizes)-1)]))
            for mar in list(cp.findall(qn('w:tcMar'))):cp.remove(mar)
            mar=OxmlElement('w:tcMar')
            for k,v in [('top',55),('bottom',55),('left',65),('right',65)]:
                e=OxmlElement('w:'+k);e.set(qn('w:type'),'dxa');e.set(qn('w:w'),str(v));mar.append(e)
            cp.append(mar)
            for p in cell.paragraphs:
                pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=pf.space_after=Pt(0);pf.line_spacing=1.15;pf.keep_with_next=False;pf.keep_together=True
                size=9 if len(cols)>=9 else 9.5 if len(cols)>=6 else 10.5
                for r in xp(p._p,'.//w:r'):run_font(r,size,bold=ri==0)

for drawing in xp(doc.element,'.//w:drawing'):
    ext=xp(drawing,'.//wp:extent')[0];cx=int(ext.get('cx'));cy=int(ext.get('cy'))
    ratio=min(1,int(width)/cx,int(Cm(20))/cy)
    for e in xp(drawing,'.//wp:extent|.//a:xfrm/a:ext'):
        if e.get('cx') and e.get('cy'):
            e.set('cx',str(round(int(e.get('cx'))*ratio)));e.set('cy',str(round(int(e.get('cy'))*ratio)))

def pnew(text,style_name='Normal'):
    p=OxmlElement('w:p');doc.element.body.insert(-1,p);p=Paragraph(p,doc._body);p.style=style_name
    if text:p.add_run(text)
    return p

# Bookmarks and a cached TOC with native page references.
# Repair the source's unmatched chapter bookmark and stray end marker.
legacy_start=xp(doc.element,'.//w:bookmarkStart[@w:id="39"]')
legacy_end=xp(doc.element,'.//w:bookmarkEnd[@w:id="39"]')
orphan_end=xp(doc.element,'.//w:bookmarkEnd[@w:id="152"]')
if legacy_start and not legacy_end:
    first_chapter._p.insert(1,legacy_start[0])
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'39');first_chapter._p.append(end)
if not xp(doc.element,'.//w:bookmarkStart[@w:id="152"]'):
    for end in orphan_end:end.getparent().remove(end)
headings=[]
for p in doc.paragraphs:
    if p is abstract or p._p is abstract._p:level=1
    elif p._p is english_abstract._p:level=1
    elif p.style.name in ['Heading 1','Heading 2','Heading 3']:level=int(p.style.name[-1])
    else:continue
    if not txt(p._p).strip():continue
    name=f'ThesisHeading_{len(headings)+1:03}'
    bs=OxmlElement('w:bookmarkStart');bs.set(qn('w:id'),str(5000+len(headings)));bs.set(qn('w:name'),name)
    be=OxmlElement('w:bookmarkEnd');be.set(qn('w:id'),str(5000+len(headings)))
    p._p.insert(1,bs);p._p.append(be)
    headings.append({'bookmark':name,'title':txt(p._p),'level':level})

toc_title=pnew('目录','Front Heading');toc_title.paragraph_format.page_break_before=True;toc_title.paragraph_format.space_before=Pt(18);toc_title.paragraph_format.space_after=Pt(18);toc_title.alignment=WD_ALIGN_PARAGRAPH.CENTER
first_chapter._p.addprevious(toc_title._p)
toc_begin=pnew('');toc_begin.paragraph_format.line_spacing=Pt(1);toc_begin.paragraph_format.space_before=Pt(0);toc_begin.paragraph_format.space_after=Pt(0)
run=toc_begin.add_run();beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');run._r.append(beg)
inst=OxmlElement('w:instrText');inst.text=' TOC \\o "1-3" \\h \\z ';run._r.append(inst)
sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');run._r.append(sep)
first_chapter._p.addprevious(toc_begin._p)
cached_path=O/'_work/toc_pages.json'
cached=json.loads(cached_path.read_text()) if cached_path.exists() else {}
for h in headings:
    p=pnew('',f"TOC {h['level']}");pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Cm(.5*(h['level']-1));pf.line_spacing=1.15;pf.space_after=Pt(3);pf.keep_with_next=False
    pf.tab_stops.add_tab_stop(width,WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),h['bookmark']);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=h['title'];r.append(t);link.append(r);p._p.append(link)
    p.add_run('\t');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),f" PAGEREF {h['bookmark']} \\h ")
    r=OxmlElement('w:r');v=OxmlElement('w:t');v.text=cached.get(h['bookmark'],'1');r.append(v);f.append(r);p._p.append(f)
    first_chapter._p.addprevious(p._p)
toc_end=pnew('');toc_end.paragraph_format.line_spacing=Pt(1)
r=toc_end.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end);first_chapter._p.addprevious(toc_end._p)

# Cover, Roman-numbered preliminaries, and Arabic-numbered body sections.
base_sect=deepcopy(doc.element.body.sectPr)
for ref in list(base_sect):
    if ref.tag in [qn('w:headerReference'),qn('w:footerReference'),qn('w:pgNumType'),qn('w:titlePg')]:base_sect.remove(ref)
def section_end(before, fmt=None):
    p=pnew('');p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(0)
    s=deepcopy(base_sect)
    typ=s.find(qn('w:type'))
    if typ is None:typ=OxmlElement('w:type');s.append(typ)
    typ.set(qn('w:val'),'nextPage')
    if fmt:
        pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),fmt);pg.set(qn('w:start'),'1');s.append(pg)
    p._p.get_or_add_pPr().append(s);before.addprevious(p._p)
section_end(abstract._p)
section_end(first_chapter._p,'upperRoman')
final=doc.element.body.sectPr
for ref in list(final):
    if ref.tag in [qn('w:headerReference'),qn('w:footerReference'),qn('w:pgNumType'),qn('w:titlePg')]:final.remove(ref)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:fmt'),'decimal');pg.set(qn('w:start'),'1');final.append(pg)
first_chapter.paragraph_format.page_break_before=False
for i,s in enumerate(doc.sections):
    s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
    for part in [s.header,s.footer]:
        for p in part.paragraphs:
            for c in list(p._p):
                if c.tag!=qn('w:pPr'):p._p.remove(c)
    if i==0:continue
    h=s.header.paragraphs[0];h.style='Running Header';h.alignment=WD_ALIGN_PARAGRAPH.CENTER;h.paragraph_format.first_line_indent=Pt(0)
    h.add_run('面向关系感知学习与依赖闭合修正的端到端论辩挖掘方法研究')
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Pt(0);f.paragraph_format.line_spacing=1
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),' PAGE ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);run_font(r,10.5);fld.append(r);f._p.append(fld)

up=doc.settings.element.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');doc.settings.element.append(up)
up.set(qn('w:val'),'true')
doc.core_properties.comments='规范排版版：去除额外术语回指，保留公式编号及原稿证据说明。'
doc.save(OUT)
(O/'_work/format_manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':sha(SOURCE),'output':str(OUT),'output_sha256':sha(OUT),'edits':edits,'headings':headings,'format':{'paper':'A4','margins_cm':2.5,'body_pt':12,'body_line_spacing_pt':22,'cjk':'宋体','latin':'Times New Roman','headings':'黑体'}},ensure_ascii=False,indent=2))
print(json.dumps({'output':str(OUT),'edits':len(edits),'toc_entries':len(headings),'sections':len(doc.sections)},ensure_ascii=False))
