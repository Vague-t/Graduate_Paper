from pathlib import Path
from copy import deepcopy
from io import BytesIO
import re, json, hashlib, ast, sys
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

B=Path(__file__).resolve().parents[2]
C=B/'chapter_revision_20261012'
O=Path(__file__).resolve().parents[1]
F=B/'model_architecture_editable_svg_20261012_v3'
OUT=O/'论文整合修订稿_图示公式与术语_v2.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','a':'http://schemas.openxmlformats.org/drawingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
xp=lambda e,s:etree.XPath(s,namespaces=ns)(e)
txt=lambda e:''.join(n.text or '' for n in xp(e,'.//w:t|.//m:t'))
wtext=lambda e:''.join(n.text or '' for n in xp(e,'.//w:t'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
sources=[B/'manuscript_review_test.docx',C/'ch01/v4/第一章_修订稿_v4.docx',C/'ch02/v2/第二章_修订稿_v2.docx',C/'ch03_ch04_revision_v1/第三章_修订稿_v1.docx',C/'ch03_ch04_revision_v1/第四章_修订稿_v1.docx']
d=Document(sources[0]); src=[Document(p) for p in sources[1:]]
ledger=[]

def replace(e,old,new,kind='术语与说明',required=True):
    nodes=xp(e,'.//w:t'); s=''.join(n.text or '' for n in nodes)
    if old not in s:
        if required: raise AssertionError((old,s[:200]))
        return False
    before=txt(e); start=s.index(old);end=start+len(old);off=0
    for n in nodes:
        t=n.text or '';a=off;b=off+len(t);off=b
        if b<=start or a>=end:continue
        n.text=t[:max(0,start-a)]+(new if a<=start<b else '')+t[max(0,end-a):]
        n.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    ledger.append({'type':kind,'before':before,'after':txt(e)})
    return True

def find(elements,prefix):
    matches=[e for e in elements if e.tag==qn('w:p') and txt(e).startswith(prefix)]
    assert len(matches)==1,(prefix,len(matches))
    return matches[0]

def pnew(text,style='Normal'):
    p=OxmlElement('w:p'); pp=OxmlElement('w:pPr');st=OxmlElement('w:pStyle');st.set(qn('w:val'),style);pp.append(st);p.append(pp)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;r.append(t);p.append(r);return p

def add_after(elements,prefix,text):
    e=find(elements,prefix);elements.insert(elements.index(e)+1,pnew(text))
    ledger.append({'type':'新增释义','after_anchor':prefix,'after':text})

def clone(e,part):
    n=deepcopy(e)
    for el in n.iter():
        for attr,rid in list(el.attrib.items()):
            if not attr.startswith('{'+ns['r']+'}'):continue
            rel=part.rels.get(rid)
            if rel is None:raise ValueError(('missing rel',rid))
            if rel.is_external:new=d.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
            elif rel.reltype==RT.IMAGE:new,_=d.part.get_or_add_image(BytesIO(rel.target_part.blob))
            else:raise ValueError(('unsupported',rel.reltype))
            el.set(attr,new)
    return n

chapters=[]
for doc in src:
    chapters.append([clone(e,doc.part) for e in doc.element.body if e.tag!=qn('w:sectPr')])
supp_start=next(i for i,e in enumerate(chapters[1]) if txt(e)=='第二章补充材料')
supp=chapters[1][supp_start:];chapters[1]=chapters[1][:supp_start]
replace(supp[0],'第二章补充材料','附录J 公共实验配置与统计补充','合并定位')
# The approved S2 labels remain stable, with one explicit destination note.
supp.insert(1,pnew('本附录承接第二章补充材料，保留S2节号与补充表号，供正文交叉引用。'))
add_after(chapters[1],'各项比较共享骨干','第二章补充说明S2.1—S2.5及补充表S2-1—S2-4统一收录于附录J。')

t=chapters[2]
replace(find(t,'在端到端论辩挖掘中'),'原本没有标注关系的因素','原本没有标注关系的组件组合')
add_after(t,'本章研究在组件位置变化','这里将同一篇原文及其改写版本称为不同“视图”；可单独决定是否参与监督的判断称为“因素”，包括一个组件的类型判断，或一组完整端点之间的关系状态判断。成对学习是在原文与改写的对应因素上比较模型输出，使通过检查的同一监督在两种表达下保持一致。')
add_after(t,'整体框架包括增强构造','因素掩码是随监督保存的0/1标记：1表示该因素参与相应的增强侧监督或对齐，0表示不参与。它不遮蔽输入文字，也不删除原始金标；组件掩码与关系掩码分别记录，具体判定见3.3.3节。')
add_after(t,txt(t[t.index(find(t,'本节说明继承的专项建模'))+1]),'式（3-1）在第二章式（2-6）的LoRA参数化上区分两组低秩矩阵：下标g表示全局分支，下标c表示结论专项分支；h为当前投影的输入表示，h′为输出表示，位置下标t表示当前序列位置。')
replace(find(t,'全局分支参与所有位置'),'字段按 id、type、span 排列','字段按 id（组件编号）、type（类型）、span（原文跨度）排列')
replace(find(t,'全局分支参与所有位置'),'教师强制训练也仅使用可见前缀。','教师强制训练沿用2.2.1节的序列条件，同样只使用当前位置之前已提供的输出前缀，不能读取后续答案来开启门控。')
replace(find(t,'生成目标对有效回答'),'token（含 EOS）','token（含序列结束符EOS）')
replace(find(t,'句子辅助任务沿用'),'沿用 ICIC 专项监督思路','沿用继承配置ICIC的句子监督')
add_after(t,'本章的文本改写是训练数据增强','改写“槽位”指已定位到原文字符区间、允许被替换的组件片段；“只读端点”指文字保持不变的其他组件。只读并不表示坐标固定，前文改写造成长度变化时，它们的字符位置也要更新。下文“正关系”指具有明确原生关系标签的关系，与无边状态NONE相对；它也包括attacks，不能理解为仅含支持关系。')
replace(find(t,'表内T4'),'T4→T7属于跨边界NONE。','T4→T7属于跨边界NONE（无边状态沿用2.3节的定义）。')
replace(find(t,'优先集合为空时'),'超过三源成员的金标超边','超过三个源成员的金标超边（多前提关系，见2.1.1节）')
replace(find(t,'记 Ω') if any(txt(e).startswith('记 Ω') for e in t) else find(t,'记 '),'在可审核因素宇宙内','在可审核的有限端点组合集合内')
replace(find(t,'训练因素宇宙包含'),'训练因素宇宙包含','训练所用的有限端点组合集合包含')
replace(find(t,'NONE 仅在上述'),'需按元数诊断','需按源元数（源成员个数）诊断')
add_after(t,'验证器分别判断原文因素','“待定”表示审核尚不能确认原有状态是否保持，“未审核”表示没有执行该项判断。两者都不能提供已确认的增强监督，掩码取0；掩码为0并不等于关系标签为NONE。下文的“语义未知”另指联合源真子集缺少可推出的单独语义标签，其定义与审核状态分开。')
for e in t:
    if '组件门' in wtext(e):replace(e,'组件门','组件一致性检查')
replace(find(t,'增强视图通过检查后'),'分布对齐必须与金标锚定配合。','分布对齐必须与金标锚定配合，即同时用已知原生标签监督预测，不能只要求两个视图彼此相似。')
replace(find(t,'组件类型头和关系状态头'),'组件类型头和关系状态头以共享生成器的输入表示为基础','组件类型头和关系状态头是读取生成器隐藏表示的辅助分类器，以共享生成器的输入表示为基础')
add_after(t,'因果模型的早期','式（3-7）中的MeanPool对组件覆盖的输入子词表示取算术均值；方括号及分号表示向量拼接。组件头输出类型分布，式（3-8）的关系头则输出完整端点组合的关系状态分布。')
# Avoid a redundant JS definition already provided in Chapter 2.
js=find(t,'其中 JS')
before=txt(js)
idx=t.index(js);t[idx]=pnew('式（3-10）中的JS散度沿用第二章式（2-7），用于衡量两个视图的预测分布差异；两侧均保留梯度，不将任一视图视为固定目标。')
ledger.append({'type':'去除重复定义并回指','before':before,'after':txt(t[idx])})
replace(find(t,'两名审核者A、B'),'stable、changed、undecided与original_disagreement','stable（状态保持）、changed（状态改变）、undecided（待定）与original_disagreement（原标签分歧）')
add_after(t,'两名审核者A、B','表3-7和表3-8用Wilson区间表示二项比例的不确定性，用κ表示扣除偶然一致后的审核者一致性。这里的κ是审核统计量，与第四章效用中的代价参数不是同一量。')

# Update stale citations only in source ranges which predate the chapter revisions.
secmap=json.loads((C/'ch03_ch04_revision_v1/章节引用映射.json').read_text())
tablemap={'2-2':'S2-1','2-3':'2-4','2-4':'S2-2','2-5':'S2-3','3-2':'3-4','3-3':'3-5','3-4':'3-6','3-5':'3-7','3-6':'3-8','3-7':'3-10','3-8':'3-9','4-1':'4-3','4-2':'4-4','4-3':'4-5','4-4':'4-6','4-5':'4-7','4-6':'4-8','4-7':'4-9','4-8a':'4-10','4-8b':'4-11','4-9':'4-12','4-10':'4-13','4-11':'4-14','4-12':'4-15','4-13':'4-16','4-14':'4-17','4-15':'4-18'}
for e in chapters[0]:
    if '表3-4' in wtext(e):replace(e,'表3-4','表3-6','跨章引用')
    if '表4-4' in wtext(e):replace(e,'表4-4','表4-6','跨章引用')
for e in supp:
    for a,b in [('表3-4','表3-6'),('表4-4','表4-6')]:
        if a in wtext(e):replace(e,a,b,'跨章引用')
base=list(d.element.body)
tail=[deepcopy(e) for e in base[682:] if e.tag!=qn('w:sectPr')]
for el in tail:
    for e in ([el] if el.tag==qn('w:p') else xp(el,'.//w:p')):
        old=wtext(e)
        new=re.sub(r'表([234]-\d+[ab]?)(?!\d)',lambda m:'表'+tablemap.get(m[1],m[1]),old)
        # All tokens in the appendix reading-index table are section references.
        index_row=el.tag==qn('w:tbl') and '方法入口' in txt(el) and '结果或解释' in txt(el)
        pattern=r'(?<![\d.])([234]\.\d+(?:\.\d+){0,2})(?![\d.])' if index_row else r'(?<![\d.])([234]\.\d+(?:\.\d+){0,2})(?=节)'
        new=re.sub(pattern,lambda m:secmap.get(m[1],m[1]),new)
        if new!=old:replace(e,old,new,'跨章引用')
        if index_row and '2.5.1' in wtext(e):replace(e,'2.5.1','附录J的S2.1','跨章引用')
        if index_row:
            for a,b in [('3.6.3真实审核诊断','3.6.3审核与漂移诊断'),('4.6.2真实候选覆盖','4.6.2候选覆盖统计'),('4.6.4逐项评分与最终输出','4.6.3评分质量与缺分诊断')]:
                if a in wtext(e):replace(e,a,b,'阅读索引同步')
old=find(tail,'集合动作的表达能力可用抽象数学构造证明')
newtext='取两个兼容旧组件c1=[0,4)、c2=[4,8)，以及与二者均冲突的新组件b=[0,8)。三者类型相同，当前与候选关系集合均为空。令κC=1、wC=wR=1、δ=0.20，qC(c1)=qC(c2)=0.65、qC(b)=0.95。旧效用为0.30+0.30=0.60，新效用为0.90，故ΔΦC=0.30、ΔΦR=0，总增量严格大于δ。组件预算至少为3时，2→1集合事务可以接受。'
replace(old,wtext(old),newtext,'同步此前附录B.6补丁')
add_after(tail,newtext,'单独删除任一旧组件的效用增量为−0.30，不能通过严格改进门。插入b或仅声明替换一个旧组件，都必须经冲突闭合删除另一个旧组件，其实际形态仍是2→1，因而不属于单对象邻域。该构造说明集合动作的可达性，不证明真实F1改善。')
for el in tail:
    if el.tag==qn('w:tbl') and txt(el).startswith('符号或术语含义使用边界'):
        row=el.findall(qn('w:tr'))[2];cells=row.findall(qn('w:tc'));before=txt(row)
        inner=xp(cells[0],'.//m:d/m:e')[0]
        xp(inner,'.//m:t')[0].text='b'
        for pos,token in enumerate(['d',',']):
            r=OxmlElement('m:r');v=OxmlElement('m:t');v.text=token;r.append(v);inner.insert(pos,r)
        replace(cells[1],'跨度起止位置及类型组成的组件','文档版本、字符起止位置及原生类型组成的规范组件','跨章符号一致')
        replace(cells[2],'显示ID不替代跨度和类型构成的身份','显示ID不替代文档版本、字符区间与类型构成的身份','跨章符号一致')
        ledger.append({'type':'跨章符号一致','before':before,'after':txt(row)})

# Assemble without losing the original package or untouched later chapters.
sect=deepcopy(d.element.body.sectPr)
for e in list(d.element.body):d.element.body.remove(e)
allblocks=[deepcopy(e) for e in base[:20]]+sum(chapters,[])+tail+supp
for e in allblocks:d.element.body.append(e)
d.element.body.append(sect)

# Original pandoc tables omit tcW and reference unavailable Table/Compact styles.
# Supply explicit geometry without touching any cell text or native mathematics.
manual_widths={
 '表5-1':[25,66,79], '表5-2':[28,35,35,36,36],
 '表5-3':[14,45,18,30,46,17], '表5-4':[13,16,18,18,18,18,18,18,13],
 '表5-5a':[14,27,20,20,36,53], '表5-5b':[14,18,18,27,27,19,23,24],
 '表5-6':[14,32,18,18,18,70],
 '表B-1':[40,15,15,15,85], '表F-1':[35,65,70],
 '表F-2':[40,59,71], '表H-4':[16,24,20,22,22,22,22,22],
 '表H-5':[19,33,20,22,22,22,32], '表I-1':[25,24,23,19,20,21,20,18],
 '表I-2':[18,37,28,67,20],
}
lastcap=''
for e in d.element.body:
    if e.tag==qn('w:p') and re.match(r'^表[5-6A-I]-',txt(e)):lastcap=txt(e).split()[0]
    if e.tag!=qn('w:tbl') or e.find(qn('w:tblPr')).find(qn('w:tblStyle')) is None:continue
    style=e.find(qn('w:tblPr')).find(qn('w:tblStyle'))
    if style.get(qn('w:val'))!='Table':continue
    table=Table(e,d._body);ncol=len(table.columns)
    weights=manual_widths.get(lastcap)
    if weights is None or len(weights)!=ncol:
        widths=[]
        for col in range(ncol):
            ts=[row.cells[col].text for row in table.rows]
            lens=[sum(1.8 if ord(c)>127 else 1 for c in s) for s in ts]
            widths.append(max(9,min(48,max(lens[:1]+sorted(lens[1:])[max(0,len(lens)//2):max(1,len(lens)//2+1)]))))
        weights=widths
    twips=[round(9638*w/sum(weights)) for w in weights]
    pr=e.find(qn('w:tblPr'))
    for child in list(pr):pr.remove(child)
    wid=OxmlElement('w:tblW');wid.set(qn('w:type'),'dxa');wid.set(qn('w:w'),'9638');pr.append(wid)
    layout=OxmlElement('w:tblLayout');layout.set(qn('w:type'),'fixed');pr.append(layout)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    pr.append(borders)
    grid=e.find(qn('w:tblGrid'))
    for c in list(grid):grid.remove(c)
    for w in twips:
        c=OxmlElement('w:gridCol');c.set(qn('w:w'),str(w));grid.append(c)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr()
        for h in list(trpr):
            if h.tag==qn('w:trHeight'):trpr.remove(h)
        if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
        if ri==0 and trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cp=cell._tc.get_or_add_tcPr();cw=cp.find(qn('w:tcW'))
            if cw is None:cw=OxmlElement('w:tcW');cp.insert(0,cw)
            cw.set(qn('w:type'),'dxa');cw.set(qn('w:w'),str(twips[ci]))
            margins=OxmlElement('w:tcMar')
            pad=40 if lastcap.startswith('表5-') else 55
            for k,v in [('top',pad),('bottom',pad),('left',65),('right',65)]:
                z=OxmlElement('w:'+k);z.set(qn('w:type'),'dxa');z.set(qn('w:w'),str(v));margins.append(z)
            cp.append(margins)
            va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');cp.append(va)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');cp.append(sh)
            for p in cell.paragraphs:
                p.style='Normal';pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.06;pf.keep_with_next=False;pf.keep_together=True
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or len(p.text)<22 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.size=Pt(8.5);r.font.name='Arial Unicode MS';r.font.bold=(ri==0)

# The original A4 width and margins are retained; source chapter hierarchy is normalized.
for name,size in [('Normal',10.5),('Body Text',10.5),('Title',17),('Heading 1',15),('Heading 2',13),('Heading 3',11.5),('Heading 4',10.5),('Caption',10),('SourceCode',8.5)]:
    if name not in d.styles:continue
    st=d.styles[name];st.font.name='Arial Unicode MS';st.font.size=Pt(size)
    rpr=st._element.get_or_add_rPr();rf=rpr.get_or_add_rFonts();rf.set(qn('w:eastAsia'),'Arial Unicode MS')
    for attr in list(rf.attrib):
        if attr.endswith('Theme') or attr.endswith('theme'):del rf.attrib[attr]
    st.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    pf=st.paragraph_format
    if name in ['Normal','Body Text']:
        pf.first_line_indent=Pt(21);pf.line_spacing=1.16;pf.space_after=Pt(5)
    elif name.startswith('Heading'):
        pf.first_line_indent=Pt(0);pf.keep_with_next=True;pf.space_before=Pt(11);pf.space_after=Pt(6)
    elif name=='Caption':pf.first_line_indent=Pt(0);pf.line_spacing=1.05
    for bdr in xp(st._element,'.//w:pBdr'):bdr.getparent().remove(bdr)

flow_section=''
for e in list(d.element.body):
    if e.tag!=qn('w:p'):continue
    p=Paragraph(e,d._body);s=txt(e).strip()
    if not s and not xp(e,'.//w:drawing|.//w:br[@w:type="page"]|.//m:oMath'):
        d.element.body.remove(e);continue
    if re.match(r'^第[一二三四五六]章 ',s) or re.match(r'^附录[A-J] ',s) or s=='参考文献':
        p.style='Heading 1';p.paragraph_format.page_break_before=True
        flow_section=s.split()[0]
    elif re.match(r'^[1-6]\.\d+\.\d+\s',s):p.style='Heading 3'
    elif re.match(r'^[1-6]\.\d+\s',s):p.style='Heading 2'
    elif re.match(r'^[A-J]\.\d+(?:\.\d+)?\s',s) or re.match(r'^S2\.\d+\s',s):p.style='Heading 2'
    if s=='B.6 固定评分性质的证明':p.paragraph_format.page_break_before=True
    if p.style.name in ['Normal','Body Text'] and not xp(e,'.//w:drawing|.//m:oMathPara'):
        p.paragraph_format.space_after=Pt(4)
        if flow_section in ['第二章','第三章']:
            p.paragraph_format.line_spacing=1.14;p.paragraph_format.space_after=Pt(3)
        elif flow_section=='第五章' or flow_section.startswith('附录'):
            p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(3)
    if s=='Abstract':p.paragraph_format.page_break_before=True
    if xp(e,'.//w:drawing'):
        p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    if re.match(r'^(图[\dA-J]-\d+|表[\dA-J]-\d+[ab]?|补充表S2-\d+)\s',s):
        p.style='Caption';p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=s.startswith(('表','补充表'))
    if p.style.name.startswith('Heading') or p.style.name=='Caption':
        for r in p.runs:
            r.font.name='Arial Unicode MS'
            rf=r._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'Arial Unicode MS')
            for attr in list(rf.attrib):
                if attr.endswith('Theme') or attr.endswith('theme'):del rf.attrib[attr]

figmap=[('图2-1 ',F/'01_总体模型架构_可编辑','图2-1 端到端论辩挖掘的总体模型架构'),('图3-2 ',F/'02_可靠成对学习模型_可编辑','图3-2 关系影响域约束的可靠成对学习模型'),('图4-1 ',F/'03_依赖闭合修正模型_可编辑','图4-1 多候选驱动的依赖闭合修正模型')]
for idx,(prefix,stem,caption) in enumerate(figmap,1):
    cap=find(list(d.element.body),prefix);old=cap.getprevious();assert xp(old,'.//w:drawing')
    p=Paragraph(old,d._body)
    for e in list(old):
        if e.tag!=qn('w:pPr'):old.remove(e)
    shape=p.add_run().add_picture(str(stem.with_suffix('.png')),width=Cm(17.0))
    blip=xp(shape._inline,'.//a:blip')[0]
    svgpart=Part(PackURI(f'/word/media/model_architecture_{idx}.svg'),'image/svg+xml',stem.with_suffix('.svg').read_bytes(),d.part.package)
    rid=d.part.relate_to(svgpart,RT.IMAGE)
    extlst=OxmlElement('a:extLst');ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    svg=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
    svg.set(qn('r:embed'),rid);ext.append(svg);extlst.append(ext);blip.append(extlst)
    replace(cap,wtext(cap),caption,'架构图替换')
    xp(shape._inline,'.//wp:docPr')[0].set('descr',caption+'；内嵌SVG矢量图，并保留PNG兼容预览。')

# Identify independent native equations; paragraph-level inline symbols remain unnumbered.
eqs=[];counts={};chapter=None
for e in d.element.body:
    s=txt(e).strip()
    mt=re.match(r'^第([一二三四五六])章 ',s)
    if mt:chapter=str('一二三四五六'.index(mt[1])+1)
    if s.startswith('附录') or s=='参考文献':chapter=None
    if e.tag!=qn('w:p') or not xp(e,'.//m:oMath'):continue
    prose=wtext(e).strip()
    numbered=bool(re.fullmatch(r'[（(]\d+-\d+[）)]',prose))
    if not (not prose or numbered or xp(e,'./m:oMathPara')):continue
    assert chapter is not None,s
    counts[chapter]=counts.get(chapter,0)+1;seq=counts[chapter];label=f'{chapter}-{seq}'
    mats=[deepcopy(m) for m in xp(e,'.//m:oMath')]
    tokens=[''.join(x.text or '' for x in xp(m,'.//m:t')) for m in mats]
    p=Paragraph(e,d._body)
    for c in list(e):
        if c.tag!=qn('w:pPr'):e.remove(c)
    pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.keep_together=True;pf.keep_with_next=False;pf.line_spacing=1.15;pf.space_before=Pt(6);pf.space_after=Pt(6);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    tabs=pf.tab_stops;tabs.clear_all();tabs.add_tab_stop(Cm(8.9),WD_TAB_ALIGNMENT.CENTER);tabs.add_tab_stop(Cm(17.8),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    for m in mats:e.append(m)
    p.add_run('\t')
    bs=OxmlElement('w:bookmarkStart');bs.set(qn('w:id'),str(2000+len(eqs)));bs.set(qn('w:name'),'Equation_'+label.replace('-','_'));e.append(bs)
    p.add_run('（'+chapter+'-')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),f' SEQ Equation{chapter} '+('\\r 1 ' if seq==1 else '')+'\\* ARABIC ')
    r=OxmlElement('w:r');v=OxmlElement('w:t');v.text=str(seq);r.append(v);fld.append(r);e.append(fld);p.add_run('）')
    be=OxmlElement('w:bookmarkEnd');be.set(qn('w:id'),str(2000+len(eqs)));e.append(be)
    for r in p.runs:r.font.size=Pt(10.5)
    eqs.append({'label':label,'native_math_tokens':tokens,'preview':''.join(tokens)})
assert counts=={'2':10,'3':14,'4':18},counts

# Add useful equation references at the first procedural handoffs.
refs=[('完整影响域DA是定义范围','完整影响域','式（3-4）的完整影响域'),('成对项先在每个非空因素组内平均','总损失为','总损失见式（3-14）：')]
for prefix,old,new in refs:
    if old:
        matches=[e for e in d.element.body if txt(e).startswith(prefix)]
        if matches:replace(matches[0],old,new,'公式引用')

# Materialized REF fields keep prose references linked to the equation bookmarks.
ref_count=0
for e in d.element.body:
    if e.tag!=qn('w:p') or xp(e,'.//w:bookmarkStart[starts-with(@w:name,"Equation_")]'):continue
    for r in list(e):
        if r.tag!=qn('w:r') or len(r.findall(qn('w:t')))!=1:continue
        v=r.find(qn('w:t'));s=v.text or '';matches=list(re.finditer(r'（([234])-(\d+)）',s))
        if not matches:continue
        cursor=0;pos=e.index(r);rpr=r.find(qn('w:rPr'))
        for match in matches:
            if match.start()>cursor:
                nr=OxmlElement('w:r')
                if rpr is not None:nr.append(deepcopy(rpr))
                nt=OxmlElement('w:t');nt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');nt.text=s[cursor:match.start()];nr.append(nt);e.insert(pos,nr);pos+=1
            fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),f' REF Equation_{match[1]}_{match[2]} \\h ')
            nr=OxmlElement('w:r')
            if rpr is not None:nr.append(deepcopy(rpr))
            nt=OxmlElement('w:t');nt.text=match[0];nr.append(nt);fld.append(nr);e.insert(pos,fld);pos+=1;ref_count+=1;cursor=match.end()
        if cursor<len(s):
            nr=OxmlElement('w:r')
            if rpr is not None:nr.append(deepcopy(rpr))
            nt=OxmlElement('w:t');nt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');nt.text=s[cursor:];nr.append(nt);e.insert(pos,nr)
        e.remove(r)

# Unique drawing IDs and safe image sizes after multi-document merge.
for n,el in enumerate(xp(d.element,'.//wp:docPr'),1):el.set('id',str(n))
for drawing in xp(d.element,'.//w:drawing'):
    extents=xp(drawing,'.//wp:extent')
    if not extents:continue
    cx=int(extents[0].get('cx'));cy=int(extents[0].get('cy'));maxcx=int(Cm(17.5))
    if cx>maxcx:
        ratio=maxcx/cx
        for el in xp(drawing,'.//wp:extent|.//a:xfrm/a:ext'):
            if el.get('cx') and el.get('cy'):
                el.set('cx',str(round(int(el.get('cx'))*ratio)));el.set('cy',str(round(int(el.get('cy'))*ratio)))

settings=d.settings.element
up=settings.find(qn('w:updateFields'))
if up is None:up=OxmlElement('w:updateFields');settings.append(up)
up.set(qn('w:val'),'true')
# A single continuous page-number field supports review of the merged manuscript.
footer=d.sections[0].footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
if not footer.text and not xp(footer._p,'.//w:fldSimple'):
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),' PAGE ');r=OxmlElement('w:r');v=OxmlElement('w:t');v.text='1';r.append(v);fld.append(r);footer._p.append(fld)
d.core_properties.title='面向关系感知学习与依赖闭合修正的端到端论辩挖掘方法研究'
d.core_properties.comments='基于最新分章修订稿整合；本轮替换三张架构图、统一公式编号并补充第三章必要术语解释。'
d.save(OUT)
manifest={'output':str(OUT),'sha256':sha(OUT),'sources':[{'path':str(p),'sha256':sha(p)} for p in sources],'figures':[{'path':str(stem.with_suffix('.svg')),'sha256':sha(stem.with_suffix('.svg'))} for _,stem,_ in figmap],'equation_counts':counts,'equation_reference_fields':ref_count,'equations':eqs,'patches':ledger}
(O/'_work/整合与逐项修改记录.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
# Equation-only fixture permits a fast check without regenerating the entire book.
fixture=Document(OUT)
for e in list(fixture.element.body):
    if e.tag==qn('w:sectPr'):continue
    if not xp(e,'.//w:bookmarkStart[starts-with(@w:name,"Equation_")]'):fixture.element.body.remove(e)
fixture.save(O/'_work/公式排版核对.docx')
print(json.dumps({'path':str(OUT),'formulas':counts,'patches':len(ledger),'paragraphs':len(d.paragraphs),'tables':len(d.tables)},ensure_ascii=False))
