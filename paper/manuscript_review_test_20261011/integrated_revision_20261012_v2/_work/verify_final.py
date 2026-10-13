from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import hashlib, json, re, posixpath
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pypdf import PdfReader

O = Path(__file__).resolve().parents[1]
manifest = json.loads((O/'_work/整合与逐项修改记录.json').read_text())
out = Path(manifest['output'])
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
txt = lambda e: ''.join(e.xpath('.//w:t/text()|.//m:t/text()'))
wtxt = lambda e: ''.join(e.xpath('.//w:t/text()'))
maths = lambda e: [''.join(n.text or '' for n in m.iter(qn('m:t'))) for m in e.iter(qn('m:oMath'))]
assert sha(out) == manifest['sha256']
assert all(sha(Path(x['path'])) == x['sha256'] for x in manifest['sources'] + manifest['figures'])
docs = [Document(x['path']) for x in manifest['sources']]
doc = Document(out)

with ZipFile(out) as z:
    assert z.testzip() is None
    names = set(z.namelist())
    rel_count = 0
    for name in names:
        if not name.endswith('.rels'): continue
        root = etree.fromstring(z.read(name))
        base = posixpath.dirname(posixpath.dirname(name))
        for rel in root:
            if rel.get('TargetMode') == 'External': continue
            target = rel.get('Target')
            resolved = posixpath.normpath(posixpath.join(base, target)).lstrip('/')
            assert resolved in names, (name, target, resolved)
            rel_count += 1
    for i, f in enumerate(manifest['figures'], 1):
        assert z.read(f'word/media/model_architecture_{i}.svg') == Path(f['path']).read_bytes()

source_chapters = []
supp = []
for i, d in enumerate(docs[1:]):
    es = [e for e in d.element.body if e.tag != qn('w:sectPr')]
    if i == 1:
        n = next(i for i,e in enumerate(es) if txt(e) == '第二章补充材料')
        supp = es[n:]
        es = es[:n]
    source_chapters.append(es)

formula_checks = []
for chapter in (2, 3, 4):
    expected = []
    for e in source_chapters[chapter-1]:
        if e.tag != qn('w:p') or not e.xpath('.//m:oMath'): continue
        prose = wtxt(e).strip()
        if not prose or re.fullmatch(r'[（(]\d+-\d+[）)]', prose) or e.xpath('./m:oMathPara'):
            expected.append(maths(e))
    actual = []
    for e in doc.element.body:
        if e.tag != qn('w:p'): continue
        if e.xpath(f'.//w:bookmarkStart[starts-with(@w:name,"Equation_{chapter}_")]'):
            actual.append(maths(e))
    assert expected == actual, chapter
    formula_checks.append({'chapter': chapter, 'count': len(actual), 'native_tokens_equal_to_source': True})

expected_blocks = list(docs[0].element.body)[:20] + sum(source_chapters, []) + list(docs[0].element.body)[682:] + supp
expected_tables = [e for e in expected_blocks if e.tag == qn('w:tbl')]
actual_tables = [t._tbl for t in doc.tables]
assert len(expected_tables) == len(actual_tables) == 73
table_diffs = []
for i, (a,b) in enumerate(zip(expected_tables, actual_tables)):
    old = [txt(c) for c in a.xpath('.//w:tc')]
    new = [txt(c) for c in b.xpath('.//w:tc')]
    assert len(old) == len(new)
    if old != new:
        table_diffs.append({'table_index': i, 'changes': [{'cell': j, 'before': x, 'after': y} for j, (x,y) in enumerate(zip(old,new)) if x != y]})
assert [x['table_index'] for x in table_diffs] == [56, 57, 66]
assert [x['cell'] for x in table_diffs[0]['changes']] == [6,7,8]
assert table_diffs[2]['changes'] == [{'cell':13,'before':'纳入表1-2、表4-4和14项Holm主要族','after':'纳入表1-2、表4-6和14项Holm主要族'}]

expected_math = Counter(x for e in expected_blocks for x in maths(e))
actual_math = Counter(maths(doc.element.body))
removed_math = list((expected_math - actual_math).elements())
added_math = list((actual_math - expected_math).elements())
# The only mathematical-content edits are removal of a repeated inline JS
# definition and the canonical component-identity correction in Appendix F.
assert added_math == ['c=d,b,e,t'], added_math
assert 'c=s,e,t' in removed_math
assert Counter(removed_math) == Counter(['JSp,q=12KLp∥p+12KLq∥p', 'p=p+q/2', 'c=s,e,t']), removed_math

bookmarks = doc.element.xpath('.//w:bookmarkStart/@w:name')
assert len(bookmarks) == len(set(bookmarks))
fields = doc.element.xpath('.//w:fldSimple/@w:instr')
refs = [re.search(r'REF\s+(\S+)', x).group(1) for x in fields if re.search(r'\bREF\s+',x)]
assert all(x in bookmarks for x in refs)
seqs = [x for x in fields if re.search(r'\bSEQ\s+Equation',x)]
assert len(seqs) == 42
ids = doc.element.xpath('.//wp:docPr/@id')
assert len(ids) == len(set(ids))
figures = [txt(e) for e in doc.element.body if e.tag == qn('w:p') and re.match(r'^图\d-\d+\s', txt(e))]
assert len(figures) == len(set(figures))

pdfs = list((O/'_render_verified').glob('*.pdf'))
assert len(pdfs) == 1
pdf = PdfReader(pdfs[0])
pages = [{'page': i+1, 'text': p.extract_text() or ''} for i,p in enumerate(pdf.pages)]
assert len(pages) == 91
assert all(len(p['text'].strip()) > 20 for p in pages)
assert not any(re.search(r'Error!|错误!|未定义书签|找不到引用源', p['text']) for p in pages)
screen_changes = [n for n in range(1,92) if sha(O/f'_render_verified/page-{n}.png') != sha(O/f'_render_screen/page-{n}.png')]
(O/'_work/页面文本_verified.json').write_text(json.dumps(pages, ensure_ascii=False, indent=2))
(O/'_work/全文核对文本_verified.txt').write_text('\n\n'.join(f"=== 第{p['page']}页 ===\n{p['text']}" for p in pages))
report = {
    'output': str(out), 'sha256': sha(out), 'zip_integrity': True,
    'internal_relationships_resolved': rel_count,
    'source_hashes_unchanged': True, 'embedded_svg_byte_exact': 3,
    'display_equations': formula_checks, 'native_math_objects': len(doc.element.xpath('.//m:oMath')),
    'inline_math_changes': {'removed': removed_math, 'added': added_math},
    'tables_total': 73, 'tables_with_all_cell_text_preserved': 70,
    'table_changes': table_diffs, 'equation_seq_fields': len(seqs),
    'equation_ref_fields': len(refs), 'equation_refs_resolved': True,
    'active_drawings': len(ids), 'figure_captions': figures,
    'rendered_pages': len(pages), 'empty_pages': 0, 'field_errors': 0,
    'final_render_changed_pages': screen_changes,
}
(O/'_work/最终一致性核对.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False, indent=2))
