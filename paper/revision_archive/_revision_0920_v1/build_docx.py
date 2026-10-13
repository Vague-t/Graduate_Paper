"""Build the 0920 thesis in the retained 0915 DOCX layout, with native OMML.

Only running main without --check-only writes a DOCX. Importing this module is
side-effect free. The parent authoring task owns the artifact-operation marker,
render/inspect loop, final TOC page map, and any review comments.

Run with the bundled interpreter, never system Python:
  /Users/channtang/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3

Rendering is deliberately NOT performed here. Any subsequent render must use
the packaged render_docx.py and this bundled LibreOffice, never desktop LO:
  /Users/channtang/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice
"""

from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
WORK = Path(__file__).resolve().parent
BUNDLED_PYTHON = Path(
    '/Users/channtang/.cache/codex-runtimes/codex-primary-runtime/'
    'dependencies/python/bin/python3'
)
SOURCE = ROOT / 'report_0920_v1.md'
TEMPLATE = ROOT / 'report_0915_v1.docx'
OUTPUT = ROOT / 'report_0920_v1.docx'
CN, LAT, MATH = 'Songti SC', 'Times New Roman', 'STIX Two Math'
TITLE = '面向关系感知学习与依赖闭合修正的端到端论辩挖掘方法研究'
ENGLISH_TITLE = (
    'End-to-End Argument Mining via Relation-Aware Learning '
    'and Dependency-Closed Repair'
)
FENCE = '`' * 3
COMMENT_RE = re.compile(r'<!--.*?-->', re.S)


def el(tag, value=None):
    node = OxmlElement(tag)
    if value is not None:
        node.set(qn('m:val'), str(value))
    return node


def box(tag, children):
    node = el(tag)
    for child in children:
        node.append(child)
    return node


def mr(text, plain=False, bold=False):
    node = el('m:r')
    if plain or bold:
        prop = el('m:rPr')
        prop.append(el('m:sty', 'b' if bold else 'p'))
        node.append(prop)
    text_node = el('m:t')
    text_node.text = text
    text_node.set(qn('xml:space'), 'preserve')
    node.append(text_node)
    return node


# All known vocabulary is explicit. An unknown command always fails the build.
SYMBOLS = {
    'alpha': 'α', 'beta': 'β', 'gamma': 'γ', 'delta': 'δ', 'epsilon': 'ϵ',
    'varepsilon': 'ε', 'eta': 'η', 'theta': 'θ', 'iota': 'ι', 'kappa': 'κ',
    'lambda': 'λ', 'mu': 'μ', 'nu': 'ν', 'xi': 'ξ', 'pi': 'π', 'rho': 'ρ',
    'sigma': 'σ', 'tau': 'τ', 'phi': 'φ', 'varphi': 'ϕ', 'chi': 'χ',
    'psi': 'ψ', 'omega': 'ω', 'Gamma': 'Γ', 'Delta': 'Δ', 'Theta': 'Θ',
    'Lambda': 'Λ', 'Xi': 'Ξ', 'Pi': 'Π', 'Sigma': 'Σ', 'Phi': 'Φ',
    'Psi': 'Ψ', 'Omega': 'Ω', 'ell': 'ℓ', 'infty': '∞',
    'in': '∈', 'notin': '∉', 'le': '≤', 'leq': '≤', 'ge': '≥', 'geq': '≥',
    'ne': '≠', 'neq': '≠', 'equiv': '≡', 'approx': '≈', 'sim': '∼',
    'times': '×', 'cdot': '·', 'pm': '±', 'mid': '∣', 'vert': '|',
    'Vert': '‖', '|': '‖', 'cup': '∪', 'cap': '∩', 'bigcup': '⋃',
    'bigcap': '⋂', 'setminus': '∖', 'subset': '⊂', 'subseteq': '⊆',
    'supseteq': '⊇', 'succ': '≻', 'preceq': '≼', 'forall': '∀',
    'exists': '∃', 'lor': '∨', 'land': '∧', 'neg': '¬',
    'rightarrow': '→', 'to': '→', 'longrightarrow': '⟶',
    'leftarrow': '←', 'leftrightarrow': '↔', 'Rightarrow': '⇒',
    'Leftrightarrow': '⇔', 'mapsto': '↦', 'longmapsto': '⟼',
    'ldots': '…', 'dots': '…', 'cdots': '⋯', 'vdots': '⋮',
    'sum': '∑', 'prod': '∏', 'triangle': '△', 'triangleq': '≜',
    'varnothing': '∅', 'emptyset': '∅', 'partial': '∂', 'nabla': '∇',
}
CALLIGRAPHIC = dict(zip(
    'ABCDEFGHIJKLMNOPQRSTUVWXYZ',
    '𝒜ℬ𝒞𝒟ℰℱ𝒢ℋℐ𝒥𝒦ℒℳ𝒩𝒪𝒫𝒬ℛ𝒮𝒯𝒰𝒱𝒲𝒳𝒴𝒵',
))
DOUBLE = dict(zip(
    'ABCDEFGHIJKLMNOPQRSTUVWXYZ',
    '𝔸𝔹ℂ𝔻𝔼𝔽𝔾ℍ𝕀𝕁𝕂𝕃𝕄ℕ𝕆ℙℚℝ𝕊𝕋𝕌𝕍𝕎𝕏𝕐ℤ',
))
DOUBLE.update(dict(zip('0123456789', '𝟘𝟙𝟚𝟛𝟜𝟝𝟞𝟟𝟠𝟡')))
DELIMITERS = {
    r'\{': '{', r'\}': '}', r'\lbrace': '{', r'\rbrace': '}',
    r'\langle': '⟨', r'\rangle': '⟩', r'\lfloor': '⌊',
    r'\rfloor': '⌋', r'\lceil': '⌈', r'\rceil': '⌉',
    r'\vert': '|', r'\Vert': '‖', r'\|': '‖',
    '(': '(', ')': ')', '[': '[', ']': ']', '|': '|', '.': '',
}


def split_tex_level(text, separator):
    """Split row/column separators outside braces and nested environments."""
    tokens = re.findall(r'\\[A-Za-z]+|\\.|[{}_^]|.', text, re.S)
    result, current = [], []
    braces = environments = 0
    for token in tokens:
        if token == separator and braces == 0 and environments == 0:
            result.append(''.join(current))
            current = []
            continue
        current.append(token)
        if token == '{':
            braces += 1
        elif token == '}':
            braces -= 1
            if braces < 0:
                raise ValueError('Unmatched closing brace while splitting math')
        elif token == r'\begin':
            environments += 1
        elif token == r'\end':
            environments -= 1
    if braces or environments:
        raise ValueError('Unbalanced braces/environments in math rows')
    result.append(''.join(current))
    return result


def matrix(rows, alignments):
    """Native OMML matrix, used for cases and stacked subscripts."""
    if not rows:
        raise ValueError('Empty math matrix')
    cols = max(len(row) for row in rows)
    node = el('m:m')
    prop = el('m:mPr')
    prop.append(el('m:baseJc', 'center'))
    col_specs = el('m:mcs')
    for index in range(cols):
        col = el('m:mc')
        col_prop = el('m:mcPr')
        col_prop.append(el('m:count', 1))
        col_prop.append(el('m:mcJc', alignments[index % len(alignments)]))
        col.append(col_prop)
        col_specs.append(col)
    prop.append(col_specs)
    node.append(prop)
    for row in rows:
        matrix_row = el('m:mr')
        for item in row + [''] * (cols - len(row)):
            matrix_row.append(box('m:e', Tex(item).convert()))
        node.append(matrix_row)
    return node


def equation_array(rows):
    """Render aligned rows without the empty matrix cells LO cannot import.

    A single TeX & becomes an alignment property on the next nonempty run;
    a double && separates paired statements with an explicit quad. Each row
    is one populated m:e, so leading/trailing alignment markers cannot create
    empty cells or discard a continuation line.
    """
    node = el('m:eqArr')
    prop = el('m:eqArrPr')
    prop.append(el('m:baseJc', 'center'))
    node.append(prop)
    for row in rows:
        children = []
        previous = None
        for index, part in enumerate(row):
            if not part.strip():
                continue
            parsed = Tex(part).convert()
            if not parsed:
                continue
            if previous is not None and index - previous >= 2:
                children.append(mr(' ', plain=True))
            elif index > 0:
                first_run = next(
                    (run for item in parsed for run in item.iter(qn('m:r'))), None
                )
                if first_run is not None:
                    run_prop = first_run.find(qn('m:rPr'))
                    if run_prop is None:
                        run_prop = el('m:rPr')
                        first_run.insert(0, run_prop)
                    if run_prop.find(qn('m:aln')) is None:
                        run_prop.insert(0, el('m:aln'))
            children.extend(parsed)
            previous = index
        if not children:
            raise ValueError('Empty aligned equation row')
        node.append(box('m:e', children))
    return node


def delimiter(begin, end, children):
    node = el('m:d')
    prop = el('m:dPr')
    prop.append(el('m:begChr', begin))
    prop.append(el('m:endChr', end))
    node.append(prop)
    node.append(box('m:e', children))
    return node


class Tex:
    """Strict, deliberately bounded TeX-to-OMML parser for this manuscript."""

    def __init__(self, source):
        self.source = source
        self.tokens = re.findall(r'\\[A-Za-z]+|\\.|[{}_^]|.', source, re.S)
        self.index = 0

    def peek(self):
        return self.tokens[self.index] if self.index < len(self.tokens) else ''

    def pop(self):
        if self.index >= len(self.tokens):
            raise ValueError('Unexpected end of formula')
        token = self.tokens[self.index]
        self.index += 1
        return token

    def skip(self):
        while self.peek() and self.peek().isspace():
            self.pop()

    def convert(self):
        result = self.seq()
        if self.peek():
            raise ValueError('Unconsumed formula tokens: ' + self.peek())
        return result

    def group(self):
        self.skip()
        if self.peek() == '{':
            self.pop()
            result = self.seq('}')
            if self.pop() != '}':
                raise ValueError('Unclosed group')
            return result
        return self.atom()

    def literal_group(self):
        self.skip()
        if self.peek() != '{':
            return self.pop()
        self.pop()
        depth, result = 1, []
        while depth:
            token = self.pop()
            if token == '{':
                depth += 1
            elif token == '}':
                depth -= 1
            if depth:
                result.append(token)
        return ''.join(result)

    def seq(self, end=None):
        result = []
        while self.peek():
            self.skip()
            if not self.peek() or self.peek() == end:
                break
            base = self.atom()
            sub = sup = None
            self.skip()
            while self.peek() in ('_', '^'):
                kind = self.pop()
                argument = self.group()
                if kind == '_':
                    if sub is not None:
                        raise ValueError('Repeated subscript')
                    sub = argument
                else:
                    if sup is not None:
                        raise ValueError('Repeated superscript')
                    sup = argument
                self.skip()
            if sub is not None or sup is not None:
                if not base:
                    raise ValueError('Script has no base')
                tag = ('m:sSubSup' if sub is not None and sup is not None
                       else 'm:sSub' if sub is not None else 'm:sSup')
                node = el(tag)
                node.append(box('m:e', base))
                if sub is not None:
                    node.append(box('m:sub', sub))
                if sup is not None:
                    node.append(box('m:sup', sup))
                base = [node]
            result.extend(base)
        return result

    def read_delimiter(self):
        self.skip()
        token = self.pop()
        if token not in DELIMITERS:
            raise ValueError('Unsupported delimiter ' + token)
        return DELIMITERS[token]

    def environment(self, name):
        if name not in ('cases', 'aligned', 'gathered', 'matrix', 'pmatrix', 'bmatrix'):
            raise ValueError('Unsupported environment ' + name)
        raw, nesting = [], 0
        while True:
            token = self.pop()
            if token == r'\begin':
                nesting += 1
            elif token == r'\end':
                if not nesting:
                    closing = self.literal_group()
                    if closing != name:
                        raise ValueError(f'Environment mismatch: {name} / {closing}')
                    break
                nesting -= 1
            raw.append(token)
        rows = [split_tex_level(row, '&')
                for row in split_tex_level(''.join(raw), r'\\') if row.strip()]
        if name == 'aligned':
            return [equation_array(rows)]
        aligns = ['left']
        if name in ('gathered', 'matrix', 'pmatrix', 'bmatrix'):
            aligns = ['center']
        node = matrix(rows, aligns)
        if name == 'cases':
            return [delimiter('{', '', [node])]
        if name == 'pmatrix':
            return [delimiter('(', ')', [node])]
        if name == 'bmatrix':
            return [delimiter('[', ']', [node])]
        return [node]

    def atom(self):
        token = self.pop()
        if token == '{':
            result = self.seq('}')
            if self.pop() != '}':
                raise ValueError('Unclosed group')
            return result
        if token in ('}', '_', '^', '&'):
            raise ValueError('Unexpected math token ' + token)
        if not token.startswith('\\'):
            return [mr('′' if token == "'" else token)]
        cmd = token[1:]
        if cmd in ('frac', 'dfrac', 'tfrac'):
            node = el('m:f')
            node.append(box('m:num', self.group()))
            node.append(box('m:den', self.group()))
            return [node]
        if cmd in ('operatorname', 'text', 'mathrm', 'mathbf', 'mathit'):
            text = self.literal_group()
            text = re.sub(r'\\([_%{}&$#])', r'\1', text)
            if '\\' in text:
                raise ValueError('Nested command in text/font group: ' + text)
            return [mr(text, plain=cmd != 'mathit', bold=cmd == 'mathbf')]
        if cmd in ('mathcal', 'mathbb'):
            text = self.literal_group()
            mapper = CALLIGRAPHIC if cmd == 'mathcal' else DOUBLE
            if any(char not in mapper for char in text):
                raise ValueError('Unsupported alphabet in ' + token + ': ' + text)
            return [mr(''.join(mapper[char] for char in text))]
        if cmd in ('hat', 'widehat', 'bar', 'overline', 'tilde', 'widetilde', 'vec'):
            node = el('m:acc')
            prop = el('m:accPr')
            accent = {'hat': '̂', 'widehat': '̂', 'bar': '̅', 'overline': '̅',
                      'tilde': '̃', 'widetilde': '̃', 'vec': '⃗'}[cmd]
            prop.append(el('m:chr', accent))
            node.append(prop)
            node.append(box('m:e', self.group()))
            return [node]
        if cmd == 'sqrt':
            self.skip()
            degree = []
            if self.peek() == '[':
                self.pop()
                degree = self.seq(']')
                if self.pop() != ']':
                    raise ValueError('Unclosed root degree')
            node = el('m:rad')
            prop = el('m:radPr')
            prop.append(el('m:degHide', 0 if degree else 1))
            node.append(prop)
            node.append(box('m:deg', degree))
            node.append(box('m:e', self.group()))
            return [node]
        if cmd == 'left':
            begin = self.read_delimiter()
            children = self.seq(r'\right')
            if self.pop() != r'\right':
                raise ValueError('Unclosed left/right delimiter')
            return [delimiter(begin, self.read_delimiter(), children)]
        if cmd in ('big', 'Big', 'bigg', 'Bigg', 'bigl', 'bigr', 'Bigl', 'Bigr'):
            return [mr(self.read_delimiter())]
        if cmd == 'begin':
            return self.environment(self.literal_group())
        if cmd == 'substack':
            raw = self.literal_group()
            return [matrix([[row] for row in split_tex_level(raw, r'\\')], ['center'])]
        if cmd in ('displaystyle', 'textstyle', 'scriptstyle', 'scriptscriptstyle'):
            return []
        if cmd in ('quad', 'qquad', ',', ';', ':', ' ', '!'):
            # Negative thin space has no semantic token in OMML; omit explicitly.
            if cmd == '!':
                return []
            spaces = {'quad': ' ', 'qquad': ' ', ',': ' ', ';': ' ', ':': ' ', ' ': ' '}
            return [mr(spaces[cmd], True)]
        if cmd in ('{', '}', '_', '%', '&', '$', '#'):
            return [mr(cmd, True)]
        if cmd in ('min', 'max', 'arg', 'log', 'exp', 'sin', 'cos', 'lim', 'Pr', 'det'):
            return [mr(cmd, True)]
        if cmd in SYMBOLS:
            return [mr(SYMBOLS[cmd])]
        raise ValueError('Unsupported LaTeX command ' + token)


def split_table_row(line):
    """Keep pipes inside math/code and escaped literal pipes in one cell."""
    text = line.strip()
    if not text.startswith('|') or not text.endswith('|'):
        raise ValueError('Table rows require outside | delimiters: ' + line)
    text = text[1:-1]
    cells, current = [], []
    in_math, code_ticks = False, 0
    index = 0
    while index < len(text):
        char = text[index]
        if char == '\\' and index + 1 < len(text):
            if text[index + 1] == '|' and not in_math:
                current.append('|')
            else:
                current.extend((char, text[index + 1]))
            index += 2
            continue
        if char == '`':
            end = index
            while end < len(text) and text[end] == '`':
                end += 1
            count = end - index
            if not code_ticks:
                code_ticks = count
            elif code_ticks == count:
                code_ticks = 0
            current.append(text[index:end])
            index = end
            continue
        if char == '$' and not code_ticks:
            in_math = not in_math
        if char == '|' and not in_math and not code_ticks:
            cells.append(''.join(current).strip())
            current = []
        else:
            current.append(char)
        index += 1
    if in_math or code_ticks:
        raise ValueError('Unclosed math/code delimiter in table: ' + line)
    cells.append(''.join(current).strip())
    return cells


def blocks(lines):
    """Markdown block iterator retaining source line numbers for strict errors."""
    index = 0
    while index < len(lines):
        number = index + 1
        line = lines[index].strip()
        index += 1
        if not line or line in ('---', '***'):
            continue
        if line.startswith(FENCE):
            kind, code = line[3:].strip(), []
            while index < len(lines) and not lines[index].strip().startswith(FENCE):
                code.append(lines[index])
                index += 1
            if index == len(lines):
                raise ValueError(f'Line {number}: unclosed code fence')
            index += 1
            yield 'code', (kind, code), number
        elif line.startswith('$$'):
            formula = line[2:]
            while not formula.rstrip().endswith('$$'):
                if index == len(lines):
                    raise ValueError(f'Line {number}: unclosed display equation')
                formula += '\n' + lines[index]
                index += 1
            yield 'math', formula.rstrip()[:-2], number
        elif line.startswith('|'):
            rows = [split_table_row(line)]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(split_table_row(lines[index]))
                index += 1
            rows = [row for row in rows
                    if not all(re.fullmatch(r':?-{3,}:?', cell) for cell in row)]
            if not rows or len({len(row) for row in rows}) != 1:
                raise ValueError(f'Line {number}: inconsistent table column counts')
            yield 'table', rows, number
        else:
            heading = re.fullmatch(r'(#{2,4}) (.+)', line)
            if heading:
                yield 'heading', (len(heading[1]) - 1, heading[2]), number
            elif line.startswith('# '):
                raise ValueError(f'Line {number}: unexpected manuscript H1 in body')
            else:
                yield 'text', line, number


INLINE = re.compile(r'(\$[^$]+\$|\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))')


def clean_inline(text):
    return COMMENT_RE.sub('', text).replace('`', '').replace('<br>', '\n').replace('<br/>', '\n')


def inline_formulas(text):
    clean = clean_inline(text)
    # Bold segments are recursively interpreted by inline(), so inspect every $.
    matches = list(re.finditer(r'\$([^$]+)\$', clean))
    residual = re.sub(r'\$[^$]+\$', '', clean)
    if '$' in residual:
        raise ValueError('Unbalanced inline math delimiters: ' + text)
    return [match[1] for match in matches]


def load_manuscript(path):
    text = COMMENT_RE.sub('', path.read_text(encoding='utf-8'))
    lines = text.splitlines()
    try:
        start = next(index for index, line in enumerate(lines) if line.strip() == '## 摘要')
    except StopIteration as exc:
        raise ValueError('Manuscript must contain an exact ## 摘要 heading') from exc
    return lines[start:]


def preflight(lines):
    counts = {'display_equations': 0, 'inline_equations': 0, 'tables': 0}
    for kind, value, number in blocks(lines):
        try:
            if kind == 'math':
                Tex(value).convert()
                counts['display_equations'] += 1
            elif kind != 'code':
                if kind == 'table':
                    counts['tables'] += 1
                    texts = [cell for row in value for cell in row]
                else:
                    texts = [value[1] if kind == 'heading' else value]
                for text in texts:
                    for formula in inline_formulas(text):
                        Tex(formula).convert()
                        counts['inline_equations'] += 1
        except Exception as exc:
            raise ValueError(f'Body line {number}: {exc}') from exc
    return counts


def font_run(run, size=None, bold=None, cn=CN, latin=LAT):
    run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), cn)
    return run


class Builder:
    def __init__(self, template, lines, toc_pages):
        self.lines, self.toc_pages = lines, toc_pages
        self.math_count = 0
        self.doc = Document(template)
        initial = deepcopy(self.doc.sections[0]._sectPr)
        # Retain the established package/styles/settings; replace requested body.
        for child in list(self.doc._element.body):
            self.doc._element.body.remove(child)
        self.doc._element.body.append(initial)
        for rel in list(self.doc.part.rels.values()):
            if rel.reltype.endswith('/comments'):
                self.doc.part.drop_rel(rel.rId)
        # The retained cover/front section has no running page furniture.
        section = self.doc.sections[0]
        for header_footer in (section.header, section.footer):
            for child in list(header_footer._element):
                header_footer._element.remove(child)
            header_footer.add_paragraph()
        self.text_width = section.page_width - section.left_margin - section.right_margin
        # Assert the expected source format; do not silently impose another layout.
        if abs(section.page_width.mm - 210) > .2 or abs(section.page_height.mm - 297) > .2:
            raise ValueError('Retained template is not the expected A4 layout')
        settings = self.doc.settings.element
        math_prop = settings.find(qn('m:mathPr'))
        if math_prop is None:
            math_prop = el('m:mathPr')
            settings.append(math_prop)
        math_font = math_prop.find(qn('m:mathFont'))
        if math_font is None:
            math_font = el('m:mathFont', MATH)
            math_prop.append(math_font)
        else:
            math_font.set(qn('m:val'), MATH)

    def math(self, paragraph, text):
        node = el('m:oMath')
        try:
            for item in Tex(text).convert():
                node.append(item)
        except Exception as exc:
            raise ValueError(f'Invalid formula {text!r}: {exc}') from exc
        paragraph._p.append(node)
        self.math_count += 1

    def inline(self, paragraph, text, size=None):
        text = clean_inline(text)
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith('$') and part.endswith('$'):
                self.math(paragraph, part[1:-1])
            elif part.startswith('**') and part.endswith('**'):
                begin = len(paragraph.runs)
                self.inline(paragraph, part[2:-2], size)
                for run in paragraph.runs[begin:]:
                    run.bold = True
            else:
                match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
                if match:
                    link = OxmlElement('w:hyperlink')
                    rel = paragraph.part.relate_to(
                        match[2],
                        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                        is_external=True,
                    )
                    link.set(qn('r:id'), rel)
                    run = OxmlElement('w:r')
                    prop = OxmlElement('w:rPr')
                    color = OxmlElement('w:color')
                    color.set(qn('w:val'), '000000')
                    prop.append(color)
                    run.append(prop)
                    value = OxmlElement('w:t')
                    value.text = match[1]
                    run.append(value)
                    link.append(run)
                    paragraph._p.append(link)
                else:
                    font_run(paragraph.add_run(part), size)

    def cover(self):
        doc = self.doc
        paragraph = doc.add_paragraph('武汉大学', style='Subtitle')
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(80)
        for run in paragraph.runs:
            font_run(run, 22, True)
        paragraph = doc.add_paragraph('学术型硕士学位论文', style='Subtitle')
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            font_run(run, 18)
        paragraph = doc.add_paragraph(
            '面向关系感知学习与依赖闭合修正的\n端到端论辩挖掘方法研究', style='Title'
        )
        paragraph.paragraph_format.space_before = Pt(45)
        paragraph.paragraph_format.space_after = Pt(30)
        paragraph = doc.add_paragraph(
            'End-to-End Argument Mining via\n'
            'Relation-Aware Learning and\nDependency-Closed Repair'
        )
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        for run in paragraph.runs:
            font_run(run, 15)
        for text in ('Chen Tang', '计算机科学与技术', '0920-v1修订稿', '2026年9月20日'):
            paragraph = doc.add_paragraph(text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(15)

    def toc(self):
        heads = [value[1] for kind, value, _ in blocks(self.lines)
                 if kind == 'heading' and value[0] == 1
                 and (re.match(r'^第.+章\s', value[1])
                      or re.match(r'^附录[A-E](?:\s|$)', value[1])
                      or value[1] == '参考文献')]
        self.doc.add_heading('目录', 1)
        for name in heads:
            paragraph = self.doc.add_paragraph()
            fmt = paragraph.paragraph_format
            fmt.first_line_indent = Pt(0)
            fmt.space_after = Pt(14)
            fmt.line_spacing = 1.4
            fmt.tab_stops.add_tab_stop(Mm(150), WD_TAB_ALIGNMENT.RIGHT)
            self.inline(paragraph, name)
            if name in self.toc_pages:
                font_run(paragraph.add_run('\t' + str(self.toc_pages[name])), 12)

    def body_section(self):
        section = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        for old in list(section._sectPr.findall(qn('w:pgNumType'))):
            section._sectPr.remove(old)
        page = OxmlElement('w:pgNumType')
        page.set(qn('w:start'), '1')
        section._sectPr.append(page)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        for item in (section.header, section.footer):
            for child in list(item._element):
                item._element.remove(child)
        paragraph = section.header.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        font_run(paragraph.add_run('武汉大学硕士学位论文　0920-v1修订稿'), 9)
        paragraph = section.footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        paragraph._p.append(field)

    def table(self, rows):
        count = len(rows[0])
        defaults = {2: [100, 57], 3: [49, 50, 58],
                    4: [64, 31, 31, 31], 5: [54, 25, 24, 24, 30]}
        widths = defaults.get(count, [157 / count] * count)
        first = rows[0][0]
        specific = {
            ('编号', 3): [15, 53, 89], ('研究问题', 4): [33, 40, 38, 46],
            ('语料', 4): [18, 52, 30, 57], ('异常对象', 3): [31, 47, 79],
            ('研究层次', 4): [26, 43, 47, 41], ('证据类别', 4): [25, 43, 44, 45],
            ('对照', 4): [40, 38, 39, 40], ('研究判断', 3): [34, 62, 61],
            ('机制', 3): [36, 49, 72],
        }
        widths = specific.get((first, count), widths)
        scale = self.text_width / Mm(sum(widths))
        widths = [width * scale for width in widths]
        compact = first == '编号'
        table = self.doc.add_table(rows=0, cols=count)
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Mm(width)
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + side)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:color'), 'D9D9D9')
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for index, row in enumerate(rows):
            cells = table.add_row().cells
            prop = table.rows[-1]._tr.get_or_add_trPr()
            prop.append(OxmlElement('w:cantSplit'))
            if index == 0:
                prop.append(OxmlElement('w:tblHeader'))
            for col_index, (cell, text, width) in enumerate(zip(cells, row, widths)):
                cell.width = Mm(width)
                cell_prop = cell._tc.get_or_add_tcPr()
                align = OxmlElement('w:vAlign')
                align.set(qn('w:val'), 'center')
                cell_prop.append(align)
                margins = OxmlElement('w:tcMar')
                for side, value in (('top', 70), ('bottom', 70), ('left', 80), ('right', 80)):
                    margin = OxmlElement('w:' + side)
                    margin.set(qn('w:w'), str(value))
                    margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                cell_prop.append(margins)
                if index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E7E7E7')
                    cell_prop.append(shade)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(15) if compact else 1.15
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 or len(text) > 28
                                       else WD_ALIGN_PARAGRAPH.CENTER)
                self.inline(paragraph, text, 10 if compact else 10.5)
                if index == 0:
                    for run in paragraph.runs:
                        run.bold = True
        # Inherit the earlier short-table / final-two-rows pagination policy.
        keep_from = 0 if len(rows) <= 6 else max(0, len(rows) - 2)
        for index, row in enumerate(table.rows):
            if keep_from <= index < len(rows) - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1

    def body(self):
        chapter, eq_counter = '0', 0
        appendix = english = body_started = False
        chinese_numbers = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6,
                           '七': 7, '八': 8, '九': 9, '十': 10}
        for kind, value, line in blocks(self.lines):
            try:
                if kind == 'heading':
                    level, title = value
                    if title.startswith('第一章') and not body_started:
                        self.toc()
                        self.body_section()
                        body_started = True
                    english = title == 'Abstract' if level == 1 else english
                    if level == 1:
                        match = re.match(r'^第([一二三四五六七八九十]+)章\s', title)
                        app = re.match(r'^附录([A-E])(?:\s|$)', title)
                        if match:
                            chapter = str(chinese_numbers.get(match[1], match[1]))
                            eq_counter = 0
                        elif app:
                            chapter, eq_counter, appendix = app[1], 0, True
                    paragraph = self.doc.add_heading(clean_inline(title), level)
                    if title.startswith('第一章'):
                        paragraph.paragraph_format.page_break_before = False
                    continue
                if kind == 'table':
                    self.table(value)
                    continue
                paragraph = self.doc.add_paragraph()
                fmt = paragraph.paragraph_format
                if appendix:
                    fmt.line_spacing = Pt(19)
                if english:
                    fmt.line_spacing = Pt(16.5)
                if kind == 'code':
                    code_type, code = value
                    fmt.first_line_indent = Pt(0)
                    fmt.left_indent = Mm(3)
                    fmt.line_spacing = 1.1
                    fmt.space_before, fmt.space_after = Pt(5), Pt(9)
                    # Long algorithms may flow naturally; short source boxes stay together.
                    fmt.keep_together = len(code) <= 18
                    if code_type == 'flow':
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    font_run(paragraph.add_run('\n'.join(code)),
                             11 if code_type == 'flow' else 9.5,
                             latin=LAT if code_type == 'flow' else 'Courier New')
                elif kind == 'math':
                    fmt.first_line_indent = Pt(0)
                    fmt.line_spacing = 1.2
                    fmt.space_before, fmt.space_after = Pt(5), Pt(9)
                    fmt.keep_together = True
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.math(paragraph, value)
                    eq_counter += 1
                    font_run(paragraph.add_run(f'　（{chapter}-{eq_counter}）'), 10)
                else:
                    text = value
                    label = clean_inline(text).replace('**', '')
                    if re.match(r'^(表|图|算法)[A-E0-9]+-\d+(?:\s|$)', label):
                        fmt.keep_with_next = True
                        fmt.first_line_indent = Pt(0)
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        self.inline(paragraph, text, 11)
                    elif re.match(r'^\[\d+\] ', text):
                        fmt.first_line_indent, fmt.left_indent = Mm(-8), Mm(8)
                        fmt.line_spacing, fmt.space_after = Pt(14), Pt(8)
                        fmt.keep_together = True
                        self.inline(paragraph, text, 10.5)
                    else:
                        if re.match(r'^(?:\d+\.|[-*] )', text) or text.startswith(('关键词', 'Keywords')):
                            fmt.first_line_indent = Pt(0)
                        size = 11 if appendix or text.startswith(('关键词', 'Keywords:')) else None
                        self.inline(paragraph, text, size)
            except Exception as exc:
                raise ValueError(f'Body line {line}: {exc}') from exc

    def save(self, output):
        self.doc.core_properties.title = TITLE
        self.doc.core_properties.subject = '0920-v1修订稿'
        self.doc.core_properties.author = 'Chen Tang'
        self.doc.core_properties.keywords = '论辩挖掘 关系感知学习 依赖闭合修正'
        # OOXML pPr must remain the first child even after late formatting calls.
        for paragraph in self.doc._element.xpath('.//w:p'):
            prop = paragraph.find(qn('w:pPr'))
            if prop is not None and paragraph.index(prop):
                paragraph.remove(prop)
                paragraph.insert(0, prop)
        self.doc.save(output)


def self_test():
    formulas = [
        r'\begin{aligned}a&=\frac1{|A|}\sum_{i\in A}q_i\\b&=\mathbb E[Y]\end{aligned}',
        r'\sum_{\substack{i\in A\\j\ne i}}q_{ij}',
        r'\left\Vert p-\bar q\right\Vert_\infty<\delta/2',
        r'\forall c\in C,\quad x_c\le1,\ y_c\ge0,\ C\ne\emptyset',
        r'f(x)=\begin{cases}1&x\ge0\\0&x<0\end{cases}',
    ]
    for formula in formulas:
        Tex(formula).convert()
    aligned = Tex(r'\begin{aligned}a={}&b\\&+c&&d=e\end{aligned}').convert()[0]
    assert aligned.tag == qn('m:eqArr')
    assert len(aligned.findall(qn('m:e'))) == 2
    assert not list(aligned.iter(qn('m:m')))
    assert ''.join(aligned.itertext()).replace(' ', '') == 'a=b+cd=e'
    assert split_table_row(r'|对象|$|R|$ 与 $p\|q$|`a|b`|') == [
        '对象', r'$|R|$ 与 $p\|q$', '`a|b`',
    ]
    try:
        Tex(r'\unsupported{x}').convert()
    except ValueError:
        pass
    else:
        raise AssertionError('Unknown TeX command was silently accepted')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE)
    parser.add_argument('--template', type=Path, default=TEMPLATE)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    parser.add_argument('--toc-map', type=Path, default=WORK / 'toc_pages.json')
    parser.add_argument('--check-only', action='store_true', help='Validate math and tables; write nothing')
    args = parser.parse_args()
    if Path(sys.executable).resolve() != BUNDLED_PYTHON.resolve():
        raise SystemExit('Use the bundled Python interpreter: ' + str(BUNDLED_PYTHON))
    self_test()
    lines = load_manuscript(args.source)
    counts = preflight(lines)
    if args.check_only:
        print(json.dumps({'status': 'OK', 'source': str(args.source), **counts}, ensure_ascii=False))
        return
    if args.output.resolve() == args.template.resolve():
        raise ValueError('The retained template must not be overwritten')
    template_hash = hashlib.sha256(args.template.read_bytes()).hexdigest()
    toc_pages = json.loads(args.toc_map.read_text(encoding='utf-8')) if args.toc_map.exists() else {}
    builder = Builder(args.template, lines, toc_pages)
    builder.cover()
    builder.body()
    builder.save(args.output)
    if hashlib.sha256(args.template.read_bytes()).hexdigest() != template_hash:
        raise RuntimeError('Retained template changed unexpectedly')
    report = {
        'output': str(args.output), 'template': str(args.template),
        'template_sha256': template_hash, 'native_equation_count': builder.math_count,
        'table_count': len(builder.doc.tables), 'paragraph_count': len(builder.doc.paragraphs),
        'cjk_font': CN, 'math_font': MATH, 'toc_page_map_loaded': bool(toc_pages),
        'rendered_or_visually_verified': False, **counts,
    }
    (WORK / 'docx_build.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
