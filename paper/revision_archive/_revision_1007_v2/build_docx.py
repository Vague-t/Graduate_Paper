"""Build the explicitly authorized v2 thesis, retaining native editable math.

The parent has run the sole artifact-operation marker. Only the new v2 DOCX
and task-local build/QA files are written. Render with packaged render_docx.py
using the bundled Python and bundled LibreOffice only:
/Users/channtang/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import sys

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_LEADER

ROOT = Path(__file__).resolve().parent.parent
WORK = Path(__file__).resolve().parent
SOURCE = ROOT / 'report_1007_v2_full.md'
TEMPLATE = ROOT / 'report_1007_v1_full.docx'
OUTPUT = ROOT / 'report_1007_v2_full.docx'
WARNING = '1007-v2 修订稿　含非实测占位　勿送审'
spec = importlib.util.spec_from_file_location('v1_layout', ROOT / '_revision_1007_v1/build_docx.py')
v1 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v1)
base = v1.base
IMAGE = v1.IMAGE
CAPTION = re.compile(r'^(表|图|算法)([A-F0-9]+-\d+P?)\s+(.+)$')
NUMERIC = re.compile(r'^[+−-]?\d+(?:\.\d+)?(?:P%?|%P?)?(?:\s*±\s*\d+(?:\.\d+)?(?:P%?|%P?)?)?$')


def is_numeric_cell(value):
    text = base.clean_inline(value).replace('**', '').strip()
    if NUMERIC.fullmatch(text):
        return True
    values = [part.strip() for part in text.split('→')]
    return len(values) == 2 and all(NUMERIC.fullmatch(part) for part in values)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def node(tag, **attrs):
    result = base.OxmlElement(tag)
    for key, value in attrs.items():
        result.set(base.qn('w:' + key), str(value))
    return result


def field_part(paragraph, kind, instruction=None):
    run = paragraph.add_run()
    if instruction is not None:
        item = node('w:instrText')
        item.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        item.text = ' ' + instruction + ' '
    else:
        item = node('w:fldChar', fldCharType=kind)
    run._r.append(item)


def bookmark(paragraph, name, number):
    paragraph._p.append(node('w:bookmarkStart', id=number, name=name))
    paragraph._p.append(node('w:bookmarkEnd', id=number))


class Builder(v1.Builder):
    def __init__(self, template, lines, page_map):
        # Retain source package/styles; explicitly replace authorized content.
        base.Builder.__init__(self, template, lines, page_map)
        self.doc.settings.odd_and_even_pages_header_footer = False
        self.page_map = page_map
        self.blocks = list(base.blocks(lines))
        self.headings = [(level, title) for kind, value, _ in self.blocks
                         if kind == 'heading' for level, title in [value]
                         if title not in ('目录', '图目录', '表目录')]
        self.captions = [base.clean_inline(value).replace('**', '')
                         for kind, value, _ in self.blocks if kind == 'text'
                         and CAPTION.fullmatch(base.clean_inline(value).replace('**', ''))]
        keys = [title for _, title in self.headings] + self.captions
        self.bookmarks = {key: 'v2_' + str(index + 10) for index, key in enumerate(keys)}
        self.attached_bookmarks = set()
        self.bookmark_counter = len(keys) + 10
        self.table_caption = ''
        self.table_records = []
        self.equation_records = []
        self.algorithm_records = []
        self.table_splits = json.loads((WORK / 'table_splits.json').read_text()) if (WORK / 'table_splits.json').exists() else {}
        self.apply_styles()
        self.furniture(self.doc.sections[0], numbered=False)

    def apply_styles(self):
        doc = self.doc
        styles = doc.styles
        normal = styles['Normal']
        normal.font.name, normal.font.size = base.LAT, base.Pt(12)
        normal.element.get_or_add_rPr().rFonts.set(base.qn('w:eastAsia'), base.CN)
        normal.paragraph_format.line_spacing = base.Pt(20.5)
        normal.paragraph_format.space_after = base.Pt(2)
        normal.paragraph_format.first_line_indent = base.Pt(24)
        normal.paragraph_format.widow_control = True
        for level, size in [(1, 18), (2, 14), (3, 12)]:
            style = styles[f'Heading {level}']
            style.font.size, style.font.bold = base.Pt(size), True
            style.font.color.rgb = base.RGBColor(0, 0, 0)
            style.paragraph_format.first_line_indent = base.Pt(0)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
            style.paragraph_format.line_spacing = 1.25
            style.paragraph_format.space_before = base.Pt(0 if level == 1 else 10)
            style.paragraph_format.space_after = base.Pt(12 if level == 1 else 5)
            style.paragraph_format.page_break_before = level == 1
            style.paragraph_format.alignment = base.WD_ALIGN_PARAGRAPH.CENTER if level == 1 else base.WD_ALIGN_PARAGRAPH.LEFT
        roles = [('Front Heading', 18, True), ('Front Contents', 18, True),
                 ('Figure Caption', 10.5, False), ('Table Caption', 10.5, False),
                 ('Continuation Caption', 10.5, False),
                 ('Algorithm Caption', 10.5, True), ('Algorithm Line', 10.5, False)]
        for name, size, bold in roles:
            style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
            style.font.name, style.font.size, style.font.bold = base.LAT, base.Pt(size), bold
            style.font.color.rgb = base.RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().rFonts.set(base.qn('w:eastAsia'), base.CN)
            fmt = style.paragraph_format
            fmt.first_line_indent = base.Pt(0)
            fmt.space_before = base.Pt(0)
            fmt.space_after = base.Pt(7)
            fmt.line_spacing = 1.15
            fmt.keep_together = True
            fmt.keep_with_next = name != 'Figure Caption'
            if name.startswith('Front'):
                fmt.page_break_before = True
                fmt.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
                fmt.space_after = base.Pt(18)
            elif name.endswith('Caption'):
                fmt.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        for level in (1, 2, 3):
            name = f'TOC {level}'
            style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
            style.font.size, style.font.bold = base.Pt(11), level == 1
            style.font.color.rgb = base.RGBColor(0, 0, 0)
            fmt = style.paragraph_format
            fmt.first_line_indent = base.Pt(0)
            fmt.left_indent = base.Mm((level - 1) * 5)
            fmt.right_indent = base.Mm(9)
            fmt.space_after = base.Pt(3 if level == 1 else 1)
            fmt.line_spacing = base.Pt(16)
            fmt.keep_with_next = False
            fmt.tab_stops.add_tab_stop(self.text_width, base.WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    def furniture(self, section, numbered=True, roman=False):
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        for container in (section.header, section.footer):
            for child in list(container._element):
                container._element.remove(child)
        p = section.header.add_paragraph()
        p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = base.Pt(0)
        p.paragraph_format.line_spacing = base.Pt(10)
        p.paragraph_format.space_after = base.Pt(2)
        base.font_run(p.add_run('武汉大学硕士学位论文'), 9)
        p = section.header.add_paragraph()
        p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = base.Pt(0)
        p.paragraph_format.line_spacing = base.Pt(9.5)
        p.paragraph_format.space_after = base.Pt(0)
        base.font_run(p.add_run(WARNING), 8, True)
        borders = node('w:pBdr')
        borders.append(node('w:bottom', val='single', sz='4', color='000000', space='3'))
        p._p.get_or_add_pPr().append(borders)
        for old in list(section._sectPr.findall(base.qn('w:pgNumType'))):
            section._sectPr.remove(old)
        if numbered:
            section._sectPr.append(node('w:pgNumType', start='1', fmt='upperRoman' if roman else 'decimal'))
        p = section.footer.add_paragraph()
        p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = base.Pt(0)
        if numbered:
            field = node('w:fldSimple', instr='PAGE')
            run = node('w:r')
            text = node('w:t')
            text.text = 'I' if roman else '1'
            run.append(text)
            field.append(run)
            p._p.append(field)

    def cover(self):
        base.Builder.cover(self)
        for p in self.doc.paragraphs:
            if p.text == '0920-v1修订稿':
                p.clear()
                base.font_run(p.add_run(WARNING), 12, True)
            elif p.text == '2026年9月20日':
                p.clear()
                base.font_run(p.add_run('2026年10月'))
            elif p.text.startswith('End-to-End Argument Mining via'):
                p._p.getparent().remove(p._p)
        p = self.doc.add_paragraph('Wuhan University', style='Subtitle')
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = base.Pt(75)
        p.paragraph_format.space_after = base.Pt(15)
        p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            base.font_run(run, 22, True)
        p = self.doc.add_paragraph("Academic Master's Thesis", style='Subtitle')
        p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            base.font_run(run, 18)
        p = self.doc.add_paragraph('End-to-End Argument Mining via\nRelation-Aware Learning and\nDependency-Closed Repair', style='Title')
        p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = base.Pt(55)
        p.paragraph_format.space_after = base.Pt(45)
        for run in p.runs:
            base.font_run(run, 20, True)
        for text in ('Chen Tang', 'Computer Science and Technology', 'October 2026'):
            p = self.doc.add_paragraph(text)
            p.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = base.Pt(0)
            p.paragraph_format.space_before = base.Pt(18)
            for run in p.runs:
                base.font_run(run, 14)
        section = self.doc.add_section(base.WD_SECTION_START.NEW_PAGE)
        self.furniture(section, roman=True)

    def body_section(self):
        section = self.doc.add_section(base.WD_SECTION_START.NEW_PAGE)
        self.furniture(section)

    def register(self, paragraph, key):
        if key not in self.bookmarks:
            self.bookmark_counter += 1
            name = 'v2_' + str(self.bookmark_counter)
            self.bookmarks[key] = name
        if key not in self.attached_bookmarks:
            name = self.bookmarks[key]
            bookmark(paragraph, name, int(name.split('_')[1]))
            self.attached_bookmarks.add(key)

    def contents_field(self, title, entries, instruction):
        self.doc.add_paragraph(title, style='Front Heading')
        for index, (level, text) in enumerate(entries):
            p = self.doc.add_paragraph(style=f'TOC {level}')
            if index == 0:
                field_part(p, 'begin')
                field_part(p, '', instruction)
                field_part(p, 'separate')
            # Cached text is deliberately complete. Field update remains native.
            begin = len(p._p)
            self.inline(p, text, 11)
            value = self.page_map.get(text, '1')
            base.font_run(p.add_run('\t' + str(value)), 11)
            if text in self.bookmarks and not p._p.xpath('./m:oMath'):
                link = node('w:hyperlink', anchor=self.bookmarks[text], history='1')
                for child in list(p._p)[begin:]:
                    link.append(child)
                p._p.append(link)
            if index == len(entries) - 1:
                field_part(p, 'end')

    def toc(self):
        self.contents_field('目录', self.headings, r'TOC \o "1-3" \h \z \t "Front Contents,1"')
        figures = [(1, t) for t in self.captions if t.startswith('图')]
        tables = [(1, t) for t in self.captions if t.startswith('表')]
        self.contents_field('图目录', figures, r'TOC \h \z \t "Figure Caption,1"')
        self.contents_field('表目录', tables, r'TOC \h \z \t "Table Caption,1"')

    def inline(self, paragraph, text, size=None):
        # Bind the two-character link label at these observed line endings;
        # the source bibliography, visible wording, and link target stay intact.
        if re.match(r'^\[(20|26)\] ', text):
            text = text.replace('[来源](', '[来\u2060源](')
        image = IMAGE.fullmatch(text.strip())
        if not image:
            return base.Builder.inline(self, paragraph, text, size)
        path = (ROOT / image[2]).resolve()
        assert path.is_relative_to(WORK / 'figures') and path.is_file(), path
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = base.Pt(0)
        fmt.line_spacing = 1
        fmt.space_before, fmt.space_after = base.Pt(4), base.Pt(3)
        fmt.keep_with_next = True
        paragraph.alignment = base.WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(path), width=self.text_width)
        shape._inline.docPr.set('descr', image[1] + '。机制示意，非实测案例。')
        shape._inline.docPr.set('name', path.name)

    def display_math(self, expression, label):
        # The equal side cells keep the native equation centered independently
        # of the right-aligned number. It is never rasterized or space-padded.
        table = self.doc.add_table(rows=1, cols=3)
        table.autofit = False
        widths = [16, self.text_width / base.Mm(1) - 32, 16]
        borders = node('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(node('w:' + side, val='nil'))
        table._tbl.tblPr.append(borders)
        for column, width in zip(table.columns, widths):
            column.width = base.Mm(width)
        table.rows[0]._tr.get_or_add_trPr().append(node('w:cantSplit'))
        for index, (cell, width) in enumerate(zip(table.rows[0].cells, widths)):
            cell.width = base.Mm(width)
            cell._tc.get_or_add_tcPr().append(node('w:vAlign', val='center'))
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = base.Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = base.Pt(4)
            p.paragraph_format.space_after = base.Pt(7)
            p.alignment = base.WD_ALIGN_PARAGRAPH.RIGHT if index == 2 else base.WD_ALIGN_PARAGRAPH.CENTER
            if index == 1:
                self.math(p, expression)
            elif index == 2:
                base.font_run(p.add_run('(' + label + ')'), 10)
        self.equation_records.append({'number': label, 'source': expression})

    def style_data_table(self, table, rows, numeric):
        borders = table._tbl.tblPr.find(base.qn('w:tblBorders'))
        for border in borders:
            side = border.tag.split('}')[-1]
            border.set(base.qn('w:val'), 'single' if not numeric or side in ('top', 'bottom') else 'nil')
            border.set(base.qn('w:color'), '000000' if numeric else 'D9D9D9')
            border.set(base.qn('w:sz'), '8' if numeric else '4')
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                prop = cell._tc.get_or_add_tcPr()
                if numeric:
                    for old in list(prop.findall(base.qn('w:shd'))):
                        prop.remove(old)
                    if index == 0:
                        cb = node('w:tcBorders')
                        cb.append(node('w:bottom', val='single', sz='5', color='000000'))
                        prop.append(cb)
                for p in cell.paragraphs:
                    # Keep at least the header plus three data rows together.
                    if index < min(3, len(rows) - 1):
                        p.paragraph_format.keep_with_next = True

    def table(self, rows):
        rows = deepcopy(rows)
        # Header line breaks are visual only; all values and labels are retained.
        header_breaks = {'AAEC有益排序AUROC': 'AAEC\n有益排序\nAUROC',
                         'CDCP有益排序AUROC': 'CDCP\n有益排序\nAUROC',
                         'AAEC关系Brier': 'AAEC\n关系\nBrier',
                         'CDCP超边Brier': 'CDCP\n超边\nBrier',
                         '自动与人工判定': '自动与人工\n判定',
                         '后续用途与影响': '后续用途\n与影响',
                         '预测效用增量': '预测效用\n增量',
                         '源文档与变体ID': '源文档与\n变体ID'}
        rows[0] = [header_breaks.get(cell, cell) for cell in rows[0]]
        numeric_cells = sum(is_numeric_cell(c) for row in rows[1:] for c in row[1:])
        numeric = numeric_cells >= max(2, (len(rows) - 1) * .75)
        key = self.table_caption.split(' ', 1)[0]
        breaks = self.table_splits.get(key, [])
        assert breaks == sorted(set(breaks)) and all(1 < row < len(rows) for row in breaks), (key, breaks)
        # Explicit split rows come from rendered QA, never alter table data.
        chunks = []
        start = 1
        for end in breaks + [len(rows)]:
            chunks.append([rows[0]] + rows[start:end])
            start = end
        for index, chunk in enumerate(chunks):
            if index:
                p = self.doc.add_paragraph(key + '（续）', style='Continuation Caption')
                p.paragraph_format.page_break_before = True
            base.Builder.table(self, chunk)
            table = self.doc.tables[-1]
            if index < len(chunks) - 1:
                # The base spacing paragraph can overflow by itself and create
                # a blank page before the continuation's explicit page break.
                spacer = self.doc.paragraphs[-1]
                assert not spacer.text and spacer._p.getprevious() is table._tbl
                spacer._p.getparent().remove(spacer._p)
            self.style_data_table(table, chunk, numeric)
            count, first = len(rows[0]), rows[0][0]
            widths = {('组件', 4): [15, 25, 86, 31], ('符号或术语', 3): [32, 56, 69],
                      ('方法', 3): [37, 49, 71], ('要回答的问题', 3): [56, 34, 67],
                      ('案例类别', 6): [24, 26, 30, 23, 27, 27],
                      ('研究', 4): [27, 39, 43, 48],
                      ('配置项目', 3): [31, 59, 67],
                      ('评分方式', 5): [49, 27, 27, 27, 27]}.get((first, count))
            if widths:
                for col, width in zip(table.columns, widths):
                    col.width = base.Mm(width)
                for row in table.rows:
                    for cell, width in zip(row.cells, widths):
                        cell.width = base.Mm(width)
        self.table_records.append({'caption': self.table_caption, 'rows': len(rows), 'numeric': numeric,
                                   'parts': len(chunks), 'row_keys': [row[0] for row in rows]})

    def algorithm(self, code, caption=''):
        visible = [line.rstrip() for line in code if line.strip()]
        manual_pattern = re.compile(r'^(\s*)(\d{1,3})(?:[.)、])?([ \t]+)(.*)$')
        manual = [manual_pattern.match(line) for line in visible]
        separators = [len(match[3]) for match in manual if match]
        base_separator = min(separators) if separators else 0
        number = 0
        for index, (line, match) in enumerate(zip(visible, manual)):
            p = self.doc.add_paragraph(style='Algorithm Line')
            fmt = p.paragraph_format
            fmt.left_indent = base.Mm(7)
            fmt.first_line_indent = base.Mm(-7)
            fmt.line_spacing = base.Pt(15)
            fmt.space_after = base.Pt(0)
            fmt.keep_with_next = index < len(visible) - 1 and len(visible) <= 40
            fmt.tab_stops.add_tab_stop(base.Mm(7))
            source = line
            if match:
                source = ' ' * (len(match[3]) - base_separator) + match[4]
            is_io = bool(re.match(r'^\s*(输入|输出|Input|Output)[：:]', source))
            if is_io or not match and separators:
                text = '\t' + source
            elif match:
                number = int(match[2])
                text = f'{number}\t' + source
            else:
                number += 1
                text = f'{number}\t' + source
            self.inline(p, text, 10.5)
            if index in (0, len(visible) - 1):
                bd = node('w:pBdr')
                bd.append(node('w:top' if index == 0 else 'w:bottom', val='single', sz='4', color='000000', space='4'))
                p._p.get_or_add_pPr().append(bd)
        self.algorithm_records.append({'caption': caption, 'lines': number, 'total_lines': len(visible)})

    def body(self):
        chapter, eq_counter = '0', 0
        appendix = english = body_started = False
        cn = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6}
        pending_algorithm = ''
        for kind, value, line in self.blocks:
            try:
                if kind == 'heading':
                    level, title = value
                    match = re.match(r'^第([一二三四五六]+|\d+)章\s', title)
                    if match and not body_started:
                        self.toc()
                        self.body_section()
                        body_started = True
                    if level == 1:
                        english = title == 'Abstract'
                        app = re.match(r'^附录\s*([A-F])(?:\s|$)', title)
                        if match:
                            chapter, eq_counter = str(cn.get(match[1], match[1])), 0
                        elif app:
                            chapter, eq_counter, appendix = app[1], 0, True
                    p = self.doc.add_paragraph(base.clean_inline(title), style=f'Heading {level}' if body_started else 'Front Contents')
                    if title == '摘要' or match and str(cn.get(match[1], match[1])) == '1':
                        p.paragraph_format.page_break_before = False
                    self.register(p, title)
                    continue
                if kind == 'table':
                    self.table(value)
                    continue
                if kind == 'math':
                    previous = self.doc.paragraphs[-1] if self.doc.paragraphs else None
                    if previous is not None and len(previous.text) < 80 and re.search(r'(?:定义|写为|可得|为|即|如下|：|:)$', previous.text.strip()):
                        previous.paragraph_format.keep_with_next = True
                    eq_counter += 1
                    self.display_math(value, f'{chapter}-{eq_counter}')
                    continue
                if kind == 'code' and (pending_algorithm or value[0] == 'algorithm'):
                    self.algorithm(value[1], pending_algorithm)
                    pending_algorithm = ''
                    continue
                p = self.doc.add_paragraph()
                fmt = p.paragraph_format
                if appendix:
                    fmt.line_spacing = base.Pt(19)
                if english:
                    fmt.line_spacing = base.Pt(16)
                    fmt.space_after = base.Pt(5)
                if kind == 'code':
                    fmt.first_line_indent = base.Pt(0)
                    fmt.left_indent = base.Mm(3)
                    fmt.line_spacing = 1.15
                    fmt.keep_together = len(value[1]) <= 15
                    base.font_run(p.add_run('\n'.join(value[1])), 10.5, latin='Courier New')
                    continue
                text = value
                label = base.clean_inline(text).replace('**', '')
                caption = CAPTION.fullmatch(label)
                if caption:
                    typ = {'图': 'Figure', '表': 'Table', '算法': 'Algorithm'}[caption[1]]
                    p.style = self.doc.styles[typ + ' Caption']
                    self.inline(p, text, 10.5)
                    self.register(p, label)
                    if caption[1] == '表':
                        self.table_caption = label
                    elif caption[1] == '算法':
                        pending_algorithm = label
                elif re.match(r'^\[\d+\] ', text):
                    fmt.first_line_indent, fmt.left_indent = base.Mm(-8), base.Mm(8)
                    fmt.line_spacing, fmt.space_after = base.Pt(14), base.Pt(6)
                    fmt.keep_together = True
                    self.inline(p, text, 10.5)
                else:
                    if re.match(r'^(?:\d+\.|[-*] )', text) or text.startswith(('关键词', 'Keywords')):
                        fmt.first_line_indent = base.Pt(0)
                    self.inline(p, text, 11 if appendix or text.startswith(('关键词', 'Keywords')) else None)
        
            except Exception as exc:
                raise ValueError(f'Body line {line}: {exc}') from exc

    def save(self, output):
        self.doc.core_properties.title = base.TITLE
        self.doc.core_properties.subject = WARNING
        self.doc.core_properties.author = 'Chen Tang'
        self.doc.core_properties.keywords = '关系感知学习 依赖闭合修正 非实测占位 勿送审'
        update = self.doc.settings.element.find(base.qn('w:updateFields'))
        if update is None:
            update = node('w:updateFields')
            self.doc.settings.element.append(update)
        update.set(base.qn('w:val'), 'true')
        for p in self.doc._element.xpath('.//w:p'):
            prop = p.find(base.qn('w:pPr'))
            if prop is not None and p.index(prop):
                p.remove(prop)
                p.insert(0, prop)
        self.doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    parser.add_argument('--check-only', action='store_true')
    args = parser.parse_args()
    assert Path(sys.executable).resolve() == base.BUNDLED_PYTHON.resolve()
    assert args.output.resolve() == OUTPUT or args.output.resolve().is_relative_to(WORK)
    old_hash = digest(TEMPLATE)
    lines = base.load_manuscript(args.source)
    counts = base.preflight(lines)
    base.self_test()
    if args.check_only:
        print(json.dumps(counts, ensure_ascii=False))
        return
    cache = WORK / 'page_map.json'
    pages = json.loads(cache.read_text()) if cache.exists() else {}
    builder = Builder(TEMPLATE, lines, pages)
    builder.cover()
    builder.body()
    builder.save(args.output)
    assert digest(TEMPLATE) == old_hash
    result = {'output': str(args.output), 'source_sha256': digest(args.source),
              'retained_template_sha256': old_hash, 'native_equations': builder.math_count,
              'tables': builder.table_records, 'display_equations': builder.equation_records,
              'algorithms': builder.algorithm_records, 'images': len(builder.doc.inline_shapes),
              'headings': builder.headings, 'captions': builder.captions,
              'page_cache_loaded': bool(pages), 'rendered_or_visually_verified': False,
              'counts': counts}
    (WORK / 'docx_build.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in result.items() if k not in ('tables', 'display_equations', 'headings', 'captions')}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
